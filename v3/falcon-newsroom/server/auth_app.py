import base64
import csv
import hashlib
import os
import re
import secrets
import zipfile
from io import BytesIO, StringIO
from datetime import timedelta, timezone, datetime
from functools import wraps
from pathlib import Path
from time import monotonic
from urllib.parse import urlencode, urlparse

import gridfs
import requests
from bson import ObjectId
from cryptography.fernet import Fernet, InvalidToken
from dotenv import load_dotenv
from flask import Flask, jsonify, redirect, request, send_file, session, url_for
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from pymongo import ASCENDING, ReturnDocument
from pymongo.errors import DuplicateKeyError
from pymongo.mongo_client import MongoClient
from pymongo.server_api import ServerApi
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash, generate_password_hash

try:
    from .article_extractor import (
        ArticleExtractionError,
        canonicalize_article_url,
        extract_article,
        normalize_publication_host,
    )
except ImportError:  # Support direct execution during local diagnostics.
    from article_extractor import (
        ArticleExtractionError,
        canonicalize_article_url,
        extract_article,
        normalize_publication_host,
    )


def _repo_root() -> Path:
    for parent in Path(__file__).resolve().parents:
        if (parent / ".env").exists() or (parent / ".env.example").exists() or (parent / ".git").exists():
            return parent
    return Path(__file__).resolve().parents[3]


load_dotenv(_repo_root() / ".env")

FLASK_ENV = (os.getenv("FLASK_ENV") or "production").strip().lower()
MONGO_URI = os.getenv("MONGO_URI", "")
if not MONGO_URI:
    raise RuntimeError("MONGO_URI is not set. Add it to your .env file.")

MONGO_DB = os.getenv("MONGO_DB", "mudu1735")
USER_COLLECTION = os.getenv("USER_COLLECTION", "loginInfov2")
INTERVIEW_COLLECTION = os.getenv("INTERVIEW_COLLECTION", "interviewRecords")
ARTICLE_COLLECTION = os.getenv("ARTICLE_COLLECTION", "articleRecords")
STORY_COLLECTION = os.getenv("STORY_COLLECTION", "storyRecords")
PITCH_COLLECTION = os.getenv("PITCH_COLLECTION", "pitchRecords")
ACTIVITY_COLLECTION = os.getenv("ACTIVITY_COLLECTION", "activityRecords")
FEEDBACK_COLLECTION = os.getenv("FEEDBACK_COLLECTION", "feedbackRecords")
WORKSPACE_COLLECTION = os.getenv("WORKSPACE_COLLECTION", "newsroomWorkspaces")
NAMES_COLLECTION = os.getenv("NAMES_COLLECTION", "names")
OAUTH_STATE_COLLECTION = os.getenv("OAUTH_STATE_COLLECTION", "googleOAuthStateNonces")
EXTRACTION_RATE_COLLECTION = os.getenv("EXTRACTION_RATE_COLLECTION", "articleExtractionRateLimits")
SECURITY_RATE_COLLECTION = os.getenv("SECURITY_RATE_COLLECTION", "securityRateLimits")
ADMIN_MUTATION_LOCK_COLLECTION = os.getenv("ADMIN_MUTATION_LOCK_COLLECTION", "adminMutationLocks")
AUTH_RATE_LIMIT_MAX = int(os.getenv("AUTH_RATE_LIMIT_MAX", "8"))
AUTH_RATE_LIMIT_WINDOW_SECONDS = int(os.getenv("AUTH_RATE_LIMIT_WINDOW_SECONDS", "900"))
AUTH_ACCOUNT_RATE_LIMIT_MAX = int(os.getenv("AUTH_ACCOUNT_RATE_LIMIT_MAX", "20"))
AUTH_IP_RATE_LIMIT_MAX = int(os.getenv("AUTH_IP_RATE_LIMIT_MAX", "500"))
REGISTRATION_RATE_LIMIT_MAX = int(os.getenv("REGISTRATION_RATE_LIMIT_MAX", "100"))
REGISTRATION_RATE_LIMIT_WINDOW_SECONDS = int(os.getenv("REGISTRATION_RATE_LIMIT_WINDOW_SECONDS", "3600"))
OAUTH_START_RATE_LIMIT_MAX = int(os.getenv("OAUTH_START_RATE_LIMIT_MAX", "20"))
OAUTH_START_IP_RATE_LIMIT_MAX = int(os.getenv("OAUTH_START_IP_RATE_LIMIT_MAX", "500"))
OAUTH_START_RATE_LIMIT_WINDOW_SECONDS = int(os.getenv("OAUTH_START_RATE_LIMIT_WINDOW_SECONDS", "900"))
JOIN_RATE_LIMIT_MAX = int(os.getenv("JOIN_RATE_LIMIT_MAX", "20"))
JOIN_IP_RATE_LIMIT_MAX = int(os.getenv("JOIN_IP_RATE_LIMIT_MAX", "500"))
JOIN_RATE_LIMIT_WINDOW_SECONDS = int(os.getenv("JOIN_RATE_LIMIT_WINDOW_SECONDS", "900"))
EXTRACTION_RATE_LIMIT_MAX = int(os.getenv("EXTRACTION_RATE_LIMIT_MAX", "10"))
EXTRACTION_RATE_LIMIT_WINDOW_SECONDS = int(os.getenv("EXTRACTION_RATE_LIMIT_WINDOW_SECONDS", "900"))
EXTRACTION_TOKEN_MAX_AGE_SECONDS = int(os.getenv("EXTRACTION_TOKEN_MAX_AGE_SECONDS", "900"))
MAX_STORY_ATTACHMENT_BYTES = int(os.getenv("MAX_STORY_ATTACHMENT_BYTES", str(10 * 1024 * 1024)))
MAX_NAMES_UPLOAD_BYTES = int(os.getenv("MAX_NAMES_UPLOAD_BYTES", str(10 * 1024 * 1024)))
MAX_REQUEST_BYTES = int(os.getenv("MAX_REQUEST_BYTES", str(12 * 1024 * 1024)))
MAX_STORY_ATTACHMENTS = int(os.getenv("MAX_STORY_ATTACHMENTS", "20"))
MAX_STORY_STORAGE_BYTES = int(os.getenv("MAX_STORY_STORAGE_BYTES", str(50 * 1024 * 1024)))
MAX_DRIVE_SHARE_RECIPIENTS = int(os.getenv("MAX_DRIVE_SHARE_RECIPIENTS", "50"))
DRIVE_SHARE_TOTAL_TIMEOUT_SECONDS = float(os.getenv("DRIVE_SHARE_TOTAL_TIMEOUT_SECONDS", "15"))
TRUST_PROXY_HEADERS = os.getenv("TRUST_PROXY_HEADERS", "0").strip().lower() in {"1", "true", "yes", "on"}

GOOGLE_CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID", "")
GOOGLE_CLIENT_SECRET = os.getenv("GOOGLE_CLIENT_SECRET", "")
OAUTH_TOKEN_ENCRYPTION_KEY = os.getenv("OAUTH_TOKEN_ENCRYPTION_KEY", "")
GOOGLE_AUTH_REDIRECT_URI = os.getenv("GOOGLE_AUTH_REDIRECT_URI", "")
FRONTEND_ORIGIN = os.getenv("FRONTEND_ORIGIN", "").rstrip("/")
DEFAULT_DEV_FRONTEND_ORIGIN = "http://127.0.0.1:5173"
GOOGLE_DRIVE_SCOPE = "https://www.googleapis.com/auth/drive.file"
GOOGLE_OAUTH_SCOPES = ("openid", "email", "profile", GOOGLE_DRIVE_SCOPE)
GOOGLE_PICKER_API_KEY = os.getenv("GOOGLE_PICKER_API_KEY", "")
GOOGLE_PICKER_APP_ID = os.getenv("GOOGLE_PICKER_APP_ID", "")
GOOGLE_PICKER_CLIENT_ID = os.getenv("GOOGLE_PICKER_CLIENT_ID", GOOGLE_CLIENT_ID)
GOOGLE_ALLOWED_DOMAINS = {
    domain.strip().lower().lstrip("@")
    for domain in os.getenv("GOOGLE_ALLOWED_DOMAINS", "").split(",")
    if domain.strip()
}
GOOGLE_AUTH_ENDPOINT = "https://accounts.google.com/o/oauth2/v2/auth"
GOOGLE_TOKEN_ENDPOINT = "https://oauth2.googleapis.com/token"
GOOGLE_USERINFO_ENDPOINT = "https://openidconnect.googleapis.com/v1/userinfo"
GOOGLE_DRIVE_FILES_ENDPOINT = "https://www.googleapis.com/drive/v3/files"
GOOGLE_INTENT_LOGIN = "login"
GOOGLE_INTENT_SIGNUP = "signup"
GOOGLE_INTENTS = {GOOGLE_INTENT_LOGIN, GOOGLE_INTENT_SIGNUP}

ROLE_GUEST = "guest"
ROLE_WRITER = "writer"
ROLE_EDITOR = "editor"
ROLE_ADMIN = "admin"
VALID_ROLES = {ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER, ROLE_GUEST}
ROLE_SCHEMA_CAPABILITY = "guest-role-v1"
PITCH_SECTIONS = {
    "News",
    "Features",
    "Sports",
    "Culture",
    "Opinion",
    "Science & Technology",
    "Photo",
}
MAX_PITCH_TITLE_LENGTH = 240
MAX_PITCH_ANGLE_LENGTH = 4000
MAX_PITCH_SECTION_LENGTH = 80
MAX_PITCH_NOTES_LENGTH = 10000
VALID_STORY_STATUSES = {
    "Assigned",
    "Reporting",
    "Drafting",
    "Submitted",
    "In Review",
    "Needs Revision",
    "Returned",
    "Ready for Publish",
    "Published",
}
WRITER_STORY_CONTENT_STATUSES = {
    "Assigned",
    "Reporting",
    "Drafting",
    "Returned",
    "Needs Revision",
}
EDITOR_STORY_CONTENT_STATUSES = WRITER_STORY_CONTENT_STATUSES | {"Submitted", "In Review"}
LOCKED_COLLABORATOR_STATUSES = {"Ready for Publish", "Published"}
WORKSPACE_PLATFORM_TYPES = {
    "SNO Sites / WordPress",
    "SNO Sites",
    "WordPress",
    "Other",
}
ROLE_RANK = {
    ROLE_GUEST: 1,
    ROLE_WRITER: 2,
    ROLE_EDITOR: 3,
    ROLE_ADMIN: 4,
}
FRONTEND_ROLE_LANDING = {
    ROLE_GUEST: "/interviewees",
    ROLE_WRITER: "/stories",
    ROLE_EDITOR: "/dashboard",
    ROLE_ADMIN: "/dashboard",
}
IMPORTANT_ACTIVITY_EVENTS = {"status_change"}
BACKEND_CAPABILITIES = {
    "admin-users",
    ROLE_SCHEMA_CAPABILITY,
    "rbac-v4",
    "stories",
    "pitches",
    "pitch-owner-submit",
    "shared-workflow-activity",
    "shared-feedback",
    "personal-dashboard",
    "story-invitations",
    "multi-workspace",
    "workspace-join-codes",
    "workspace-settings",
    "names-database",
}

TRUSTED_CSRF_ORIGINS = {
    origin.rstrip("/")
    for origin in os.getenv("TRUSTED_CSRF_ORIGINS", os.getenv("TRUSTED_ORIGINS", "")).split(",")
    if origin.strip()
}
if FRONTEND_ORIGIN:
    TRUSTED_CSRF_ORIGINS.add(FRONTEND_ORIGIN)
if FLASK_ENV != "production":
    TRUSTED_CSRF_ORIGINS.update({
        "http://127.0.0.1:5173",
        "http://127.0.0.1:5174",
        "http://localhost:5173",
        "http://localhost:5174",
    })

app = Flask(__name__)
_configured_secret = os.getenv("FLASK_SECRET_KEY", "").strip()
_known_insecure_secrets = {
    "dev-secret-change-me",
    "change-me",
    "changeme",
    "replace-with-a-random-secret",
    "secret",
}
_has_strong_secret = (
    len(_configured_secret) >= 32
    and len(set(_configured_secret)) >= 8
    and _configured_secret.lower() not in _known_insecure_secrets
)
_has_strong_oauth_encryption_key = (
    len(OAUTH_TOKEN_ENCRYPTION_KEY) >= 32
    and len(set(OAUTH_TOKEN_ENCRYPTION_KEY)) >= 8
    and OAUTH_TOKEN_ENCRYPTION_KEY.lower() not in _known_insecure_secrets
)
_has_persistent_oauth_encryption_key = bool(_has_strong_oauth_encryption_key or _has_strong_secret)
if FLASK_ENV == "production" and not _has_strong_secret:
    raise RuntimeError("FLASK_SECRET_KEY must be a unique secret of at least 32 characters in production.")
if FLASK_ENV == "production" and OAUTH_TOKEN_ENCRYPTION_KEY and not _has_strong_oauth_encryption_key:
    raise RuntimeError("OAUTH_TOKEN_ENCRYPTION_KEY must be a unique secret of at least 32 characters in production.")
app.secret_key = _configured_secret if _has_strong_secret else secrets.token_urlsafe(48)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = FLASK_ENV == "production"
app.config["SESSION_COOKIE_PATH"] = "/"
app.config["SESSION_COOKIE_NAME"] = "__Host-falcon_session" if FLASK_ENV == "production" else "falcon_session"
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=30)
app.config["MAX_CONTENT_LENGTH"] = MAX_REQUEST_BYTES

mongo_client = MongoClient(MONGO_URI, server_api=ServerApi("1"))
db = mongo_client[MONGO_DB]
users_col = db[USER_COLLECTION]
interviews_col = db[INTERVIEW_COLLECTION]
articles_col = db[ARTICLE_COLLECTION]
stories_col = db[STORY_COLLECTION]
pitches_col = db[PITCH_COLLECTION]
activity_col = db[ACTIVITY_COLLECTION]
feedback_col = db[FEEDBACK_COLLECTION]
workspaces_col = db[WORKSPACE_COLLECTION]
names_col = db[NAMES_COLLECTION]
oauth_state_col = db[OAUTH_STATE_COLLECTION]
extraction_rate_col = db[EXTRACTION_RATE_COLLECTION]
security_rate_col = db[SECURITY_RATE_COLLECTION]
admin_mutation_locks_col = db[ADMIN_MUTATION_LOCK_COLLECTION]
story_files = gridfs.GridFS(db, collection="storyAttachments")
AUTH_FAILURES: dict[str, list[float]] = {}
AUTH_FAILURES_MAX_KEYS = 10000
JOIN_FAILURES: dict[str, list[float]] = {}
GOOGLE_STATE_MAX_AGE_SECONDS = 600
DUMMY_PASSWORD_HASH = generate_password_hash(secrets.token_urlsafe(32))

try:
    users_col.create_index([("email", ASCENDING)], unique=True)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    users_col.create_index([("googleSub", ASCENDING)], unique=True, sparse=True)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    users_col.create_index([("googleId", ASCENDING)], unique=True, sparse=True)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    workspaces_col.create_index([("publicId", ASCENDING)], unique=True)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    workspaces_col.create_index([("joinCode", ASCENDING)], unique=True)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    names_col.create_index([
        ("workspaceId", ASCENDING),
        ("uploadBatchId", ASCENDING),
        ("firstName", ASCENDING),
        ("lastName", ASCENDING),
    ])
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    oauth_state_col.create_index([("nonce", ASCENDING)], unique=True)
    oauth_state_col.create_index([("createdAt", ASCENDING)], expireAfterSeconds=GOOGLE_STATE_MAX_AGE_SECONDS)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    extraction_rate_col.create_index(
        [("workspaceId", ASCENDING), ("userId", ASCENDING), ("windowStart", ASCENDING)],
        unique=True,
    )
    extraction_rate_col.create_index([("expiresAt", ASCENDING)], expireAfterSeconds=0)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    security_rate_col.create_index(
        [("kind", ASCENDING), ("rateKey", ASCENDING), ("windowStart", ASCENDING)],
        unique=True,
    )
    security_rate_col.create_index([("expiresAt", ASCENDING)], expireAfterSeconds=0)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    admin_mutation_locks_col.create_index([("workspaceId", ASCENDING)], unique=True)
    admin_mutation_locks_col.create_index([("expiresAt", ASCENDING)], expireAfterSeconds=0)
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    stories_col.create_index([("writerEmail", ASCENDING), ("status", ASCENDING)])
except Exception:
    pass

try:
    stories_col.create_index([("collaborators.email", ASCENDING), ("status", ASCENDING)])
except Exception:
    pass

try:
    stories_col.create_index(
        [("workspaceId", ASCENDING), ("sourcePitchId", ASCENDING)],
        unique=True,
        partialFilterExpression={"sourcePitchId": {"$type": "string"}},
    )
except Exception:
    if FLASK_ENV == "production":
        raise


try:
    pitches_col.create_index([("ownerEmail", ASCENDING), ("status", ASCENDING)])
except Exception:
    pass

try:
    articles_col.create_index(
        [("workspaceId", ASCENDING), ("sourceStoryId", ASCENDING)],
        unique=True,
        partialFilterExpression={"sourceStoryId": {"$type": "string"}},
    )
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    articles_col.create_index(
        [("workspaceId", ASCENDING), ("urlNorm", ASCENDING)],
        unique=True,
        partialFilterExpression={"urlNorm": {"$type": "string"}},
    )
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    interviews_col.create_index(
        [
            ("workspaceId", ASCENDING),
            ("urlNorm", ASCENDING),
            ("firstNameNorm", ASCENDING),
            ("lastNameNorm", ASCENDING),
        ],
        unique=True,
        partialFilterExpression={
            "workspaceId": {"$type": "string"},
            "urlNorm": {"$type": "string"},
            "firstNameNorm": {"$type": "string"},
            "lastNameNorm": {"$type": "string"},
        },
    )
except Exception:
    if FLASK_ENV == "production":
        raise

try:
    activity_col.create_index([("entityType", ASCENDING), ("entityId", ASCENDING), ("createdAt", -1)])
except Exception:
    pass

try:
    feedback_col.create_index([("entityType", ASCENDING), ("entityId", ASCENDING), ("createdAt", -1)])
except Exception:
    pass


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds")


def _month_day_year(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        dt = value.astimezone(timezone.utc)
    else:
        raw = str(value or "").strip()
        if not raw:
            return ""
        normalized = raw.replace("Z", "+00:00")
        try:
            dt = datetime.fromisoformat(normalized)
        except Exception:
            for fmt in ("%B %d, %Y, %I:%M %p", "%B %d, %Y", "%b %d, %Y", "%Y-%m-%d"):
                try:
                    dt = datetime.strptime(raw, fmt).replace(tzinfo=timezone.utc)
                    break
                except Exception:
                    dt = None
            if dt is None:
                parsed = raw.split("T", 1)[0]
                if re.match(r"^\d{4}-\d{2}-\d{2}$", parsed):
                    try:
                        dt = datetime.strptime(parsed, "%Y-%m-%d").replace(tzinfo=timezone.utc)
                    except Exception:
                        return raw
                else:
                    return raw
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        else:
            dt = dt.astimezone(timezone.utc)
    return dt.strftime("%B %d, %Y").replace(" 0", " ")


def normalize_email(email: str) -> str:
    return (email or "").strip().lower()


def normalize_role(role_value) -> str:
    role = str(role_value or "").strip().lower()
    return role if role in VALID_ROLES else ROLE_GUEST


def normalize_workspace_code(value) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def _new_workspace_code() -> str:
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    for _attempt in range(20):
        code = "".join(secrets.choice(alphabet) for _ in range(7))
        if not workspaces_col.find_one({"joinCode": code}):
            return code
    raise RuntimeError("Could not allocate a unique workspace join code.")


def _serialize_workspace(doc: dict | None, include_join_code: bool = False) -> dict | None:
    if not doc:
        return None
    workspace = {
        "id": str(doc.get("publicId") or doc.get("_id") or ""),
        "name": str(doc.get("name") or "Workspace"),
        "publicationUrl": str(doc.get("publicationUrl") or ""),
        "articleDomain": str(doc.get("articleDomain") or ""),
        "platformType": str(doc.get("platformType") or "SNO Sites / WordPress"),
    }
    if include_join_code:
        workspace["joinCode"] = str(doc.get("joinCode") or "")
    return workspace


def _workspace_id_for_user(user_doc: dict | None) -> str:
    return str((user_doc or {}).get("workspaceId") or "").strip()


def _workspace_for_user(user_doc: dict | None):
    workspace_id = _workspace_id_for_user(user_doc)
    if not workspace_id:
        return None
    clauses = [{"publicId": workspace_id}]
    try:
        clauses.append({"_id": ObjectId(workspace_id)})
    except Exception:
        pass
    return workspaces_col.find_one({"$or": clauses})


def _workspace_query(user_doc: dict | None) -> dict:
    workspace_id = _workspace_id_for_user(user_doc)
    return {"workspaceId": workspace_id} if workspace_id else {"_id": None}


def _scoped_query(user_doc: dict | None, query: dict | None = None) -> dict:
    workspace = _workspace_query(user_doc)
    if not query:
        return workspace
    return {"$and": [workspace, query]}


def _ensure_default_workspace():
    public_id = os.getenv("DEFAULT_WORKSPACE_ID", "poolesville-pulse").strip() or "poolesville-pulse"
    existing = workspaces_col.find_one({"publicId": public_id})
    if not existing:
        now_iso = _now_iso()
        doc = {
            "publicId": public_id,
            "name": os.getenv("DEFAULT_WORKSPACE_NAME", "Poolesville Pulse").strip() or "Poolesville Pulse",
            "joinCode": normalize_workspace_code(os.getenv("DEFAULT_WORKSPACE_JOIN_CODE")) or _new_workspace_code(),
            "publicationUrl": os.getenv("DEFAULT_PUBLICATION_URL", "https://poolesvillepulse.org").strip(),
            "articleDomain": os.getenv("DEFAULT_ARTICLE_DOMAIN", "poolesvillepulse.org").strip().lower(),
            "platformType": os.getenv("DEFAULT_PLATFORM_TYPE", "SNO Sites / WordPress").strip(),
            "createdAt": now_iso,
            "updatedAt": now_iso,
        }
        try:
            result = workspaces_col.insert_one(doc)
            doc["_id"] = result.inserted_id
            existing = doc
        except DuplicateKeyError:
            existing = workspaces_col.find_one({"publicId": public_id})

    if existing and not existing.get("legacyMembershipMigratedAt"):
        users_col.update_many({"workspaceId": {"$exists": False}}, {"$set": {"workspaceId": public_id}})
        for collection in (interviews_col, articles_col, stories_col, pitches_col, activity_col, feedback_col):
            collection.update_many({"workspaceId": {"$exists": False}}, {"$set": {"workspaceId": public_id}})
        workspaces_col.update_one({"_id": existing["_id"]}, {"$set": {"legacyMembershipMigratedAt": _now_iso()}})

    if existing and not existing.get("namesDatabaseLegacyMigratedAt"):
        migrated_at = _now_iso()
        names_metadata = existing.get("namesDatabase") if isinstance(existing.get("namesDatabase"), dict) else {}
        active_batch_id = str(names_metadata.get("activeBatchId") or f"legacy::{public_id}")[:160]
        legacy_query = {"$or": [
            {"workspaceId": {"$exists": False}},
            {"workspaceId": None},
            {"workspaceId": ""},
        ]}
        legacy_count = names_col.count_documents(legacy_query)
        if legacy_count:
            names_col.update_many(legacy_query, {"$set": {
                "workspaceId": public_id,
                "uploadBatchId": active_batch_id,
                "uploadedAt": migrated_at,
            }})
            names_metadata = {
                **names_metadata,
                "activeBatchId": active_batch_id,
                "count": legacy_count,
                "updatedAt": str(names_metadata.get("updatedAt") or migrated_at),
            }
        workspace_update = {"namesDatabaseLegacyMigratedAt": migrated_at}
        if legacy_count:
            workspace_update["namesDatabase"] = names_metadata
        workspaces_col.update_one({"_id": existing["_id"]}, {"$set": workspace_update})


def _canonicalize_workspace_user_roles():
    """Keep every workspace membership on the current v3 role vocabulary."""
    users_col.update_many(
        {
            "workspaceId": {"$exists": True, "$nin": [None, ""]},
            "role": {"$nin": sorted(VALID_ROLES)},
        },
        {"$set": {"role": ROLE_GUEST}},
    )


_ensure_default_workspace()
_canonicalize_workspace_user_roles()


def is_valid_email(email: str) -> bool:
    candidate = normalize_email(email)
    if not candidate or len(candidate) > 254:
        return False
    return bool(re.match(r"^[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+$", candidate))


def _sanitize_next_url(next_url: str) -> str:
    raw = str(next_url or "")
    if raw != raw.strip() or "\\" in raw or any(ord(char) < 32 or ord(char) == 127 for char in raw):
        return ""
    candidate = raw
    if not candidate:
        return ""
    parsed = urlparse(candidate)
    if parsed.scheme or parsed.netloc or not candidate.startswith("/") or candidate.startswith("//"):
        return ""
    if candidate in {"/login", "/signup"}:
        return ""
    if candidate == "/":
        return "/dashboard"
    if candidate == "/records":
        return "/interviewees"
    return candidate


def _auth_redirect_for_role(role_value: str) -> str:
    next_url = _sanitize_next_url(request.args.get("next") or "")
    return next_url or FRONTEND_ROLE_LANDING.get(normalize_role(role_value), "/interviewees")


def _request_payload() -> dict:
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict) or not payload:
        payload = request.form.to_dict(flat=True)
    return payload if isinstance(payload, dict) else {}


def _payload_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "on"}


def _origin_from_url(value: str) -> str:
    raw = (value or "").strip()
    if not raw:
        return ""
    parsed = urlparse(raw)
    if not parsed.scheme or not parsed.netloc:
        return ""
    return f"{parsed.scheme.lower()}://{parsed.netloc.lower()}".rstrip("/")


def _trusted_request_origins() -> set[str]:
    origins = {origin.lower().rstrip("/") for origin in TRUSTED_CSRF_ORIGINS if origin}
    current = _origin_from_url(request.host_url)
    if current:
        origins.add(current)
    return origins


def _is_same_origin_request() -> bool:
    origin = (request.headers.get("Origin") or "").strip()
    referer = (request.headers.get("Referer") or "").strip()
    candidate = origin or referer
    if not candidate:
        return False
    return _origin_from_url(candidate) in _trusted_request_origins()


def _is_allowed_auth_request_origin() -> bool:
    fetch_site = (request.headers.get("Sec-Fetch-Site") or "").strip().lower()
    if fetch_site == "cross-site":
        return False
    if request.headers.get("Origin") or request.headers.get("Referer"):
        return _is_same_origin_request()
    return True


def _ensure_csrf_token() -> str:
    token = session.get("csrf_token")
    if not isinstance(token, str) or len(token) < 32:
        token = secrets.token_urlsafe(32)
        session["csrf_token"] = token
    return token


def _has_valid_csrf_token() -> bool:
    expected = session.get("csrf_token")
    provided = request.headers.get("X-CSRF-Token") or request.headers.get("X-CSRFToken") or ""
    if not isinstance(expected, str) or not provided:
        return False
    return secrets.compare_digest(expected, str(provided))


@app.before_request
def csrf_guard_for_authenticated_users():
    if request.method not in {"POST", "PUT", "PATCH", "DELETE"}:
        return None
    if request.path in {"/api/auth/login", "/api/auth/register"}:
        if not _is_allowed_auth_request_origin():
            return jsonify({"ok": False, "error": "Request origin is not allowed."}), 403
        return None
    if session.get("user_id") and not (_has_valid_csrf_token() or _is_same_origin_request()):
        return jsonify({"ok": False, "error": "CSRF validation failed."}), 403
    return None


@app.after_request
def add_api_security_headers(response):
    if request.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["Content-Security-Policy"] = "default-src 'none'; frame-ancestors 'none'; base-uri 'none'"
    if FLASK_ENV == "production":
        response.headers.setdefault("Strict-Transport-Security", "max-age=31536000; includeSubDomains")
    return response


@app.errorhandler(413)
def request_too_large(_error):
    return jsonify({"ok": False, "error": "Request body is too large."}), 413


def _client_ip() -> str:
    remote = (request.remote_addr or "unknown").strip()
    if TRUST_PROXY_HEADERS:
        remote = (request.headers.get("X-Forwarded-For", "").split(",")[0] or remote).strip()
    return remote[:128]


def _auth_rate_key(email: str) -> str:
    return f"{_client_ip()}:{normalize_email(email)}"


def _security_rate_key(value: str) -> str:
    material = f"{app.secret_key}:{value}".encode("utf-8", errors="ignore")
    return hashlib.sha256(material).hexdigest()


def _security_rate_window(window_seconds: int, now: datetime | None = None) -> tuple[datetime, datetime]:
    seconds = max(60, min(int(window_seconds), 86400))
    current = now or datetime.now(timezone.utc)
    if current.tzinfo is None:
        current = current.replace(tzinfo=timezone.utc)
    epoch = int(current.timestamp())
    window_epoch = epoch - (epoch % seconds)
    start = datetime.fromtimestamp(window_epoch, tz=timezone.utc)
    return start, start + timedelta(seconds=seconds * 2)


def _shared_rate_count(kind: str, key: str, window_seconds: int) -> int:
    window_start, _expires_at = _security_rate_window(window_seconds)
    bucket = security_rate_col.find_one({
        "kind": kind,
        "rateKey": _security_rate_key(key),
        "windowStart": window_start,
    }) or {}
    return int(bucket.get("count") or 0)


def _record_shared_rate_failure(kind: str, key: str, window_seconds: int):
    window_start, expires_at = _security_rate_window(window_seconds)
    query = {
        "kind": kind,
        "rateKey": _security_rate_key(key),
        "windowStart": window_start,
    }
    update = {
        "$inc": {"count": 1},
        "$setOnInsert": {"createdAt": datetime.now(timezone.utc), "expiresAt": expires_at},
    }
    try:
        bucket = security_rate_col.find_one_and_update(
            query, update, upsert=True, return_document=ReturnDocument.AFTER,
        ) or {}
    except DuplicateKeyError:
        bucket = security_rate_col.find_one_and_update(
            query, {"$inc": {"count": 1}}, return_document=ReturnDocument.AFTER,
        ) or {}
    return int(bucket.get("count") or 0)


def _clear_shared_rate_failures(kind: str, key: str):
    security_rate_col.delete_many({"kind": kind, "rateKey": _security_rate_key(key)})


def _acquire_admin_mutation_lock(workspace_id: str) -> str:
    token = secrets.token_urlsafe(24)
    now = datetime.now(timezone.utc)
    try:
        lock = admin_mutation_locks_col.find_one_and_update(
            {
                "workspaceId": workspace_id,
                "$or": [
                    {"expiresAt": {"$lte": now}},
                    {"expiresAt": {"$exists": False}},
                ],
            },
            {"$set": {"token": token, "acquiredAt": now, "expiresAt": now + timedelta(seconds=20)}},
            upsert=True,
            return_document=ReturnDocument.AFTER,
        )
    except DuplicateKeyError:
        return ""
    except Exception:
        return ""
    return token if lock and secrets.compare_digest(str(lock.get("token") or ""), token) else ""


def _release_admin_mutation_lock(workspace_id: str, token: str):
    if token:
        admin_mutation_locks_col.delete_one({"workspaceId": workspace_id, "token": token})


def _active_auth_failures(key: str, now: float) -> list[float]:
    attempts = [ts for ts in AUTH_FAILURES.get(key, []) if now - ts < AUTH_RATE_LIMIT_WINDOW_SECONDS]
    if attempts:
        AUTH_FAILURES[key] = attempts
    else:
        AUTH_FAILURES.pop(key, None)
    return attempts


def _is_auth_rate_limited(email: str) -> bool:
    if AUTH_RATE_LIMIT_MAX <= 0:
        return False
    try:
        normalized_email = normalize_email(email)
        counts = (
            (_record_shared_rate_failure("auth", _auth_rate_key(email), AUTH_RATE_LIMIT_WINDOW_SECONDS), AUTH_RATE_LIMIT_MAX),
            (_record_shared_rate_failure("auth-account", normalized_email, AUTH_RATE_LIMIT_WINDOW_SECONDS), AUTH_ACCOUNT_RATE_LIMIT_MAX),
            (_record_shared_rate_failure("auth-ip", _client_ip(), AUTH_RATE_LIMIT_WINDOW_SECONDS), AUTH_IP_RATE_LIMIT_MAX),
        )
        return any(limit > 0 and count > limit for count, limit in counts)
    except Exception:
        if FLASK_ENV == "production":
            return True
    key = _auth_rate_key(email)
    attempts = _active_auth_failures(key, monotonic())
    attempts.append(monotonic())
    AUTH_FAILURES[key] = attempts
    return len(attempts) > AUTH_RATE_LIMIT_MAX


def _record_auth_failure(email: str):
    key = _auth_rate_key(email)
    try:
        _record_shared_rate_failure("auth", key, AUTH_RATE_LIMIT_WINDOW_SECONDS)
        return
    except Exception:
        if FLASK_ENV == "production":
            return
    now = monotonic()
    if key not in AUTH_FAILURES and len(AUTH_FAILURES) >= AUTH_FAILURES_MAX_KEYS:
        for existing_key in list(AUTH_FAILURES):
            _active_auth_failures(existing_key, now)
        while len(AUTH_FAILURES) >= AUTH_FAILURES_MAX_KEYS:
            AUTH_FAILURES.pop(next(iter(AUTH_FAILURES)), None)
    attempts = _active_auth_failures(key, now)
    attempts.append(now)
    AUTH_FAILURES[key] = attempts


def _clear_auth_failures(email: str):
    key = _auth_rate_key(email)
    try:
        _clear_shared_rate_failures("auth", key)
        _clear_shared_rate_failures("auth-account", normalize_email(email))
    except Exception:
        pass
    AUTH_FAILURES.pop(key, None)


def _active_join_failures(now: float) -> list[float]:
    key = _client_ip()
    attempts = [ts for ts in JOIN_FAILURES.get(key, []) if now - ts < JOIN_RATE_LIMIT_WINDOW_SECONDS]
    if attempts:
        JOIN_FAILURES[key] = attempts
    else:
        JOIN_FAILURES.pop(key, None)
    return attempts


def _record_join_failure():
    key = _client_ip()
    try:
        _record_shared_rate_failure("join", key, JOIN_RATE_LIMIT_WINDOW_SECONDS)
        return
    except Exception:
        if FLASK_ENV == "production":
            return
    now = monotonic()
    attempts = _active_join_failures(now)
    attempts.append(now)
    JOIN_FAILURES[key] = attempts
    while len(JOIN_FAILURES) > AUTH_FAILURES_MAX_KEYS:
        JOIN_FAILURES.pop(next(iter(JOIN_FAILURES)), None)


def _is_join_rate_limited(user_doc: dict | None = None) -> bool:
    if JOIN_RATE_LIMIT_MAX <= 0:
        return False
    try:
        user_key = str((user_doc or {}).get("_id") or _client_ip())
        counts = (
            (_record_shared_rate_failure("join-user", user_key, JOIN_RATE_LIMIT_WINDOW_SECONDS), JOIN_RATE_LIMIT_MAX),
            (_record_shared_rate_failure("join-ip", _client_ip(), JOIN_RATE_LIMIT_WINDOW_SECONDS), JOIN_IP_RATE_LIMIT_MAX),
        )
        return any(limit > 0 and count > limit for count, limit in counts)
    except Exception:
        if FLASK_ENV == "production":
            return True
    attempts = _active_join_failures(monotonic())
    attempts.append(monotonic())
    JOIN_FAILURES[_client_ip()] = attempts
    return len(attempts) > JOIN_RATE_LIMIT_MAX


def _is_ip_action_rate_limited(kind: str, maximum: int, window_seconds: int) -> bool:
    if maximum <= 0:
        return False
    try:
        return _record_shared_rate_failure(kind, _client_ip(), window_seconds) > maximum
    except Exception:
        return FLASK_ENV == "production"


def _is_oauth_start_rate_limited() -> bool:
    browser_key = session.get("rate_limit_browser_id")
    if not isinstance(browser_key, str) or not 16 <= len(browser_key) <= 128:
        browser_key = secrets.token_urlsafe(24)
        session["rate_limit_browser_id"] = browser_key
    try:
        counts = (
            (_record_shared_rate_failure("oauth-browser", browser_key, OAUTH_START_RATE_LIMIT_WINDOW_SECONDS), OAUTH_START_RATE_LIMIT_MAX),
            (_record_shared_rate_failure("oauth-ip", _client_ip(), OAUTH_START_RATE_LIMIT_WINDOW_SECONDS), OAUTH_START_IP_RATE_LIMIT_MAX),
        )
        return any(limit > 0 and count > limit for count, limit in counts)
    except Exception:
        return FLASK_ENV == "production"


def find_user_by_email(email: str):
    return users_col.find_one({"email": normalize_email(email)})


def _google_id_from_doc(doc: dict) -> str:
    return str(doc.get("googleId") or doc.get("googleSub") or "").strip()


def find_user_by_google_id(google_id: str):
    value = str(google_id or "").strip()
    if not value:
        return None
    return users_col.find_one({"$or": [{"googleId": value}, {"googleSub": value}]})


def find_user_by_google_sub(google_sub: str):
    return find_user_by_google_id(google_sub)


def _serialize_user(doc: dict) -> dict:
    first_name = str(doc.get("firstName", "") or "").strip()
    last_name = str(doc.get("lastName", "") or "").strip()
    email = doc.get("email", "")
    name = str(doc.get("name") or f"{first_name} {last_name}".strip() or email).strip()
    return {
        "id": str(doc.get("_id")),
        "email": email,
        "firstName": first_name,
        "lastName": last_name,
        "name": name,
        "role": normalize_role(doc.get("role")) if _workspace_id_for_user(doc) else "",
        "workspaceId": _workspace_id_for_user(doc),
        "lastSeen": _month_day_year(doc.get("lastLoginAt") or doc.get("updatedAt") or doc.get("createdAt")),
    }


def _auth_version_for_doc(doc: dict | None) -> int:
    try:
        return max(1, int((doc or {}).get("authVersion") or 1))
    except (TypeError, ValueError):
        return 1


def _current_user_doc():
    user_id = session.get("user_id")
    if not user_id:
        return None
    try:
        oid = ObjectId(user_id)
    except Exception:
        session.clear()
        return None
    doc = users_col.find_one({"_id": oid})
    if not doc:
        session.clear()
        return None
    session_auth_version = session.get("auth_version")
    if not isinstance(session_auth_version, int) or session_auth_version != _auth_version_for_doc(doc):
        session.clear()
        return None
    return doc


def require_auth(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not _current_user_doc():
            return jsonify({"ok": False, "error": "Unauthorized"}), 401
        return fn(*args, **kwargs)

    return wrapper


def require_role_at_least(role_name: str):
    target = normalize_role(role_name)

    def decorator(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            user_doc = _current_user_doc()
            if not user_doc:
                return jsonify({"ok": False, "error": "Unauthorized"}), 401
            if not _workspace_id_for_user(user_doc):
                return jsonify({"ok": False, "error": "Join a workspace first."}), 403
            if ROLE_RANK.get(normalize_role(user_doc.get("role")), 1) < ROLE_RANK.get(target, 1):
                return jsonify({"ok": False, "error": "Editor access required."}), 403
            return fn(*args, **kwargs)

        return wrapper

    return decorator


def require_roles(*role_names: str):
    allowed = {normalize_role(role_name) for role_name in role_names}

    def decorator(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            user_doc = _current_user_doc()
            if not user_doc:
                return jsonify({"ok": False, "error": "Unauthorized"}), 401
            if not _workspace_id_for_user(user_doc):
                return jsonify({"ok": False, "error": "Join a workspace first."}), 403
            if normalize_role(user_doc.get("role")) not in allowed:
                return jsonify({"ok": False, "error": "Access denied."}), 403
            return fn(*args, **kwargs)

        return wrapper

    return decorator


def _current_user_role(user_doc: dict) -> str:
    return normalize_role(user_doc.get("role"))


def _user_display_name(user_doc: dict) -> str:
    first_name = str(user_doc.get("firstName") or "").strip()
    last_name = str(user_doc.get("lastName") or "").strip()
    return str(user_doc.get("name") or f"{first_name} {last_name}".strip() or user_doc.get("email") or "Newsroom user").strip()


def _owner_match_values(user_doc: dict) -> list:
    values = [str(user_doc.get("_id")), normalize_email(user_doc.get("email") or "")]
    try:
        values.append(user_doc["_id"])
    except Exception:
        pass
    return [value for value in values if value]


def _owned_story_query(user_doc: dict) -> dict:
    values = _owner_match_values(user_doc)
    return {"$or": [
        {"writerUserId": {"$in": values}},
        {"writerId": {"$in": values}},
        {"ownerUserId": {"$in": values}},
        {"ownerId": {"$in": values}},
        {"writerEmail": {"$in": values}},
        {"ownerEmail": {"$in": values}},
    ]}


def _collaborator_story_query(user_doc: dict) -> dict:
    values = _owner_match_values(user_doc)
    return {"collaborators": {"$elemMatch": {"$and": [
        {"status": "accepted"},
        {"$or": [
            {"userId": {"$in": values}},
            {"id": {"$in": values}},
            {"email": {"$in": values}},
        ]},
    ]}}}


def _editable_collaborator_story_query(user_doc: dict) -> dict:
    values = _owner_match_values(user_doc)
    return {"collaborators": {"$elemMatch": {
        "status": "accepted",
        "role": "edit",
        "$or": [
            {"userId": {"$in": values}},
            {"id": {"$in": values}},
            {"email": {"$in": values}},
        ],
    }}}


def _story_primary_owned_by_user(story: dict, user_doc: dict) -> bool:
    if not story or not user_doc:
        return False
    values = {str(value).lower() for value in _owner_match_values(user_doc) if value}
    fields = [
        story.get("writerUserId"),
        story.get("writerId"),
        story.get("ownerUserId"),
        story.get("ownerId"),
        normalize_email(story.get("writerEmail") or ""),
        normalize_email(story.get("ownerEmail") or ""),
    ]
    return any(str(value).lower() in values for value in fields if value)


def _story_collaborators(doc: dict) -> list[dict]:
    raw = doc.get("collaborators") if isinstance(doc.get("collaborators"), list) else []
    collaborators = []
    seen = set()
    for item in raw:
        if not isinstance(item, dict):
            continue
        email = normalize_email(item.get("email") or "")
        if not email or email in seen:
            continue
        seen.add(email)
        role = str(item.get("role") or "comment").strip().lower()
        if role not in {"comment", "edit"}:
            role = "comment"
        name = str(item.get("name") or email).strip()
        collaborators.append({
            "id": str(item.get("userId") or item.get("id") or email),
            "userId": str(item.get("userId") or item.get("id") or ""),
            "email": email,
            "name": name,
            "role": role,
            "status": str(item.get("status") or "invalid").strip().lower(),
            "message": str(item.get("message") or "").strip(),
            "invitedBy": str(item.get("invitedBy") or "").strip(),
            "invitedAt": _date_for_api(item.get("invitedAt")),
            "invitedAtIso": _datetime_for_api(item.get("invitedAt")),
        })
    return collaborators


def _story_collaborator_for_user(story: dict, user_doc: dict) -> dict | None:
    if not story or not user_doc:
        return None
    values = {str(value).lower() for value in _owner_match_values(user_doc) if value}
    for collaborator in _story_collaborators(story):
        candidate_values = {
            str(collaborator.get("id") or "").lower(),
            str(collaborator.get("userId") or "").lower(),
            normalize_email(collaborator.get("email") or ""),
        }
        if collaborator.get("status") == "accepted" and values.intersection({value for value in candidate_values if value}):
            return collaborator
    return None


def _can_manage_story_collaborators(story: dict, user_doc: dict) -> bool:
    if str(story.get("status") or "Assigned").strip() in LOCKED_COLLABORATOR_STATUSES:
        return False
    role = _current_user_role(user_doc)
    collaborator = _story_collaborator_for_user(story, user_doc)
    return role in {ROLE_ADMIN, ROLE_EDITOR} or (
        role == ROLE_WRITER
        and (_story_primary_owned_by_user(story, user_doc) or (collaborator is not None and collaborator.get("role") == "edit"))
    )


def _can_edit_story_content(story: dict, user_doc: dict) -> bool:
    role = _current_user_role(user_doc)
    status = str(story.get("status") or "Assigned").strip()
    if role in {ROLE_ADMIN, ROLE_EDITOR}:
        return status in EDITOR_STORY_CONTENT_STATUSES
    if role != ROLE_WRITER:
        return False
    if status not in WRITER_STORY_CONTENT_STATUSES:
        return False
    if _story_primary_owned_by_user(story, user_doc):
        return True
    collaborator = _story_collaborator_for_user(story, user_doc)
    return collaborator is not None and collaborator.get("role") == "edit"

def _owned_pitch_query(user_doc: dict) -> dict:
    values = _owner_match_values(user_doc)
    return {"$or": [
        {"ownerUserId": {"$in": values}},
        {"ownerId": {"$in": values}},
        {"writerUserId": {"$in": values}},
        {"writerId": {"$in": values}},
        {"ownerEmail": {"$in": values}},
        {"writerEmail": {"$in": values}},
        {"owner": {"$in": values}},
        {"writer": {"$in": values}},
    ]}


def _pitch_owned_by_user(pitch: dict, user_doc: dict) -> bool:
    if not pitch or not user_doc:
        return False
    values = {str(value).lower() for value in _owner_match_values(user_doc) if value}
    fields = [
        pitch.get("ownerUserId"),
        pitch.get("ownerId"),
        pitch.get("writerUserId"),
        pitch.get("writerId"),
        normalize_email(pitch.get("ownerEmail") or ""),
        normalize_email(pitch.get("writerEmail") or ""),
    ]
    return any(str(value).lower() in values for value in fields if value)

def _story_query_for_user(user_doc: dict):
    role = _current_user_role(user_doc)
    if not _workspace_id_for_user(user_doc) or role == ROLE_GUEST:
        return None
    if role == ROLE_WRITER:
        return _scoped_query(user_doc, {"$or": _owned_story_query(user_doc)["$or"] + [_collaborator_story_query(user_doc)]})
    return _workspace_query(user_doc)


def _editable_story_query(user_doc: dict):
    role = _current_user_role(user_doc)
    if not _workspace_id_for_user(user_doc) or role == ROLE_GUEST:
        return None
    if role == ROLE_WRITER:
        return _scoped_query(user_doc, {"$or": _owned_story_query(user_doc)["$or"] + [_editable_collaborator_story_query(user_doc)]})
    return _workspace_query(user_doc)


def _story_content_mutation_query(user_doc: dict, expected_status: str) -> dict:
    role = _current_user_role(user_doc)
    allowed_statuses = EDITOR_STORY_CONTENT_STATUSES if role in {ROLE_ADMIN, ROLE_EDITOR} else WRITER_STORY_CONTENT_STATUSES
    base_query = _editable_story_query(user_doc)
    if base_query is None or expected_status not in allowed_statuses:
        return {"_id": None}
    return {"$and": [base_query, {"status": expected_status}]}


def _story_collaborator_mutation_query(user_doc: dict, expected_status: str) -> dict:
    if expected_status in LOCKED_COLLABORATOR_STATUSES:
        return {"_id": None}
    return {"$and": [_editable_story_query(user_doc) or {"_id": None}, {"status": expected_status}]}


def _pitch_query_for_user(user_doc: dict):
    role = _current_user_role(user_doc)
    if not _workspace_id_for_user(user_doc) or role == ROLE_GUEST:
        return None
    if role == ROLE_WRITER:
        return _scoped_query(user_doc, _owned_pitch_query(user_doc))
    return _workspace_query(user_doc)


def _find_owned_story_or_404(story_id: str, user_doc: dict):
    query = _story_query_for_user(user_doc)
    if query is None:
        return None
    clauses = [{"_id": _object_id_or_none(story_id)}, {"id": story_id}, {"storyId": story_id}]
    clauses = [clause for clause in clauses if list(clause.values())[0]]
    final_query = {"$and": [query, {"$or": clauses}]} if query else {"$or": clauses}
    return stories_col.find_one(final_query)


def _find_owned_pitch_or_404(pitch_id: str, user_doc: dict):
    query = _pitch_query_for_user(user_doc)
    if query is None:
        return None
    clauses = [{"_id": _object_id_or_none(pitch_id)}, {"id": pitch_id}, {"pitchId": pitch_id}]
    clauses = [clause for clause in clauses if list(clause.values())[0]]
    final_query = {"$and": [query, {"$or": clauses}]} if query else {"$or": clauses}
    return pitches_col.find_one(final_query)


def _object_id_or_none(value: str):
    try:
        return ObjectId(value)
    except Exception:
        return None


def _login_user_doc(doc: dict, remember: bool = False):
    session.clear()
    session.permanent = remember
    session["user_id"] = str(doc["_id"])
    session["auth_version"] = _auth_version_for_doc(doc)
    session["csrf_token"] = secrets.token_urlsafe(32)


def _coerce_list_field(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str):
        return [part.strip() for part in value.replace("\n", ",").split(",") if part.strip()]
    return [str(value).strip()] if str(value).strip() else []


def _article_url_lookup_candidates(url: str) -> set[str]:
    raw = str(url or "").strip()
    if not raw:
        return set()
    candidates = {raw, raw.rstrip("/")}
    try:
        parsed = urlparse(raw)
        trimmed_path = (parsed.path or "").rstrip("/")
        candidates.add(parsed._replace(path=trimmed_path).geturl())
        if trimmed_path:
            candidates.add(parsed._replace(path=f"{trimmed_path}/").geturl())
    except Exception:
        pass
    return {candidate for candidate in candidates if candidate}


def _article_join_url_key(url: str) -> str:
    raw = str(url or "").strip()
    if not raw:
        return ""
    try:
        parsed = urlparse(raw)
        scheme = (parsed.scheme or "https").lower()
        host = (parsed.netloc or "").lower()
        path = (parsed.path or "").rstrip("/")
        query = f"?{parsed.query}" if parsed.query else ""
        return f"{scheme}://{host}{path}{query}"
    except Exception:
        return raw.rstrip("/")


def _canonical_article_url(url: str) -> str:
    safe_url = _safe_http_url(url)
    if not safe_url:
        return ""
    parsed = urlparse(safe_url)
    host = (parsed.hostname or "").lower().rstrip(".")
    if ":" in host and not host.startswith("["):
        host = f"[{host}]"
    port = parsed.port
    if port and not ((parsed.scheme.lower() == "http" and port == 80) or (parsed.scheme.lower() == "https" and port == 443)):
        host = f"{host}:{port}"
    path = (parsed.path or "/").rstrip("/") or "/"
    return parsed._replace(
        scheme=parsed.scheme.lower(),
        netloc=host,
        path=path,
        fragment="",
    ).geturl()


def _publication_url_conflict_query(workspace_id: str, source_story_id: str, publication_url: str) -> dict:
    canonical_url = _canonical_article_url(publication_url)
    candidates = sorted(_article_url_lookup_candidates(publication_url))
    url_clauses = [{"urlNorm": canonical_url}, {"canonicalUrl": canonical_url}]
    if candidates:
        url_clauses.extend([
            {"url": {"$in": candidates}},
            {"articleUrl": {"$in": candidates}},
            {"article_url": {"$in": candidates}},
        ])
    return {
        "workspaceId": str(workspace_id or "").strip(),
        "sourceStoryId": {"$ne": str(source_story_id or "").strip()},
        "$or": url_clauses,
    }


def _publication_url_conflicts(story: dict, publication_url: str) -> bool:
    query = _publication_url_conflict_query(
        story.get("workspaceId") or "",
        _doc_public_id(story, "storyId"),
        publication_url,
    )
    return bool(articles_col.find_one(query, {"_id": 1}))


def _interviewee_search_query(workspace_id: str, search: str) -> dict:
    escaped = re.escape(str(search or "").strip())
    rx = {"$regex": escaped, "$options": "i"}
    full_name_expressions = []
    for first_field, last_field in (("$firstName", "$lastName"), ("$first_name", "$last_name")):
        full_name_expressions.append({
            "$expr": {
                "$regexMatch": {
                    "input": {"$trim": {"input": {"$concat": [
                        {"$convert": {"input": first_field, "to": "string", "onError": "", "onNull": ""}},
                        " ",
                        {"$convert": {"input": last_field, "to": "string", "onError": "", "onNull": ""}},
                    ]}}},
                    "regex": escaped,
                    "options": "i",
                }
            }
        })
    return {
        "workspaceId": str(workspace_id or "").strip(),
        "$or": [
            {"firstName": rx},
            {"first_name": rx},
            {"lastName": rx},
            {"last_name": rx},
            {"name": rx},
            {"fullName": rx},
            *full_name_expressions,
        ],
    }


def _article_search_clause(search: str, interview_urls) -> dict:
    rx = {"$regex": re.escape(str(search or "").strip()), "$options": "i"}
    search_fields = ["title", "articleTitle", "author", "authors", "tags", "categories", "section", "category"]
    clauses = [{field: rx} for field in search_fields]
    urls = sorted({str(value or "").strip() for value in interview_urls or [] if str(value or "").strip()})
    if urls:
        clauses.extend([
            {"url": {"$in": urls}},
            {"articleUrl": {"$in": urls}},
            {"article_url": {"$in": urls}},
        ])
    return {"$or": clauses}


def _published_date_for_api(value) -> str:
    return _month_day_year(value)


def _interview_to_api(doc: dict) -> dict:
    first_name = str(doc.get("firstName") or doc.get("first_name") or "").strip()
    last_name = str(doc.get("lastName") or doc.get("last_name") or "").strip()
    name = str(doc.get("name") or doc.get("fullName") or f"{first_name} {last_name}").strip() or "Unknown"
    return {
        "id": str(doc.get("_id")),
        "firstName": first_name,
        "lastName": last_name,
        "name": name,
        "grade": doc.get("grade", "") or "",
        "house": doc.get("house", "") or "",
        "url": doc.get("url") or doc.get("articleUrl") or doc.get("article_url") or "",
        "dateAdded": _month_day_year(doc.get("dateAdded") or doc.get("date") or doc.get("dateInterviewed")),
        "time": _month_day_year(doc.get("time") or doc.get("dateAdded") or doc.get("date") or doc.get("dateInterviewed")),
        "addedBy": doc.get("addedBy") or doc.get("added_by") or doc.get("createdBy") or "",
        "quote": doc.get("quote") or doc.get("evidence") or doc.get("quoteEvidence") or "",
    }


def _article_to_api(doc: dict, people_by_url: dict[str, list[dict]] | None = None) -> dict:
    url = str(doc.get("url") or doc.get("articleUrl") or doc.get("article_url") or "").strip()
    authors = _coerce_list_field(doc.get("authors") or doc.get("author"))
    tags = _coerce_list_field(doc.get("tags") or doc.get("categories"))
    section = str(doc.get("section") or doc.get("category") or (tags[0] if tags else "") or "").strip()
    return {
        "id": str(doc.get("_id")),
        "title": doc.get("title") or doc.get("articleTitle") or "",
        "url": url,
        "outlet": doc.get("outlet", "") or "",
        "authors": authors,
        "author": ", ".join(authors),
        "section": section,
        "category": doc.get("category", "") or "",
        "tags": tags,
        "publishedAt": _published_date_for_api(doc.get("datePublished") or doc.get("publishedAt") or doc.get("date") or ""),
        "interviewees": (people_by_url or {}).get(_article_join_url_key(url), []),
    }


def _doc_public_id(doc: dict, fallback_field: str = "id") -> str:
    return str(doc.get(fallback_field) or doc.get("id") or doc.get("_id") or "")


def _date_for_api(value) -> str:
    return _month_day_year(value)


def _datetime_for_api(value) -> str:
    if isinstance(value, datetime):
        parsed = value
    else:
        raw = str(value or "").strip()
        if not raw:
            return ""
        try:
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        except ValueError:
            return ""
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    else:
        parsed = parsed.astimezone(timezone.utc)
    return parsed.isoformat(timespec="milliseconds")


def _int_field(doc: dict, field: str, default: int = 0) -> int:
    try:
        return int(doc.get(field, default) or default)
    except Exception:
        return default


def _bool_field(doc: dict, field: str, default: bool = False) -> bool:
    value = doc.get(field, default)
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "on"}


def _safe_http_url(value: str, max_length: int = 2048) -> str:
    raw = str(value or "")
    if not raw or raw != raw.strip() or len(raw) > max_length:
        return ""
    if "\\" in raw or any(char.isspace() or ord(char) < 32 or ord(char) == 127 for char in raw):
        return ""
    try:
        parsed = urlparse(raw)
        _ = parsed.port
    except Exception:
        return ""
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.hostname:
        return ""
    if parsed.username or parsed.password:
        return ""
    return raw


def _is_valid_http_url(value: str) -> bool:
    return bool(_safe_http_url(value))


def _normalized_publication_hostname(value: str) -> str:
    raw = str(value or "").strip().lower().rstrip(".")
    if not raw:
        return ""
    try:
        host = raw.encode("idna").decode("ascii")
    except (UnicodeError, ValueError):
        return ""
    if len(host) > 253:
        return ""
    labels = host.split(".")
    if len(labels) < 2:
        return ""
    if any(not re.fullmatch(r"[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?", label) for label in labels):
        return ""
    return host


def _validated_publication_url(value: str, workspace: dict | None) -> str:
    safe_url = _safe_http_url(value)
    if not safe_url or not workspace:
        return ""
    parsed_publication = urlparse(safe_url)
    if parsed_publication.port is not None and parsed_publication.port not in {80, 443}:
        return ""
    allowed_hosts = set()
    configured_domain = str(workspace.get("articleDomain") or "").strip().lower().rstrip(".")
    if configured_domain:
        if "://" in configured_domain:
            try:
                configured_domain = (urlparse(configured_domain).hostname or "").lower().rstrip(".")
            except Exception:
                configured_domain = ""
        else:
            configured_domain = configured_domain.split("/", 1)[0].split(":", 1)[0].rstrip(".")
        configured_domain = _normalized_publication_hostname(configured_domain)
        if configured_domain:
            allowed_hosts.add(configured_domain)
    configured_url = _safe_http_url(workspace.get("publicationUrl") or "")
    if configured_url:
        configured_url_host = _normalized_publication_hostname(urlparse(configured_url).hostname or "")
        if configured_url_host:
            allowed_hosts.add(configured_url_host)
    candidate_host = _normalized_publication_hostname(parsed_publication.hostname or "")
    if not candidate_host or not allowed_hosts:
        return ""
    if any(candidate_host == host or candidate_host.endswith(f".{host}") for host in allowed_hosts if host):
        return safe_url
    return ""


DRIVE_FILE_ID_RE = re.compile(r"^[A-Za-z0-9_-]{10,256}$")
DANGEROUS_UPLOAD_SUFFIXES = {
    ".app", ".bat", ".cmd", ".com", ".cpl", ".dll", ".exe", ".hta",
    ".html", ".htm", ".jar", ".js", ".jse", ".lnk", ".msi", ".msp",
    ".ps1", ".reg", ".scr", ".svg", ".vbs", ".vbe", ".wsf",
}


def _drive_type_label(mime_type: str = "", name: str = "") -> str:
    mime = str(mime_type or "").strip().lower()
    filename = str(name or "").strip().lower()
    if mime == "application/vnd.google-apps.document":
        return "Doc"
    if mime == "application/vnd.google-apps.presentation":
        return "Slides"
    if mime == "application/vnd.google-apps.spreadsheet":
        return "Sheet"
    if mime == "application/vnd.google-apps.folder":
        return "Folder"
    if mime == "application/pdf" or filename.endswith(".pdf"):
        return "PDF"
    if "wordprocessingml" in mime or filename.endswith((".doc", ".docx", ".odt")):
        return "Document"
    if "presentationml" in mime or filename.endswith((".ppt", ".pptx", ".odp")):
        return "Slides"
    if "spreadsheetml" in mime or filename.endswith((".xls", ".xlsx", ".ods", ".csv")):
        return "Spreadsheet"
    if mime.startswith("image/"):
        return "Image"
    if mime.startswith("video/"):
        return "Video"
    if mime.startswith("audio/"):
        return "Audio"
    if mime.startswith("text/") or filename.endswith((".txt", ".md", ".rtf")):
        return "Text"
    return "Drive file"


def _new_attachment_id(prefix: str = "att") -> str:
    return f"{prefix}_{secrets.token_urlsafe(8)}"


def _attachment_added_by(user_doc: dict) -> dict:
    return {
        "userId": str(user_doc.get("_id")),
        "email": normalize_email(user_doc.get("email") or ""),
        "name": _user_display_name(user_doc),
    }


def _story_attachment_item_to_api(story_id: str, item: dict) -> dict | None:
    if not isinstance(item, dict):
        return None
    attachment_type = str(item.get("type") or "").strip().lower()
    attachment_id = str(item.get("id") or item.get("attachmentId") or "").strip()

    if attachment_type == "drive":
        drive_file_id = str(item.get("fileId") or "").strip()
        if not DRIVE_FILE_ID_RE.fullmatch(drive_file_id):
            return None
        drive_url = _safe_http_url(item.get("webViewLink") or item.get("url") or "")
        return {
            "id": attachment_id or drive_file_id,
            "type": "drive",
            "provider": "google-drive",
            "fileId": drive_file_id,
            "name": item.get("name") or "Drive file",
            "mimeType": item.get("mimeType") or "",
            "typeLabel": item.get("typeLabel") or _drive_type_label(item.get("mimeType"), item.get("name")),
            "url": drive_url,
            "webViewLink": drive_url,
            "iconUrl": _safe_http_url(item.get("iconLink") or ""),
            "addedBy": item.get("addedBy") or {},
            "addedAt": _date_for_api(item.get("addedAt")),
            "permissionStatus": item.get("permissionStatus") or "not_shared",
            "shareResults": item.get("shareResults") if isinstance(item.get("shareResults"), list) else [],
            "lastShareAttemptAt": _date_for_api(item.get("lastShareAttemptAt")),
            "copyable": True,
        }

    if attachment_type == "file":
        file_id = str(item.get("fileId") or item.get("attachmentFileId") or "").strip()
        file_name = str(item.get("name") or item.get("attachmentName") or "").strip()
        if not file_id or not file_name:
            return None
        return {
            "id": attachment_id or file_id,
            "type": "file",
            "name": file_name,
            "contentType": item.get("contentType") or item.get("attachmentContentType") or "application/octet-stream",
            "size": int(item.get("size") or item.get("attachmentSize") or 0),
            "url": f"/api/stories/{story_id}/attachments/{attachment_id}" if attachment_id else f"/api/stories/{story_id}/attachment",
            "uploadedAt": _date_for_api(item.get("uploadedAt") or item.get("attachmentUploadedAt")),
            "addedBy": item.get("addedBy") or {},
            "permissionStatus": "app_access",
        }

    if attachment_type == "link":
        link_url = str(item.get("url") or item.get("webViewLink") or "").strip()
        if not _is_valid_http_url(link_url):
            return None
        return {
            "id": attachment_id or link_url,
            "type": "link",
            "provider": "manual-link",
            "name": item.get("name") or "Story link",
            "url": link_url,
            "uploadedAt": _date_for_api(item.get("uploadedAt") or item.get("addedAt")),
            "addedBy": item.get("addedBy") or {},
            "permissionStatus": "manual",
        }
    return None


def _legacy_story_attachment_items(doc: dict) -> list[dict]:
    items = []
    drive_attachment = doc.get("driveAttachment") if isinstance(doc.get("driveAttachment"), dict) else {}
    if drive_attachment.get("fileId"):
        items.append({
            **drive_attachment,
            "id": drive_attachment.get("id") or "legacy-drive",
            "type": "drive",
        })

    file_id = str(doc.get("attachmentFileId") or "").strip()
    file_name = str(doc.get("attachmentName") or "").strip()
    if file_id and file_name:
        items.append({
            "id": "legacy-file",
            "type": "file",
            "fileId": file_id,
            "name": file_name,
            "contentType": doc.get("attachmentContentType") or "application/octet-stream",
            "size": _int_field(doc, "attachmentSize"),
            "uploadedAt": doc.get("attachmentUploadedAt") or doc.get("updatedAt"),
        })

    link_url = str(doc.get("googleDocUrl") or doc.get("docUrl") or "").strip()
    if _is_valid_http_url(link_url):
        items.append({
            "id": "legacy-link",
            "type": "link",
            "provider": "manual-link",
            "name": doc.get("attachmentName") or "Story link",
            "url": link_url,
            "uploadedAt": _date_for_api(doc.get("updatedAt")),
            "permissionStatus": "manual",
        })
    return items


def _story_attachments_to_api(doc: dict) -> list[dict]:
    story_id = _doc_public_id(doc, "storyId")
    raw_items = []
    if isinstance(doc.get("attachments"), list):
        raw_items.extend(item for item in doc.get("attachments") if isinstance(item, dict))
    raw_items.extend(_legacy_story_attachment_items(doc))
    attachments = []
    seen = set()
    for item in raw_items:
        api_item = _story_attachment_item_to_api(story_id, item)
        if not api_item:
            continue
        key = _attachment_item_key(api_item)
        if key in seen:
            continue
        seen.add(key)
        attachments.append(api_item)
    return attachments


def _story_attachment_usage(doc: dict) -> tuple[int, int]:
    attachments = _story_attachments_to_api(doc)
    stored_bytes = sum(
        max(0, int(item.get("size") or 0))
        for item in attachments
        if item.get("type") == "file"
    )
    return len(attachments), stored_bytes


def _story_attachment_to_api(doc: dict) -> dict | None:
    attachments = _story_attachments_to_api(doc)
    return attachments[0] if attachments else None


def _attachment_item_key(item: dict) -> str:
    item_type = str(item.get("type") or "").strip().lower()
    url = str(item.get("webViewLink") or item.get("url") or "").strip()
    if item_type in {"drive", "link"} and url:
        return f"url:{url}"
    if item_type == "file":
        file_id = str(item.get("fileId") or item.get("attachmentFileId") or "").strip()
        if file_id:
            return f"{item_type}:{file_id}"
    return str(item.get("id") or item.get("attachmentId") or item.get("url") or "").strip()


def _attachment_identifiers(item: dict) -> set[str]:
    identifiers = {
        str(item.get("id") or "").strip(),
        str(item.get("attachmentId") or "").strip(),
        str(item.get("fileId") or "").strip(),
        str(item.get("attachmentFileId") or "").strip(),
        str(item.get("webViewLink") or "").strip(),
        str(item.get("url") or "").strip(),
    }
    key = _attachment_item_key(item)
    if key:
        identifiers.add(key)
    return {identifier for identifier in identifiers if identifier}


def _attachment_pull_condition(attachment_id: str) -> dict:
    target = str(attachment_id or "").strip()
    return {
        "$or": [
            {"id": target},
            {"attachmentId": target},
            {"fileId": target},
            {"attachmentFileId": target},
            {"webViewLink": target},
            {"url": target},
        ]
    }


def _merged_attachment_items(story: dict, *new_items: dict) -> list[dict]:
    additions = []
    if isinstance(story.get("attachments"), list):
        additions.extend(item for item in story.get("attachments") if isinstance(item, dict))
    additions.extend(_legacy_story_attachment_items(story))
    additions.extend(item for item in new_items if isinstance(item, dict))

    deduped = []
    seen = set()
    for item in additions:
        key = _attachment_item_key(item)
        if key and key in seen:
            continue
        if key:
            seen.add(key)
        deduped.append(item)
    return deduped


def _attachment_push_value(story: dict, *new_items: dict) -> dict:
    return {"$each": _merged_attachment_items(story, *new_items)}


DUE_DATE_FIELDS = ("deadline", "dueDate", "approvalDueDate", "due", "dateDue", "deadlineDate", "due_date", "deadline_date", "approval_due_date")


def _first_text_value(*values) -> str:
    for value in values:
        text = str(value or "").strip()
        if text:
            return text
    return ""


def _due_date_from_doc(doc: dict) -> str:
    return _first_text_value(*(doc.get(field) for field in DUE_DATE_FIELDS))


def _story_deadline(doc: dict) -> str:
    raw_deadline = _due_date_from_doc(doc)
    return _date_for_api(raw_deadline) or raw_deadline


def _story_to_api(doc: dict) -> dict:
    writer = str(doc.get("writer") or doc.get("owner") or doc.get("author") or "").strip()
    authors = _coerce_list_field(doc.get("authors"))
    if not writer and authors:
        writer = authors[0]
    attachments = _story_attachments_to_api(doc)
    deadline = _story_deadline(doc)
    accepted_collaborators = [
        item
        for item in _story_collaborators(doc)
        if item.get("status") == "accepted" and item.get("role") == "edit"
    ]
    author_names = []
    for name in [writer, *authors, *[item.get("name") for item in accepted_collaborators]]:
        clean_name = str(name or "").strip()
        if clean_name and clean_name not in author_names:
            author_names.append(clean_name)
    return {
        "id": _doc_public_id(doc, "storyId"),
        "title": doc.get("title") or doc.get("storyTitle") or "Untitled story",
        "section": doc.get("section", "") or "",
        "writer": writer or doc.get("writerEmail", "") or "Unassigned",
        "writerEmail": doc.get("writerEmail") or doc.get("ownerEmail") or "",
        "writerUserId": str(doc.get("writerUserId") or doc.get("ownerUserId") or ""),
        "authors": author_names,
        "authorProfiles": ([{
            "name": writer or doc.get("writerEmail", "") or "Unassigned",
            "email": doc.get("writerEmail") or doc.get("ownerEmail") or "",
            "primary": True,
        }] if writer or doc.get("writerEmail") or doc.get("ownerEmail") else []) + [{
            "name": item.get("name") or item.get("email"),
            "email": item.get("email"),
            "primary": False,
        } for item in accepted_collaborators],
        "editor": doc.get("editor", "") or "",
        "status": doc.get("status") or "Assigned",
        "priority": doc.get("priority") or "Normal",
        "deadline": deadline,
        "dueDate": deadline,
        "dueSoon": _bool_field(doc, "dueSoon"),
        "submittedAt": _date_for_api(doc.get("submittedAt") or doc.get("createdAt")),
        "lastEdited": _date_for_api(doc.get("lastEdited") or doc.get("updatedAt")),
        "googleDocUrl": doc.get("googleDocUrl") or doc.get("docUrl") or "",
        "revisionCount": _int_field(doc, "revisionCount"),
        "wordCount": _int_field(doc, "wordCount"),
        "sourceCount": _int_field(doc, "sourceCount"),
        "unread": _bool_field(doc, "unread"),
        "summary": doc.get("summary", "") or "",
        "nextStep": doc.get("nextStep", "") or "",
        "editorNote": doc.get("editorNote", "") or "",
        "publicationUrl": doc.get("publicationUrl", "") or "",
        "publishedAt": _date_for_api(doc.get("publishedAt")),
        "attachments": attachments,
        "attachment": attachments[0] if attachments else None,
        "feedback": doc.get("feedback") if isinstance(doc.get("feedback"), list) else [],
        "comments": doc.get("comments") if isinstance(doc.get("comments"), list) else [],
        "sources": doc.get("sources") if isinstance(doc.get("sources"), list) else [],
        "collaborators": _story_collaborators(doc),
    }


def _normalize_pitch_status(value) -> str:
    status = str(value or "").strip()
    return {
        "New": "In Progress",
        "Submitted": "Ready for Review",
        "Needs Review": "Ready for Review",
    }.get(status, status or "In Progress")


def _pitch_detail_updates(payload: dict, pitch: dict) -> tuple[dict, str]:
    fields = {"title", "angle", "section", "notes"}
    if not any(field in payload for field in fields):
        return {}, "No editable pitch details provided."
    merged = {
        field: str(payload.get(field) if field in payload else pitch.get(field) or "").strip()
        for field in fields
    }
    if not merged["title"]:
        return {}, "Pitch title is required."
    if not merged["angle"]:
        return {}, "Pitch angle is required."
    existing_section = str(pitch.get("section") or "").strip()
    if not merged["section"]:
        return {}, "Pitch section is required."
    if merged["section"] not in PITCH_SECTIONS and merged["section"] != existing_section:
        return {}, "Choose a valid pitch section."
    limits = {
        "title": (MAX_PITCH_TITLE_LENGTH, "Pitch title"),
        "angle": (MAX_PITCH_ANGLE_LENGTH, "Pitch angle"),
        "section": (MAX_PITCH_SECTION_LENGTH, "Pitch section"),
        "notes": (MAX_PITCH_NOTES_LENGTH, "Pitch notes"),
    }
    for field, (limit, label) in limits.items():
        if len(merged[field]) > limit:
            return {}, f"{label} must be {limit} characters or fewer."
    return {field: merged[field] for field in fields if field in payload}, ""

def _pitch_to_api(doc: dict) -> dict:
    deadline = _story_deadline(doc)
    return {
        "id": _doc_public_id(doc, "pitchId"),
        "title": doc.get("title") or "Untitled pitch",
        "angle": doc.get("angle") or doc.get("summary") or "",
        "status": _normalize_pitch_status(doc.get("status")),
        "section": doc.get("section", "") or "",
        "owner": doc.get("owner") or doc.get("writer") or doc.get("ownerEmail") or "Unassigned",
        "ownerEmail": doc.get("ownerEmail") or doc.get("writerEmail") or "",
        "ownerUserId": str(doc.get("ownerUserId") or doc.get("writerUserId") or ""),
        "submittedAt": _date_for_api(doc.get("submittedAt") or doc.get("createdAt")),
        "notes": doc.get("notes", "") or "",
        "deadline": deadline,
        "dueDate": deadline,
        "editorFeedback": doc.get("editorFeedback", "") or "",
        "feedback": doc.get("feedback") if isinstance(doc.get("feedback"), list) else [],
        "comments": doc.get("comments") if isinstance(doc.get("comments"), list) else [],
        "updatedAt": _date_for_api(doc.get("updatedAt") or doc.get("submittedAt") or doc.get("createdAt")),
    }


def _activity_to_api(doc: dict) -> dict:
    text = doc.get("text") or ""
    from_status = doc.get("fromStatus") or ""
    to_status = doc.get("toStatus") or ""
    if not text and to_status:
        text = f"Status changed from {from_status or 'none'} to {to_status}."
    return {
        "id": str(doc.get("_id")),
        "entityType": doc.get("entityType") or "",
        "entityId": doc.get("entityId") or "",
        "eventType": doc.get("eventType") or "",
        "text": text,
        "time": _date_for_api(doc.get("createdAt")),
        "occurredAt": _datetime_for_api(doc.get("createdAt")),
        "actorName": doc.get("actorName") or "",
        "fromStatus": from_status,
        "toStatus": to_status,
    }


def _feedback_to_api(doc: dict) -> dict:
    return {
        "id": str(doc.get("_id")),
        "entityType": doc.get("entityType") or "",
        "entityId": doc.get("entityId") or "",
        "author": doc.get("authorName") or doc.get("actorName") or "Editor",
        "authorEmail": doc.get("authorEmail") or "",
        "text": doc.get("text") or "",
        "time": _date_for_api(doc.get("createdAt")),
        "updatedAt": _date_for_api(doc.get("updatedAt") or doc.get("createdAt")),
    }


def _can_access_entity(entity_type: str, entity_id: str, user_doc: dict):
    if entity_type == "story":
        return _find_owned_story_or_404(entity_id, user_doc)
    if entity_type == "pitch":
        return _find_owned_pitch_or_404(entity_id, user_doc)
    return None


def _require_feedback_entity(entity_type: str, entity_id: str, user_doc: dict):
    entity_type = str(entity_type or "").strip().lower()
    entity_id = str(entity_id or "").strip()
    if entity_type not in {"story", "pitch"} or not entity_id:
        return None, "Feedback requires story or pitch entity information.", 400
    entity_doc = _can_access_entity(entity_type, entity_id, user_doc)
    if not entity_doc:
        return None, "Feedback is not available for this item.", 403
    return entity_doc, "", 200


def _record_status_activity(entity_type: str, entity_id: str, from_status: str, to_status: str, actor_doc: dict):
    if from_status == to_status:
        return
    activity_col.insert_one({
        "workspaceId": _workspace_id_for_user(actor_doc),
        "entityType": entity_type,
        "entityId": str(entity_id),
        "eventType": "status_change",
        "fromStatus": from_status or "",
        "toStatus": to_status or "",
        "text": f"Status changed from {from_status or 'none'} to {to_status}.",
        "actorId": str(actor_doc.get("_id")),
        "actorEmail": actor_doc.get("email", ""),
        "actorName": _user_display_name(actor_doc),
        "createdAt": datetime.now(timezone.utc),
    })


def _delete_story_file(file_id_value):
    file_id = _object_id_or_none(str(file_id_value or ""))
    if not file_id:
        return
    try:
        story_files.delete(file_id)
    except Exception:
        pass


def _create_story_from_pitch(pitch_doc: dict, actor_doc: dict, deadline: str = "", editor_note: str = ""):
    pitch_id = _doc_public_id(pitch_doc, "pitchId")
    now_iso = _now_iso()
    story_deadline = _first_text_value(deadline, _due_date_from_doc(pitch_doc))
    workspace_id = _workspace_id_for_user(actor_doc)
    key = {"sourcePitchId": pitch_id, "workspaceId": workspace_id}
    insert_doc = {
        "workspaceId": workspace_id,
        "title": pitch_doc.get("title") or "Untitled story",
        "section": pitch_doc.get("section", "") or "",
        "writer": pitch_doc.get("owner") or pitch_doc.get("writer") or "Unassigned",
        "writerEmail": pitch_doc.get("ownerEmail") or pitch_doc.get("writerEmail") or "",
        "writerUserId": pitch_doc.get("ownerUserId") or pitch_doc.get("writerUserId") or "",
        "editor": _user_display_name(actor_doc),
        "status": "Assigned",
        "priority": "Normal",
        "deadline": story_deadline,
        "dueDate": story_deadline,
        "submittedAt": pitch_doc.get("submittedAt") or pitch_doc.get("createdAt") or now_iso,
        "createdAt": now_iso,
        "summary": pitch_doc.get("angle") or pitch_doc.get("summary") or "",
        "nextStep": editor_note or "Begin reporting from the approved pitch.",
        "editorNote": editor_note or pitch_doc.get("editorFeedback", "") or "",
        "googleDocUrl": "",
        "revisionCount": 0,
        "wordCount": 0,
        "sourceCount": 0,
        "sourcePitchId": pitch_id,
    }
    existing_update = {"updatedAt": now_iso}
    try:
        before = stories_col.find_one_and_update(
            key,
            {"$setOnInsert": insert_doc, "$set": existing_update},
            upsert=True,
            return_document=ReturnDocument.BEFORE,
        )
    except DuplicateKeyError:
        before = stories_col.find_one(key)
        if before:
            stories_col.update_one({"_id": before["_id"], "workspaceId": workspace_id}, {"$set": existing_update})
    story = stories_col.find_one(key)
    if not story:
        raise RuntimeError("Story creation did not complete.")
    followup_update = {}
    if story_deadline and _due_date_from_doc(story) != story_deadline:
        followup_update["deadline"] = story_deadline
        followup_update["dueDate"] = story_deadline
    if editor_note and str(story.get("editorNote") or "") != editor_note:
        followup_update["editorNote"] = editor_note
        followup_update["nextStep"] = editor_note
    if followup_update:
        stories_col.update_one({"_id": story["_id"], "workspaceId": workspace_id}, {"$set": followup_update})
        story = stories_col.find_one({"_id": story["_id"], "workspaceId": workspace_id}) or {**story, **followup_update}
    if before is None:
        try:
            _record_status_activity("story", _doc_public_id(story, "storyId"), "", "Assigned", actor_doc)
        except Exception:
            pass
    return story


def create_user(email: str, password: str, first_name: str, last_name: str):
    now_iso = _now_iso()
    doc = {
        "email": normalize_email(email),
        "passwordHash": generate_password_hash(password),
        "authVersion": 1,
        "firstName": (first_name or "").strip(),
        "lastName": (last_name or "").strip(),
        "createdAt": now_iso,
        "lastLoginAt": now_iso,
        "authProviders": {"password": True},
    }
    res = users_col.insert_one(doc)
    doc["_id"] = res.inserted_id
    return doc


def _split_google_name(profile: dict) -> tuple[str, str]:
    first_name = str(profile.get("given_name") or "").strip()
    last_name = str(profile.get("family_name") or "").strip()
    if first_name or last_name:
        return first_name, last_name
    name = str(profile.get("name") or "").strip()
    if not name:
        return "", ""
    parts = name.split()
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def _google_redirect_uri() -> str:
    configured_redirect = str(GOOGLE_AUTH_REDIRECT_URI or "").strip()
    if configured_redirect:
        return configured_redirect
    if FRONTEND_ORIGIN:
        return f"{FRONTEND_ORIGIN}/api/auth/google/callback"
    if FLASK_ENV != "production":
        return f"{DEFAULT_DEV_FRONTEND_ORIGIN}/api/auth/google/callback"
    return url_for("api_google_callback", _external=True)


def _google_scope_string() -> str:
    return " ".join(GOOGLE_OAUTH_SCOPES)


def _google_state_serializer():
    return URLSafeTimedSerializer(app.secret_key, salt="falcon-google-oauth-state-v1")


def _google_missing_config() -> list[str]:
    missing = []
    if not GOOGLE_CLIENT_ID:
        missing.append("GOOGLE_CLIENT_ID")
    if not GOOGLE_CLIENT_SECRET:
        missing.append("GOOGLE_CLIENT_SECRET")
    if not _has_persistent_oauth_encryption_key:
        missing.append("OAUTH_TOKEN_ENCRYPTION_KEY (or a strong FLASK_SECRET_KEY)")
    return missing


def _normalize_google_intent(value: str) -> str:
    candidate = str(value or "").strip().lower()
    return candidate if candidate in GOOGLE_INTENTS else GOOGLE_INTENT_LOGIN


def _is_google_configured() -> bool:
    return not _google_missing_config()


def _oauth_expires_at(expires_in) -> datetime:
    seconds = max(60, int(expires_in or 3600) - 60)
    return datetime.now(timezone.utc) + timedelta(seconds=seconds)


def _parse_oauth_expiry(value) -> datetime | None:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except Exception:
        return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


OAUTH_TOKEN_PREFIX = "fernet:v1:"


def _oauth_token_cipher() -> Fernet:
    material = OAUTH_TOKEN_ENCRYPTION_KEY if _has_strong_oauth_encryption_key else str(app.secret_key)
    key = base64.urlsafe_b64encode(hashlib.sha256(material.encode("utf-8")).digest())
    return Fernet(key)


def _encrypt_oauth_token(value: str) -> str:
    token = str(value or "").strip()
    if not token:
        return ""
    encrypted = _oauth_token_cipher().encrypt(token.encode("utf-8")).decode("ascii")
    return f"{OAUTH_TOKEN_PREFIX}{encrypted}"


def _decrypt_oauth_token(value: str) -> str:
    stored = str(value or "").strip()
    if not stored:
        return ""
    if not stored.startswith(OAUTH_TOKEN_PREFIX):
        return stored
    try:
        encrypted = stored[len(OAUTH_TOKEN_PREFIX):].encode("ascii")
        return _oauth_token_cipher().decrypt(encrypted).decode("utf-8")
    except (InvalidToken, UnicodeError, ValueError):
        return ""


def _migrate_legacy_oauth_tokens():
    if not _has_persistent_oauth_encryption_key:
        return
    projection = {"googleOAuth.accessToken": 1, "googleOAuth.refreshToken": 1}
    for user_doc in users_col.find({"googleOAuth": {"$exists": True}}, projection):
        oauth = user_doc.get("googleOAuth") if isinstance(user_doc.get("googleOAuth"), dict) else {}
        updates = {}
        guard = {"_id": user_doc["_id"]}
        for field in ("accessToken", "refreshToken"):
            value = str(oauth.get(field) or "").strip()
            if value and not value.startswith(OAUTH_TOKEN_PREFIX):
                updates[f"googleOAuth.{field}"] = _encrypt_oauth_token(value)
                guard[f"googleOAuth.{field}"] = value
        if updates:
            users_col.update_one(guard, {"$set": updates})


try:
    _migrate_legacy_oauth_tokens()
except Exception:
    if FLASK_ENV == "production":
        raise


def _store_google_oauth_tokens(user_doc: dict, token_payload: dict):
    access_token = str(token_payload.get("access_token") or "").strip()
    if not access_token:
        return
    scopes = str(token_payload.get("scope") or _google_scope_string()).split()
    update = {
        "googleOAuth.accessToken": _encrypt_oauth_token(access_token),
        "googleOAuth.tokenType": token_payload.get("token_type") or "Bearer",
        "googleOAuth.expiresAt": _oauth_expires_at(token_payload.get("expires_in")),
        "googleOAuth.scopes": scopes,
        "googleOAuth.updatedAt": datetime.now(timezone.utc),
    }
    refresh_token = str(token_payload.get("refresh_token") or "").strip()
    if refresh_token:
        update["googleOAuth.refreshToken"] = _encrypt_oauth_token(refresh_token)
    users_col.update_one({"_id": user_doc["_id"]}, {"$set": update})
    user_doc.setdefault("googleOAuth", {}).update({
        "accessToken": access_token,
        "tokenType": update["googleOAuth.tokenType"],
        "expiresAt": update["googleOAuth.expiresAt"],
        "scopes": scopes,
        "updatedAt": update["googleOAuth.updatedAt"],
    })
    if refresh_token:
        user_doc["googleOAuth"]["refreshToken"] = refresh_token


def _google_drive_config_missing() -> list[str]:
    missing = []
    if not GOOGLE_PICKER_API_KEY:
        missing.append("GOOGLE_PICKER_API_KEY")
    if not GOOGLE_PICKER_APP_ID:
        missing.append("GOOGLE_PICKER_APP_ID")
    if not GOOGLE_PICKER_CLIENT_ID:
        missing.append("GOOGLE_PICKER_CLIENT_ID")
    return missing


def _google_reauth_url(next_url: str = "") -> str:
    params = {
        "origin": _sanitize_frontend_origin(FRONTEND_ORIGIN) or DEFAULT_DEV_FRONTEND_ORIGIN,
        "intent": GOOGLE_INTENT_LOGIN,
    }
    if next_url:
        params["next"] = _sanitize_next_url(next_url)
    return f"/api/auth/google/start?{urlencode(params)}"


def _has_drive_scope(user_doc: dict) -> bool:
    oauth = user_doc.get("googleOAuth") if isinstance(user_doc.get("googleOAuth"), dict) else {}
    scopes = oauth.get("scopes") if isinstance(oauth.get("scopes"), list) else str(oauth.get("scope") or "").split()
    return GOOGLE_DRIVE_SCOPE in set(str(scope) for scope in scopes)


def _google_oauth_error(message: str, status: int = 409):
    return jsonify({
        "ok": False,
        "error": message,
        "reauthUrl": _google_reauth_url(request.args.get("next") or request.path),
    }), status


def _google_access_token_for_user(user_doc: dict) -> tuple[str, str]:
    oauth = user_doc.get("googleOAuth") if isinstance(user_doc.get("googleOAuth"), dict) else {}
    if not oauth:
        return "", "Sign in with Google again to connect Drive."
    if not _has_drive_scope(user_doc):
        return "", "Sign in with Google again to grant Drive attachment access."

    access_token = _decrypt_oauth_token(oauth.get("accessToken") or "")
    expires_at = _parse_oauth_expiry(oauth.get("expiresAt"))
    if access_token and expires_at and expires_at > datetime.now(timezone.utc):
        return access_token, ""

    refresh_token = _decrypt_oauth_token(oauth.get("refreshToken") or "")
    if not refresh_token:
        return "", "Sign in with Google again so Falcon can refresh Drive access."

    try:
        response = requests.post(
            GOOGLE_TOKEN_ENDPOINT,
            data={
                "client_id": GOOGLE_CLIENT_ID,
                "client_secret": GOOGLE_CLIENT_SECRET,
                "refresh_token": refresh_token,
                "grant_type": "refresh_token",
            },
            timeout=10,
        )
        response.raise_for_status()
        payload = response.json()
        payload.setdefault("refresh_token", refresh_token)
        if "scope" not in payload:
            payload["scope"] = " ".join(oauth.get("scopes") or GOOGLE_OAUTH_SCOPES)
        _store_google_oauth_tokens(user_doc, payload)
        return str(payload.get("access_token") or ""), ""
    except Exception:
        return "", "Google Drive access expired. Sign in with Google again."


def _drive_headers(access_token: str) -> dict:
    return {
        "Authorization": f"Bearer {access_token}",
        "Accept": "application/json",
    }


def _drive_file_metadata(file_id: str, access_token: str) -> tuple[dict, str]:
    clean_file_id = str(file_id or "").strip()
    if not DRIVE_FILE_ID_RE.fullmatch(clean_file_id):
        return {}, "Drive file id is invalid."
    try:
        response = requests.get(
            f"{GOOGLE_DRIVE_FILES_ENDPOINT}/{clean_file_id}",
            params={
                "fields": "id,name,mimeType,webViewLink,iconLink",
                "supportsAllDrives": "true",
            },
            headers=_drive_headers(access_token),
            timeout=10,
        )
        response.raise_for_status()
        return response.json(), ""
    except Exception:
        return {}, "Could not read the selected Google Drive file. Reconnect Google Drive and try again."


def _editor_permission_emails(actor_doc: dict) -> list[str]:
    emails = set()
    cursor = users_col.find(
        {"role": {"$in": [ROLE_EDITOR, ROLE_ADMIN]}, "workspaceId": _workspace_id_for_user(actor_doc)},
        {"email": 1},
    ).sort([("email", ASCENDING)]).limit(max(1, MAX_DRIVE_SHARE_RECIPIENTS) + 1)
    for doc in cursor:
        email = normalize_email(doc.get("email") or "")
        if email:
            emails.add(email)
    return sorted(emails)


def _share_drive_attachment_with_editors(drive_attachment: dict, actor_doc: dict) -> tuple[bool, list[dict], str]:
    drive_attachment = drive_attachment if isinstance(drive_attachment, dict) else {}
    file_id = str(drive_attachment.get("fileId") or "").strip()
    if not DRIVE_FILE_ID_RE.fullmatch(file_id):
        return False, [], "Drive file id is invalid."
    access_token, token_error = _google_access_token_for_user(actor_doc)
    if token_error:
        return False, [], token_error

    actor_email = normalize_email(actor_doc.get("email") or "")
    all_editor_emails = [email for email in _editor_permission_emails(actor_doc) if email and email != actor_email]
    recipient_cap = max(1, min(MAX_DRIVE_SHARE_RECIPIENTS, 200))
    recipients_truncated = len(all_editor_emails) > recipient_cap
    editor_emails = all_editor_emails[:recipient_cap]
    if not editor_emails:
        return False, [], "No editor or admin Google accounts are available to share this file with."

    results = []
    deadline = monotonic() + max(1.0, min(DRIVE_SHARE_TOTAL_TIMEOUT_SECONDS, 30.0))
    for email in editor_emails:
        result = {
            "email": email,
            "role": "writer",
            "ok": False,
            "permissionId": "",
            "error": "",
            "attemptedAt": datetime.now(timezone.utc),
        }
        remaining = deadline - monotonic()
        if remaining <= 0:
            result["error"] = "Drive sharing time budget was exhausted."
            results.append(result)
            continue
        try:
            response = requests.post(
                f"{GOOGLE_DRIVE_FILES_ENDPOINT}/{file_id}/permissions",
                params={
                    "sendNotificationEmail": "false",
                    "supportsAllDrives": "true",
                },
                json={
                    "type": "user",
                    "role": "writer",
                    "emailAddress": email,
                },
                headers={
                    **_drive_headers(access_token),
                    "Content-Type": "application/json",
                },
                timeout=max(0.5, min(5.0, remaining)),
            )
            if response.status_code in {200, 201}:
                payload = response.json()
                result["ok"] = True
                result["permissionId"] = str(payload.get("id") or "")
            elif response.status_code == 409:
                result["ok"] = True
                result["error"] = "Permission already exists."
            else:
                result["error"] = "Google Drive permission update failed."
        except Exception:
            result["error"] = "Google Drive permission update failed."
        results.append(result)

    ok = not recipients_truncated and all(item.get("ok") for item in results)
    message = "" if ok else "Could not share the attached Drive file with every editor/admin within the safe request budget."
    return ok, results, message


def _share_drive_file_with_editors(story: dict, actor_doc: dict) -> tuple[bool, list[dict], str]:
    drive_attachment = story.get("driveAttachment") if isinstance(story.get("driveAttachment"), dict) else {}
    return _share_drive_attachment_with_editors(drive_attachment, actor_doc)


def _sanitize_frontend_origin(origin: str) -> str:
    candidate = _origin_from_url(origin)
    if not candidate:
        return ""
    trusted = _trusted_request_origins()
    if candidate.lower().rstrip("/") in trusted:
        return candidate.rstrip("/")
    return ""


def _origin_from_redirect_uri() -> str:
    return _sanitize_frontend_origin(_google_redirect_uri())


def _frontend_redirect_url(path: str, origin: str = "") -> str:
    target = _sanitize_next_url(path) or "/dashboard"
    frontend_origin = _sanitize_frontend_origin(origin) or _sanitize_frontend_origin(FRONTEND_ORIGIN)
    if frontend_origin:
        return f"{frontend_origin}{target}"
    return target


def _make_google_state(next_url: str, frontend_origin: str, intent: str = GOOGLE_INTENT_LOGIN) -> str:
    nonce = ""
    for _attempt in range(5):
        candidate = secrets.token_urlsafe(24)
        try:
            oauth_state_col.insert_one({"nonce": candidate, "createdAt": datetime.now(timezone.utc)})
            nonce = candidate
            break
        except DuplicateKeyError:
            continue
    if not nonce:
        raise RuntimeError("Could not create a unique Google OAuth nonce.")
    pending = [
        value for value in session.get("google_oauth_nonces", [])
        if isinstance(value, str) and 16 <= len(value) <= 128 and value != nonce
    ]
    session["google_oauth_nonces"] = [*pending[-4:], nonce]
    return _google_state_serializer().dumps({
        "nonce": nonce,
        "next": _sanitize_next_url(next_url),
        "frontendOrigin": _sanitize_frontend_origin(frontend_origin),
        "intent": _normalize_google_intent(intent),
    })


def _load_google_state(state: str) -> tuple[dict, str]:
    try:
        data = _google_state_serializer().loads(state, max_age=GOOGLE_STATE_MAX_AGE_SECONDS)
    except SignatureExpired:
        return {}, "Google sign-in expired. Please try again."
    except BadSignature:
        return {}, "Google sign-in could not be verified."
    if not isinstance(data, dict):
        return {}, "Google sign-in could not be verified."
    nonce = data.get("nonce")
    if not isinstance(nonce, str) or not 16 <= len(nonce) <= 128:
        return {}, "Google sign-in could not be verified."

    pending = [value for value in session.get("google_oauth_nonces", []) if isinstance(value, str)]
    if not any(secrets.compare_digest(nonce, value) for value in pending):
        return {}, "Google sign-in was not started in this browser."

    try:
        consumed = oauth_state_col.find_one_and_delete({"nonce": nonce})
    except Exception:
        return {}, "Google sign-in could not be verified."
    if not consumed:
        return {}, "Google sign-in has already been used. Please try again."

    session["google_oauth_nonces"] = [
        value for value in pending if not secrets.compare_digest(nonce, value)
    ][-5:]
    return data, ""


def _google_email_domain_allowed(email: str) -> bool:
    if not GOOGLE_ALLOWED_DOMAINS:
        return True
    domain = normalize_email(email).rsplit("@", 1)[-1]
    return domain in GOOGLE_ALLOWED_DOMAINS


def _google_error_redirect(message: str, frontend_origin: str = ""):
    login_path = f"/login?{urlencode({'auth': 'google', 'error': message})}"
    frontend_origin = _sanitize_frontend_origin(frontend_origin) or _sanitize_frontend_origin(FRONTEND_ORIGIN)
    if frontend_origin:
        return redirect(f"{frontend_origin}{login_path}")
    return redirect(login_path)


def upsert_google_user(profile: dict, allow_create: bool = True):
    email = normalize_email(profile.get("email") or "")
    google_id = str(profile.get("sub") or profile.get("googleId") or profile.get("googleSub") or "").strip()
    if not email or not google_id or profile.get("email_verified") is not True:
        raise ValueError("Google profile did not include a verified identity.")
    if len(google_id) > 256:
        raise ValueError("Google profile identity is invalid.")

    now_iso = _now_iso()
    first_name, last_name = _split_google_name(profile)
    first_name = first_name[:80]
    last_name = last_name[:80]
    picture = str(profile.get("picture") or "").strip()
    if picture and not _is_valid_http_url(picture):
        picture = ""
    update = {
        "email": email,
        "googleId": google_id,
        "googleSub": google_id,
        "googleEmailVerified": True,
        "googlePicture": picture,
        "lastLoginAt": now_iso,
    }
    if first_name:
        update["firstName"] = first_name
    if last_name:
        update["lastName"] = last_name

    user_by_google = find_user_by_google_id(google_id)
    if user_by_google:
        existing_email = normalize_email(user_by_google.get("email") or "")
        if existing_email and existing_email != email:
            email_owner = find_user_by_email(email)
            if email_owner and str(email_owner.get("_id")) != str(user_by_google.get("_id")):
                raise ValueError("This Google account is already linked to another user.")
        return users_col.find_one_and_update(
            {"_id": user_by_google["_id"]},
            {"$set": {**update, "authProviders.google": True}},
            return_document=ReturnDocument.AFTER,
        ) or user_by_google

    user_by_email = find_user_by_email(email)
    if user_by_email:
        existing_google_id = _google_id_from_doc(user_by_email)
        if existing_google_id and existing_google_id != google_id:
            raise ValueError("This email is already linked to a different Google account.")
        if existing_google_id == google_id:
            return users_col.find_one_and_update(
                {"_id": user_by_email["_id"]},
                {"$set": {**update, "authProviders.google": True}},
                return_document=ReturnDocument.AFTER,
            ) or user_by_email

        claimed = users_col.find_one_and_update(
            {
                "_id": user_by_email["_id"],
                "$and": [
                    {"$or": [{"googleId": {"$exists": False}}, {"googleId": None}, {"googleId": ""}]},
                    {"$or": [{"googleSub": {"$exists": False}}, {"googleSub": None}, {"googleSub": ""}]},
                ],
            },
            [
                {"$set": {
                    **update,
                    "authProviders.google": True,
                    "authProviders.password": False,
                    "passwordCredentialRevokedAt": now_iso,
                    "passwordCredentialRevokedReason": "verified_google_reclaim",
                    "authVersion": {
                        "$add": [
                            {
                                "$cond": [
                                    {
                                        "$and": [
                                            {"$isNumber": "$authVersion"},
                                            {"$gte": ["$authVersion", 1]},
                                        ],
                                    },
                                    "$authVersion",
                                    1,
                                ],
                            },
                            1,
                        ],
                    },
                }},
                {"$unset": "passwordHash"},
            ],
            return_document=ReturnDocument.AFTER,
        )
        if claimed:
            return claimed

        concurrent_owner = find_user_by_email(email)
        if not concurrent_owner or _google_id_from_doc(concurrent_owner) != google_id:
            raise ValueError("This email is already linked to a different Google account.")
        return concurrent_owner

    if not allow_create:
        raise ValueError("No Falcon account is registered for this Google email. Create an account first.")

    doc = {
        **update,
        "firstName": first_name,
        "lastName": last_name,
        "createdAt": now_iso,
        "authVersion": 1,
        "authProviders": {"google": True},
    }
    res = users_col.insert_one(doc)
    doc["_id"] = res.inserted_id
    return doc


@app.get("/api/health")
def api_health():
    return jsonify({
        "ok": True,
        "status": "healthy",
        "service": "falcon-newsroom-v3-auth",
        "version": "v3-rbac-guest-2026-07-20",
        "capabilities": sorted(BACKEND_CAPABILITIES),
    })


def _story_status_transition_decision(
    role: str,
    current_status: str,
    next_status: str,
    is_author: bool,
    publication_url: str = "",
    current_publication_url: str = "",
) -> dict:
    role = str(role or "").strip().lower()
    current_status = str(current_status or "Assigned").strip()
    next_status = str(next_status or "").strip()
    if next_status not in VALID_STORY_STATUSES:
        return {"ok": False, "status": 400, "error": "Choose a valid story status."}

    safe_publication_url = _safe_http_url(publication_url) if publication_url else ""
    if next_status == "Published" and not safe_publication_url:
        return {"ok": False, "status": 400, "error": "A valid http or https publication URL is required."}

    if current_status == "Published" and next_status == "Published":
        if role != ROLE_ADMIN:
            return {"ok": False, "status": 403, "error": "Only admins can publish stories."}
        if safe_publication_url != str(current_publication_url or "").strip():
            return {"ok": False, "status": 409, "error": "Published stories cannot be republished at a different URL."}
        return {"ok": True, "idempotent": True, "publicationUrl": safe_publication_url}

    author_transitions = {
        ("Assigned", "Submitted"),
        ("Reporting", "Submitted"),
        ("Drafting", "Submitted"),
        ("Returned", "Submitted"),
        ("Needs Revision", "Submitted"),
        ("Submitted", "Drafting"),
    }
    editorial_transitions = {
        ("Submitted", "In Review"),
        ("In Review", "Returned"),
        ("In Review", "Ready for Publish"),
    }
    transition = (current_status, next_status)
    if transition in author_transitions:
        if not is_author:
            return {"ok": False, "status": 403, "error": "Only an accepted story author can submit or unsubmit this story."}
        return {"ok": True, "idempotent": False, "publicationUrl": ""}
    if transition in editorial_transitions:
        if role not in {ROLE_ADMIN, ROLE_EDITOR}:
            return {"ok": False, "status": 403, "error": "Only editors and admins can move stories through review."}
        return {"ok": True, "idempotent": False, "publicationUrl": ""}
    if transition == ("Ready for Publish", "Published"):
        if role != ROLE_ADMIN:
            return {"ok": False, "status": 403, "error": "Only admins can publish stories."}
        return {"ok": True, "idempotent": False, "publicationUrl": safe_publication_url}
    return {"ok": False, "status": 409, "error": "That story status change skips the required newsroom workflow."}


def _story_archive_authors(story: dict) -> list[str]:
    names = []
    primary = str(story.get("writer") or story.get("owner") or story.get("author") or "").strip()
    for value in [primary, *_coerce_list_field(story.get("authors"))]:
        name = str(value or "").strip()
        if name and name not in names:
            names.append(name)
    collaborators = story.get("collaborators") if isinstance(story.get("collaborators"), list) else []
    for collaborator in collaborators:
        if not isinstance(collaborator, dict) or str(collaborator.get("status") or "").lower() != "accepted":
            continue
        name = str(collaborator.get("name") or collaborator.get("email") or "").strip()
        if name and name not in names:
            names.append(name)
    if not names:
        fallback = str(story.get("writerEmail") or story.get("ownerEmail") or "Staff").strip()
        names.append(fallback or "Staff")
    return names


def _article_archive_document(story: dict, publication_url: str, published_at: datetime) -> dict:
    workspace_id = str(story.get("workspaceId") or "").strip()
    source_story_id = _doc_public_id(story, "storyId")
    title = str(story.get("title") or story.get("storyTitle") or "Untitled story").strip()
    section = str(story.get("section") or "News").strip() or "News"
    authors = _story_archive_authors(story)
    return {
        "workspaceId": workspace_id,
        "sourceStoryId": source_story_id,
        "title": title,
        "articleTitle": title,
        "url": publication_url,
        "articleUrl": publication_url,
        "canonicalUrl": _canonical_article_url(publication_url),
        "urlNorm": _canonical_article_url(publication_url),
        "authors": authors,
        "author": ", ".join(authors),
        "byline": ", ".join(authors),
        "section": section,
        "category": section,
        "categories": [section],
        "tags": [section],
        "status": "Published",
        "datePublished": published_at,
        "publishedAt": published_at,
        "datePublishedSort": published_at,
        "createdAt": published_at,
        "updatedAt": published_at,
    }


def _upsert_published_article(story: dict, publication_url: str, published_at: datetime) -> dict:
    article = _article_archive_document(story, publication_url, published_at)
    key = {"workspaceId": article["workspaceId"], "sourceStoryId": article["sourceStoryId"]}
    created_at = article.pop("createdAt")
    articles_col.update_one(
        key,
        {"$set": article, "$setOnInsert": {"createdAt": created_at}},
        upsert=True,
    )
    return {**article, "createdAt": created_at}


@app.post("/api/auth/register")
def api_register():
    payload = _request_payload()
    email = str(payload.get("email") or "").strip()
    first_name = str(payload.get("firstName") or "").strip()
    last_name = str(payload.get("lastName") or "").strip()
    password = payload.get("password") if isinstance(payload.get("password"), str) else ""
    confirm = payload.get("confirmPassword") if isinstance(payload.get("confirmPassword"), str) else ""

    if not first_name or not last_name:
        return jsonify({"ok": False, "error": "First and last name are required."}), 400
    if len(first_name) > 80 or len(last_name) > 80:
        return jsonify({"ok": False, "error": "Names must be 80 characters or fewer."}), 400
    if not email or not password:
        return jsonify({"ok": False, "error": "Email and password are required."}), 400
    if _is_auth_rate_limited(email):
        return jsonify({"ok": False, "error": "Too many failed attempts. Please wait and try again."}), 429
    if not is_valid_email(email):
        return jsonify({"ok": False, "error": "Please enter a valid email address."}), 400
    if confirm and confirm != password:
        return jsonify({"ok": False, "error": "Passwords do not match."}), 400
    if len(password) < 8:
        return jsonify({"ok": False, "error": "Password must be at least 8 characters."}), 400
    if len(password) > 256:
        return jsonify({"ok": False, "error": "Password must be 256 characters or fewer."}), 400
    if _is_ip_action_rate_limited(
        "registration", REGISTRATION_RATE_LIMIT_MAX, REGISTRATION_RATE_LIMIT_WINDOW_SECONDS,
    ):
        return jsonify({"ok": False, "error": "Too many registration attempts. Please wait and try again."}), 429
    if find_user_by_email(email):
        return jsonify({"ok": False, "error": "Email is already registered."}), 409

    try:
        user_doc = create_user(email, password, first_name, last_name)
    except DuplicateKeyError:
        return jsonify({"ok": False, "error": "Email is already registered."}), 409

    _login_user_doc(user_doc, remember=False)
    _clear_auth_failures(email)
    return jsonify({"ok": True, "redirect": _auth_redirect_for_role(user_doc.get("role"))})


@app.post("/api/auth/login")
def api_login():
    payload = _request_payload()
    email = str(payload.get("email") or "").strip()
    password = payload.get("password") if isinstance(payload.get("password"), str) else ""
    remember = _payload_bool(payload.get("remember"))

    if not email or not password:
        return jsonify({"ok": False, "error": "Email and password are required."}), 400
    if _is_auth_rate_limited(email):
        return jsonify({"ok": False, "error": "Too many failed attempts. Please wait and try again."}), 429
    if not is_valid_email(email):
        check_password_hash(DUMMY_PASSWORD_HASH, password[:256])
        return jsonify({"ok": False, "error": "Invalid email or password."}), 401
    if len(password) > 256:
        check_password_hash(DUMMY_PASSWORD_HASH, password[:256])
        return jsonify({"ok": False, "error": "Invalid email or password."}), 401
    user_doc = find_user_by_email(email)
    if not user_doc:
        check_password_hash(DUMMY_PASSWORD_HASH, password)
        return jsonify({"ok": False, "error": "Invalid email or password."}), 401

    password_hash = str(user_doc.get("passwordHash") or "")
    candidate_hash = password_hash or DUMMY_PASSWORD_HASH
    try:
        password_matches = check_password_hash(candidate_hash, password)
    except (TypeError, ValueError):
        password_matches = False
    if not password_hash or not password_matches:
        return jsonify({"ok": False, "error": "Invalid email or password."}), 401
    _login_user_doc(user_doc, remember=remember)
    _clear_auth_failures(email)
    users_col.update_one({"_id": user_doc["_id"]}, {"$set": {"lastLoginAt": _now_iso(), "authProviders.password": True}})
    return jsonify({"ok": True, "redirect": _auth_redirect_for_role(user_doc.get("role"))})


@app.get("/api/auth/session")
def api_session():
    user_doc = _current_user_doc()
    if not user_doc:
        return jsonify({"ok": True, "authenticated": False})
    return jsonify({
        "ok": True,
        "authenticated": True,
        "user": _serialize_user(user_doc),
        "workspace": _serialize_workspace(_workspace_for_user(user_doc)),
        "csrfToken": _ensure_csrf_token(),
    })


@app.post("/api/workspaces/join")
@require_auth
def api_join_workspace():
    user_doc = _current_user_doc()
    code = normalize_workspace_code(_request_payload().get("code"))
    if _is_join_rate_limited(user_doc):
        return jsonify({"ok": False, "error": "Too many workspace join attempts. Please wait and try again."}), 429
    if not code or len(code) > 32:
        return jsonify({"ok": False, "error": "Enter a workspace code."}), 400
    workspace = workspaces_col.find_one({"joinCode": code})
    if not workspace:
        return jsonify({"ok": False, "error": "That workspace code was not found."}), 404

    workspace_id = str(workspace.get("publicId") or workspace.get("_id"))
    current_workspace_id = _workspace_id_for_user(user_doc)
    if current_workspace_id:
        if current_workspace_id != workspace_id:
            return jsonify({"ok": False, "error": "You already belong to a different workspace."}), 409
        return jsonify({
            "ok": True,
            "alreadyMember": True,
            "user": _serialize_user(user_doc),
            "workspace": _serialize_workspace(workspace),
        })

    result = users_col.update_one({
        "_id": user_doc["_id"],
        "$or": [
            {"workspaceId": {"$exists": False}},
            {"workspaceId": None},
            {"workspaceId": ""},
        ],
    }, {"$set": {
        "workspaceId": workspace_id,
        "role": ROLE_GUEST,
        "joinedWorkspaceAt": _now_iso(),
        "updatedAt": _now_iso(),
    }})
    updated = users_col.find_one({"_id": user_doc["_id"]}) or user_doc
    if result.matched_count == 0 and _workspace_id_for_user(updated) != workspace_id:
        return jsonify({"ok": False, "error": "Your workspace membership changed. Refresh and try again."}), 409
    return jsonify({
        "ok": True,
        "alreadyMember": result.matched_count == 0,
        "user": _serialize_user(updated),
        "workspace": _serialize_workspace(workspace),
    })


@app.get("/api/workspace")
@require_auth
def api_workspace():
    current_user = _current_user_doc()
    workspace = _workspace_for_user(current_user)
    if not workspace:
        return jsonify({"ok": False, "error": "Workspace not found."}), 404
    include_join_code = normalize_role(current_user.get("role")) == ROLE_ADMIN
    return jsonify({"ok": True, "workspace": _serialize_workspace(workspace, include_join_code=include_join_code)})


@app.patch("/api/workspace")
@require_roles(ROLE_ADMIN)
def api_update_workspace():
    current_user = _current_user_doc()
    workspace = _workspace_for_user(current_user)
    if not workspace:
        return jsonify({"ok": False, "error": "Workspace not found."}), 404

    payload = _request_payload()
    name = str(payload.get("name") or "").strip()
    publication_url = str(payload.get("publicationUrl") or "").strip().rstrip("/")

    if not name:
        return jsonify({"ok": False, "error": "Workspace name is required."}), 400
    if len(name) > 120:
        return jsonify({"ok": False, "error": "Workspace name must be 120 characters or fewer."}), 400

    safe_publication_url = _safe_http_url(publication_url, max_length=500)
    parsed_publication_url = urlparse(safe_publication_url) if safe_publication_url else None
    article_domain = _normalized_publication_hostname(parsed_publication_url.hostname or "") if parsed_publication_url else ""
    if (
        not safe_publication_url
        or not article_domain
        or (parsed_publication_url.port is not None and parsed_publication_url.port not in {80, 443})
    ):
        return jsonify({"ok": False, "error": "Enter a valid publication URL beginning with http:// or https://."}), 400

    workspaces_col.update_one({"_id": workspace["_id"]}, {"$set": {
        "name": name,
        "publicationUrl": safe_publication_url,
        "articleDomain": article_domain,
        "updatedAt": _now_iso(),
    }})
    updated = workspaces_col.find_one({"_id": workspace["_id"]}) or workspace
    return jsonify({"ok": True, "workspace": _serialize_workspace(updated, include_join_code=True)})


@app.post("/api/workspace/join-code/rotate")
@require_roles(ROLE_ADMIN)
def api_rotate_workspace_join_code():
    current_user = _current_user_doc()
    workspace = _workspace_for_user(current_user)
    if not workspace:
        return jsonify({"ok": False, "error": "Workspace not found."}), 404
    try:
        next_code = _new_workspace_code()
        result = workspaces_col.update_one(
            {"_id": workspace["_id"], "joinCode": workspace.get("joinCode")},
            {"$set": {
                "joinCode": next_code,
                "joinCodeRotatedAt": _now_iso(),
                "joinCodeRotatedBy": normalize_email(current_user.get("email") or ""),
                "updatedAt": _now_iso(),
            }},
        )
    except DuplicateKeyError:
        return jsonify({"ok": False, "error": "Could not rotate the join code. Please try again."}), 409
    except Exception:
        return jsonify({"ok": False, "error": "Could not rotate the join code. Please try again."}), 503
    if result.modified_count != 1:
        return jsonify({"ok": False, "error": "The join code changed in another session. Refresh and try again."}), 409
    updated = workspaces_col.find_one({"_id": workspace["_id"]}) or {**workspace, "joinCode": next_code}
    try:
        activity_col.insert_one({
            "workspaceId": _workspace_id_for_user(current_user),
            "entityType": "workspace",
            "entityId": _workspace_id_for_user(current_user),
            "eventType": "join_code_rotated",
            "text": "Workspace join code rotated.",
            "actorId": str(current_user.get("_id")),
            "actorEmail": normalize_email(current_user.get("email") or ""),
            "actorName": _user_display_name(current_user),
            "createdAt": datetime.now(timezone.utc),
        })
    except Exception:
        pass
    return jsonify({"ok": True, "workspace": _serialize_workspace(updated, include_join_code=True)})


@app.post("/api/auth/logout")
def api_logout():
    user_doc = _current_user_doc()
    if user_doc:
        users_col.update_one(
            {"_id": user_doc["_id"]},
            [{"$set": {
                "authVersion": {
                    "$add": [
                        {"$cond": [{"$isNumber": "$authVersion"}, "$authVersion", 1]},
                        1,
                    ],
                },
                "lastLogoutAt": _now_iso(),
            }}],
        )
    session.clear()
    return jsonify({"ok": True})


@app.get("/api/drive/picker-config")
@require_auth
def api_drive_picker_config():
    missing = _google_drive_config_missing()
    return jsonify({
        "ok": not missing,
        "enabled": not missing,
        "missing": missing,
        "apiKey": GOOGLE_PICKER_API_KEY if not missing else "",
        "appId": GOOGLE_PICKER_APP_ID if not missing else "",
        "clientId": GOOGLE_PICKER_CLIENT_ID if not missing else "",
        "scope": GOOGLE_DRIVE_SCOPE,
        "reauthUrl": _google_reauth_url(request.args.get("next") or request.path),
        "error": f"Google Drive picker is missing {', '.join(missing)}." if missing else "",
    })


@app.get("/api/drive/picker-token")
@require_auth
def api_drive_picker_token():
    user_doc = _current_user_doc()
    missing = _google_drive_config_missing()
    if missing:
        return jsonify({
            "ok": False,
            "error": f"Google Drive picker is missing {', '.join(missing)}.",
            "missing": missing,
        }), 503
    access_token, token_error = _google_access_token_for_user(user_doc)
    if token_error:
        return _google_oauth_error(token_error)
    oauth = user_doc.get("googleOAuth") if isinstance(user_doc.get("googleOAuth"), dict) else {}
    return jsonify({
        "ok": True,
        "accessToken": access_token,
        "expiresAt": _parse_oauth_expiry(oauth.get("expiresAt")).isoformat() if _parse_oauth_expiry(oauth.get("expiresAt")) else "",
        "scope": GOOGLE_DRIVE_SCOPE,
    })


@app.get("/api/auth/google/start")
def api_google_start():
    missing_config = _google_missing_config()
    if missing_config:
        return jsonify({
            "ok": False,
            "error": f"Google sign-in is missing {', '.join(missing_config)}.",
            "missing": missing_config,
        }), 503
    if _is_oauth_start_rate_limited():
        return jsonify({"ok": False, "error": "Too many sign-in attempts. Please wait and try again."}), 429

    redirect_uri = _google_redirect_uri()
    frontend_origin = _origin_from_redirect_uri() or _sanitize_frontend_origin(request.args.get("origin") or "")
    intent = _normalize_google_intent(request.args.get("intent") or "")
    try:
        state = _make_google_state(request.args.get("next") or "", frontend_origin, intent)
    except Exception:
        return jsonify({"ok": False, "error": "Could not start Google sign-in. Please try again."}), 503
    params = {
        "client_id": GOOGLE_CLIENT_ID,
        "redirect_uri": redirect_uri,
        "response_type": "code",
        "scope": _google_scope_string(),
        "state": state,
        "access_type": "offline",
        "include_granted_scopes": "true",
        "prompt": "select_account consent",
    }
    return jsonify({"ok": True, "authUrl": f"{GOOGLE_AUTH_ENDPOINT}?{urlencode(params)}"})


@app.get("/api/auth/google/callback")
def api_google_callback():
    state = (request.args.get("state") or "").strip()
    state_data, state_error = _load_google_state(state)
    frontend_origin = state_data.get("frontendOrigin", "") if state_data else _origin_from_redirect_uri()
    next_url = state_data.get("next", "") if state_data else ""
    intent = _normalize_google_intent(state_data.get("intent", "") if state_data else "")

    if request.args.get("error"):
        return _google_error_redirect("Google sign-in was cancelled.", frontend_origin)

    if state_error:
        return _google_error_redirect(state_error, frontend_origin)

    code = (request.args.get("code") or "").strip()
    if not code:
        return _google_error_redirect("Google did not return an authorization code.", frontend_origin)
    if not _is_google_configured():
        return _google_error_redirect("Google sign-in is not configured.", frontend_origin)

    try:
        token_response = requests.post(
            GOOGLE_TOKEN_ENDPOINT,
            data={
                "client_id": GOOGLE_CLIENT_ID,
                "client_secret": GOOGLE_CLIENT_SECRET,
                "code": code,
                "grant_type": "authorization_code",
                "redirect_uri": _google_redirect_uri(),
            },
            timeout=10,
        )
        token_response.raise_for_status()
        token_payload = token_response.json()
        access_token = token_payload.get("access_token")
        if not access_token:
            return _google_error_redirect("Google did not return an access token.", frontend_origin)

        profile_response = requests.get(
            GOOGLE_USERINFO_ENDPOINT,
            headers={"Authorization": f"Bearer {access_token}"},
            timeout=10,
        )
        profile_response.raise_for_status()
        profile = profile_response.json()
    except Exception:
        return _google_error_redirect("Google sign-in failed. Please try again.", frontend_origin)

    email = normalize_email(profile.get("email") or "")
    if not email or not profile.get("email_verified"):
        return _google_error_redirect("Google account email is not verified.", frontend_origin)
    if not _google_email_domain_allowed(email):
        return _google_error_redirect("This Google account is not allowed for Falcon Newsroom.", frontend_origin)

    try:
        user_doc = upsert_google_user(profile, allow_create=intent == GOOGLE_INTENT_SIGNUP)
    except ValueError as exc:
        return _google_error_redirect(str(exc) or "Could not link this Google account.", frontend_origin)
    except DuplicateKeyError:
        google_id = str(profile.get("sub") or "").strip()
        existing = find_user_by_google_id(google_id)
        if not existing or normalize_email(existing.get("email") or "") != email or _google_id_from_doc(existing) != google_id:
            return _google_error_redirect("Could not link this Google account.", frontend_origin)
        user_doc = existing
    except Exception:
        return _google_error_redirect("Could not create a Falcon account from Google.", frontend_origin)

    _store_google_oauth_tokens(user_doc, token_payload)
    _login_user_doc(user_doc, remember=False)
    _clear_auth_failures(email)
    return redirect(_frontend_redirect_url(
        next_url or "/dashboard",
        frontend_origin,
    ))


@app.get("/api/admin/users")
@require_roles(ROLE_ADMIN)
def api_admin_users():
    user_doc = _current_user_doc()
    docs = list(users_col.find(_workspace_query(user_doc)).sort([("role", ASCENDING), ("email", ASCENDING)]))
    return jsonify({"ok": True, "users": [_serialize_user(doc) for doc in docs], "roles": sorted(VALID_ROLES)})


@app.patch("/api/admin/users/<user_id>/role")
@require_roles(ROLE_ADMIN)
def api_admin_update_user_role(user_id: str):
    current_user = _current_user_doc()
    oid = _object_id_or_none(user_id)
    if not oid:
        return jsonify({"ok": False, "error": "Invalid user id."}), 400

    payload = _request_payload()
    next_role = normalize_role(payload.get("role"))
    if str(payload.get("role") or "").strip().lower() not in VALID_ROLES:
        return jsonify({"ok": False, "error": "Role must be admin, editor, writer, or guest."}), 400
    if str(current_user.get("_id")) == str(oid):
        return jsonify({"ok": False, "error": "You cannot change your own role while signed in."}), 400

    target = users_col.find_one({"_id": oid, "workspaceId": _workspace_id_for_user(current_user)})
    if not target:
        return jsonify({"ok": False, "error": "User not found."}), 404
    workspace_id = _workspace_id_for_user(current_user)
    lock_token = _acquire_admin_mutation_lock(workspace_id)
    if not lock_token:
        return jsonify({"ok": False, "error": "Another access change is in progress. Please try again."}), 409
    try:
        target = users_col.find_one({"_id": oid, "workspaceId": workspace_id})
        if not target:
            return jsonify({"ok": False, "error": "User not found."}), 404
        if normalize_role(target.get("role")) == ROLE_ADMIN and next_role != ROLE_ADMIN:
            admin_count = users_col.count_documents({"role": ROLE_ADMIN, "workspaceId": workspace_id})
            if admin_count <= 1:
                return jsonify({"ok": False, "error": "At least one admin must remain."}), 400
        result = users_col.update_one(
            {"_id": oid, "workspaceId": workspace_id},
            {"$set": {"role": next_role, "updatedAt": _now_iso()}},
        )
        if result.matched_count != 1:
            return jsonify({"ok": False, "error": "Membership changed in another session. Refresh and try again."}), 409
    finally:
        try:
            _release_admin_mutation_lock(workspace_id, lock_token)
        except Exception:
            pass
    updated = users_col.find_one({"_id": oid}) or {}
    return jsonify({"ok": True, "user": _serialize_user(updated)})


@app.delete("/api/admin/users/<user_id>/membership")
@require_roles(ROLE_ADMIN)
def api_admin_remove_user_membership(user_id: str):
    current_user = _current_user_doc()
    oid = _object_id_or_none(user_id)
    if not oid:
        return jsonify({"ok": False, "error": "Invalid user id."}), 400
    if str(current_user.get("_id")) == str(oid):
        return jsonify({"ok": False, "error": "You cannot remove your own workspace membership."}), 400
    workspace_id = _workspace_id_for_user(current_user)
    target = users_col.find_one({"_id": oid, "workspaceId": workspace_id})
    if not target:
        return jsonify({"ok": False, "error": "User not found."}), 404
    lock_token = _acquire_admin_mutation_lock(workspace_id)
    if not lock_token:
        return jsonify({"ok": False, "error": "Another access change is in progress. Please try again."}), 409
    try:
        target = users_col.find_one({"_id": oid, "workspaceId": workspace_id})
        if not target:
            return jsonify({"ok": False, "error": "User not found."}), 404
        if normalize_role(target.get("role")) == ROLE_ADMIN:
            admin_count = users_col.count_documents({"role": ROLE_ADMIN, "workspaceId": workspace_id})
            if admin_count <= 1:
                return jsonify({"ok": False, "error": "At least one admin must remain."}), 400
        removed_at = _now_iso()
        result = users_col.update_one(
            {"_id": oid, "workspaceId": workspace_id},
            [
                {"$set": {
                    "authVersion": {
                        "$add": [
                            {"$cond": [{"$isNumber": "$authVersion"}, "$authVersion", 1]},
                            1,
                        ],
                    },
                    "workspaceRemovedAt": removed_at,
                    "workspaceRemovedBy": normalize_email(current_user.get("email") or ""),
                    "updatedAt": removed_at,
                }},
                {"$unset": ["workspaceId", "role", "joinedWorkspaceAt"]},
            ],
        )
        if result.modified_count != 1:
            return jsonify({"ok": False, "error": "Membership changed in another session. Refresh and try again."}), 409
    finally:
        try:
            _release_admin_mutation_lock(workspace_id, lock_token)
        except Exception:
            pass
    try:
        activity_col.insert_one({
            "workspaceId": workspace_id,
            "entityType": "user",
            "entityId": str(oid),
            "eventType": "membership_removed",
            "text": f"{_user_display_name(target)} was removed from the workspace.",
            "actorId": str(current_user.get("_id")),
            "actorEmail": normalize_email(current_user.get("email") or ""),
            "actorName": _user_display_name(current_user),
            "createdAt": datetime.now(timezone.utc),
        })
    except Exception:
        pass
    return jsonify({"ok": True, "removedUserId": str(oid)})


@app.get("/api/stories")
@require_auth
def api_stories():
    user_doc = _current_user_doc()
    query = _story_query_for_user(user_doc)
    if query is None:
        return jsonify({"ok": False, "error": "Stories are not available to guests."}), 403
    docs = list(stories_col.find(query).sort([("updatedAt", -1), ("_id", -1)]))
    return jsonify({"ok": True, "stories": [_story_to_api(doc) for doc in docs]})


@app.patch("/api/stories/<story_id>")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_update_story(story_id: str):
    user_doc = _current_user_doc()
    story = _find_owned_story_or_404(story_id, user_doc)
    if not story:
        return jsonify({"ok": False, "error": "Story not found."}), 404

    payload = _request_payload()
    role = _current_user_role(user_doc)
    expected_status = str(story.get("status") or "Assigned").strip()
    now = datetime.now(timezone.utc)
    now_iso = now.isoformat(timespec="milliseconds")
    update = {"updatedAt": now_iso}
    unset = {}
    increment = {}
    push_attachment = None
    transition = None
    next_status = ""

    if "publicationUrl" in payload and str(payload.get("status") or "").strip() != "Published":
        return jsonify({"ok": False, "error": "Publication URL is only accepted while publishing a story."}), 400

    if "status" in payload:
        if "googleDocUrl" in payload or "documentUrl" in payload:
            return jsonify({"ok": False, "error": "Update attached work separately from workflow status."}), 400
        next_status = str(payload.get("status") or "").strip()
        collaborator = _story_collaborator_for_user(story, user_doc)
        is_story_author = _story_primary_owned_by_user(story, user_doc) or (
            collaborator is not None and collaborator.get("role") == "edit"
        )
        raw_publication_url = str(payload.get("publicationUrl") or "").strip()
        publication_url = _validated_publication_url(raw_publication_url, _workspace_for_user(user_doc)) if raw_publication_url else ""
        transition = _story_status_transition_decision(
            role,
            expected_status,
            next_status,
            is_story_author,
            publication_url,
            story.get("publicationUrl") or "",
        )
        if not transition.get("ok"):
            return jsonify({"ok": False, "error": transition.get("error")}), int(transition.get("status") or 400)
        if next_status == "Published" and _publication_url_conflicts(story, transition["publicationUrl"]):
            return jsonify({"ok": False, "error": "That publication URL is already archived for another story."}), 409
        if transition.get("idempotent"):
            published_at = story.get("publishedAt")
            if not isinstance(published_at, datetime):
                try:
                    published_at = datetime.fromisoformat(str(published_at or "").replace("Z", "+00:00"))
                except Exception:
                    published_at = now
            if published_at.tzinfo is None:
                published_at = published_at.replace(tzinfo=timezone.utc)
            try:
                _upsert_published_article(story, transition["publicationUrl"], published_at)
            except DuplicateKeyError:
                return jsonify({"ok": False, "error": "That publication URL is already archived for another story."}), 409
            except Exception:
                return jsonify({"ok": False, "error": "The article archive could not be synchronized. Please try again."}), 503
            return jsonify({"ok": True, "story": _story_to_api(story)})
        if next_status == "Submitted" and not _story_attachments_to_api(story):
            return jsonify({"ok": False, "error": "Attach work before submitting this story."}), 400
        update["status"] = next_status
        timestamp_fields = {
            "Submitted": "submittedAt",
            "Drafting": "unsubmittedAt",
            "In Review": "reviewStartedAt",
            "Returned": "returnedAt",
            "Ready for Publish": "readyForPublishAt",
            "Published": "publishedAt",
        }
        if next_status in timestamp_fields:
            update[timestamp_fields[next_status]] = now_iso
        if next_status == "Returned":
            increment["revisionCount"] = 1
        if next_status == "Published":
            update["publicationUrl"] = transition["publicationUrl"]
    if any(field in payload for field in DUE_DATE_FIELDS):
        if role not in {ROLE_ADMIN, ROLE_EDITOR}:
            return jsonify({"ok": False, "error": "Only editors and admins can update story due dates."}), 403
        next_deadline = _first_text_value(*(payload.get(field) for field in DUE_DATE_FIELDS))
        if len(next_deadline) > 80:
            return jsonify({"ok": False, "error": "Due date is too long."}), 400
        update["deadline"] = next_deadline
        update["dueDate"] = next_deadline
    if "googleDocUrl" in payload or "documentUrl" in payload:
        if not _can_edit_story_content(story, user_doc):
            return jsonify({"ok": False, "error": "You can comment on this story, but cannot edit its attached work."}), 403
        raw_url = str((payload.get("documentUrl") if "documentUrl" in payload else payload.get("googleDocUrl")) or "")
        next_url = _safe_http_url(raw_url) if raw_url else ""
        if raw_url and not next_url:
            return jsonify({"ok": False, "error": "Enter a valid http or https link."}), 400
        if next_url:
            attachment_count, _stored_bytes = _story_attachment_usage(story)
            existing_urls = {
                str(item.get("url") or item.get("webViewLink") or "")
                for item in _story_attachments_to_api(story)
            }
            if attachment_count >= MAX_STORY_ATTACHMENTS and next_url not in existing_urls:
                return jsonify({"ok": False, "error": f"A story can have at most {MAX_STORY_ATTACHMENTS} attachments."}), 409
            update["googleDocUrl"] = next_url
            update["attachmentType"] = "multiple"
            push_attachment = {
                "id": _new_attachment_id("link"),
                "type": "link",
                "provider": "manual-link",
                "name": "Story link",
                "url": next_url,
                "addedBy": _attachment_added_by(user_doc),
                "addedAt": datetime.now(timezone.utc),
                "permissionStatus": "manual",
            }
        else:
            update["googleDocUrl"] = ""
            update["attachmentType"] = ""
            unset.update({
                "attachments": "",
                "attachmentFileId": "",
                "attachmentName": "",
                "attachmentContentType": "",
                "attachmentSize": "",
                "attachmentUploadedAt": "",
                "driveAttachment": "",
            })
    if len(update) == 1 and not unset and not push_attachment:
        return jsonify({"ok": False, "error": "No editable fields provided."}), 400

    previous_file_id = story.get("attachmentFileId")
    operation = {"$set": update}
    if unset:
        operation["$unset"] = unset
    if increment:
        operation["$inc"] = increment
    if push_attachment:
        operation["$set"]["attachments"] = _merged_attachment_items(story, push_attachment)
    if transition:
        mutation_scope = _editable_story_query(user_doc) if is_story_author and (expected_status, next_status) in {
            ("Assigned", "Submitted"),
            ("Reporting", "Submitted"),
            ("Drafting", "Submitted"),
            ("Returned", "Submitted"),
            ("Needs Revision", "Submitted"),
            ("Submitted", "Drafting"),
        } else _workspace_query(user_doc)
        mutation_scope = {"$and": [mutation_scope, {"status": expected_status}]}
    elif "googleDocUrl" in payload or "documentUrl" in payload:
        mutation_scope = _story_content_mutation_query(user_doc, expected_status)
    else:
        mutation_scope = {"$and": [_workspace_query(user_doc), {"status": expected_status}]}
    result = stories_col.update_one({"$and": [{"_id": story["_id"]}, mutation_scope]}, operation)
    if result.matched_count == 0:
        current = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {}
        return jsonify({"ok": False, "error": "Story access or status changed. Refresh and try again.", "story": _story_to_api(current) if current else None}), 409
    if unset and previous_file_id:
        try:
            stored = story_files.get(_object_id_or_none(str(previous_file_id)))
            if _gridfs_file_matches_story(stored, story):
                story_files.delete(stored._id)
        except Exception:
            pass
    updated = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {**story, **update}
    if update.get("status") == "Published":
        archive_error = ""
        archive_status = 503
        try:
            _upsert_published_article(updated, update["publicationUrl"], now)
        except DuplicateKeyError:
            archive_error = "That publication URL is already archived for another story."
            archive_status = 409
        except Exception:
            archive_error = "The article archive could not be synchronized. Please try again."
        if archive_error:
            rollback = stories_col.update_one(
                {
                    "_id": story["_id"],
                    "workspaceId": _workspace_id_for_user(user_doc),
                    "status": "Published",
                    "publicationUrl": update["publicationUrl"],
                    "publishedAt": now_iso,
                },
                {
                    "$set": {"status": "Ready for Publish", "updatedAt": _now_iso()},
                    "$unset": {"publicationUrl": "", "publishedAt": ""},
                },
            )
            current = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or updated
            if rollback.matched_count == 0:
                return jsonify({
                    "ok": False,
                    "error": "Publication could not be finalized safely. Refresh before trying again.",
                    "story": _story_to_api(current),
                }), 500
            return jsonify({
                "ok": False,
                "error": archive_error,
                "story": _story_to_api(current),
            }), archive_status
    if "status" in update and update["status"] != expected_status:
        _record_status_activity("story", _doc_public_id(story, "storyId"), expected_status, update["status"], user_doc)
    return jsonify({"ok": True, "story": _story_to_api(updated)})


VALID_STORY_COLLABORATOR_ROLES = {"comment", "edit"}


def _story_collaborator_payload(item: dict, role: str, message: str, actor_doc: dict) -> dict:
    email = normalize_email(item.get("email") or "")
    return {
        "userId": str(item.get("userId") or item.get("id") or ""),
        "email": email,
        "name": str(item.get("name") or email).strip(),
        "role": "edit",
        "status": "pending",
        "message": message,
        "invitedBy": _user_display_name(actor_doc),
        "invitedByEmail": normalize_email(actor_doc.get("email") or ""),
        "invitedAt": datetime.now(timezone.utc),
    }


def _approval_invite_emails(raw_value) -> list[str]:
    if isinstance(raw_value, str):
        candidates = re.split(r"[,;\s]+", raw_value)
    elif isinstance(raw_value, list):
        candidates = raw_value
    else:
        candidates = []
    emails = []
    for value in candidates:
        email = normalize_email(str(value or ""))
        if email and email not in emails:
            emails.append(email)
    return emails


def _add_approval_collaborators(story: dict, emails: list[str], message: str, actor_doc: dict) -> str:
    if not emails:
        return ""
    invalid = [email for email in emails if not is_valid_email(email)]
    if invalid:
        return f"Invite skipped. Enter valid email addresses: {', '.join(invalid)}."

    actor_email = normalize_email(actor_doc.get("email") or "")
    primary_emails = {normalize_email(story.get("writerEmail") or ""), normalize_email(story.get("ownerEmail") or "")}
    current_collaborators = _story_collaborators(story)
    collaborator_by_email = {item["email"]: item for item in current_collaborators}
    invited = []
    skipped = []
    for email in emails:
        if email == actor_email or email in primary_emails or email in collaborator_by_email:
            skipped.append(email)
            continue
        invited_user = find_user_by_email(email)
        if not invited_user or _workspace_id_for_user(invited_user) != _workspace_id_for_user(actor_doc):
            skipped.append(email)
            continue
        invited_role = _current_user_role(invited_user)
        if invited_role != ROLE_WRITER:
            skipped.append(email)
            continue
        collaborator = _story_collaborator_payload(_serialize_user(invited_user), "comment", message, actor_doc)
        collaborator_by_email[email] = collaborator
        invited.append(collaborator)

    if invited:
        next_collaborators = sorted(collaborator_by_email.values(), key=lambda item: item.get("email", ""))
        stories_col.update_one({"_id": story["_id"]}, {"$set": {"collaborators": next_collaborators, "updatedAt": _now_iso()}})
        activity_col.insert_one({
            "workspaceId": _workspace_id_for_user(actor_doc),
            "entityType": "story",
            "entityId": _doc_public_id(story, "storyId"),
            "eventType": "collaborator_invite",
            "text": f"Invited {len(invited)} collaborator{'s' if len(invited) != 1 else ''} from pitch approval.",
            "actorId": str(actor_doc.get("_id")),
            "actorEmail": actor_doc.get("email", ""),
            "actorName": _user_display_name(actor_doc),
            "createdAt": datetime.now(timezone.utc),
        })
    if skipped:
        return f"Approved, but {len(skipped)} invite{'s were' if len(skipped) != 1 else ' was'} skipped."
    return ""
@app.post("/api/stories/<story_id>/collaborators")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_invite_story_collaborators(story_id: str):
    user_doc = _current_user_doc()
    story = _find_owned_story_or_404(story_id, user_doc)
    if not story:
        return jsonify({"ok": False, "error": "Story not found."}), 404
    expected_status = str(story.get("status") or "Assigned").strip()
    if expected_status in LOCKED_COLLABORATOR_STATUSES:
        return jsonify({"ok": False, "error": "Collaborators are locked once a story is ready for publication."}), 409
    if not _can_manage_story_collaborators(story, user_doc):
        return jsonify({"ok": False, "error": "Only story authors or editors can invite co-authors."}), 403

    payload = _request_payload()
    raw_emails = payload.get("emails")
    if isinstance(raw_emails, str):
        candidates = re.split(r"[,;\s]+", raw_emails)
    elif isinstance(raw_emails, list):
        candidates = raw_emails
    else:
        candidates = []
    emails = []
    for value in candidates:
        email = normalize_email(str(value or ""))
        if email and email not in emails:
            emails.append(email)
    if not emails:
        return jsonify({"ok": False, "error": "Enter at least one email address."}), 400
    invalid = [email for email in emails if not is_valid_email(email)]
    if invalid:
        return jsonify({"ok": False, "error": f"Enter valid email addresses: {', '.join(invalid)}."}), 400

    role = "edit"
    message = str(payload.get("message") or "").strip()
    if len(message) > 200:
        return jsonify({"ok": False, "error": "Invite message must be 200 characters or fewer."}), 400

    actor_email = normalize_email(user_doc.get("email") or "")
    primary_emails = {normalize_email(story.get("writerEmail") or ""), normalize_email(story.get("ownerEmail") or "")}
    current_collaborators = _story_collaborators(story)
    collaborator_by_email = {item["email"]: item for item in current_collaborators}
    invited = []
    skipped = []
    for email in emails:
        if email == actor_email or email in primary_emails:
            skipped.append(email)
            continue
        existing_collaborator = collaborator_by_email.get(email)
        if existing_collaborator and existing_collaborator.get("status") in {"pending", "accepted"}:
            skipped.append(email)
            continue
        invited_user = find_user_by_email(email)
        if not invited_user or _workspace_id_for_user(invited_user) != _workspace_id_for_user(user_doc):
            return jsonify({"ok": False, "error": f"No account in this workspace is registered for {email}."}), 404
        invited_role = _current_user_role(invited_user)
        if invited_role in {ROLE_ADMIN, ROLE_EDITOR}:
            skipped.append(email)
            continue
        if invited_role != ROLE_WRITER:
            return jsonify({"ok": False, "error": f"{email} must have writer access before they can collaborate on stories."}), 400
        collaborator = _story_collaborator_payload(_serialize_user(invited_user), role, message, user_doc)
        collaborator_by_email[email] = collaborator
        invited.append(collaborator)

    if not invited and skipped:
        return jsonify({"ok": False, "error": "Those users already have access or a pending invitation."}), 400
    if not invited:
        return jsonify({"ok": False, "error": "No collaborators were added."}), 400

    next_collaborators = sorted(collaborator_by_email.values(), key=lambda item: item.get("email", ""))
    result = stories_col.update_one(
        {"$and": [{"_id": story["_id"]}, _story_collaborator_mutation_query(user_doc, expected_status)]},
        {"$set": {"collaborators": next_collaborators, "updatedAt": _now_iso()}},
    )
    if result.matched_count == 0:
        return jsonify({"ok": False, "error": "Story access or status changed. Refresh and try again."}), 409
    activity_col.insert_one({
        "workspaceId": _workspace_id_for_user(user_doc),
        "entityType": "story",
        "entityId": _doc_public_id(story, "storyId"),
        "eventType": "collaborator_invite",
        "text": f"Invited {len(invited)} author{'s' if len(invited) != 1 else ''}.",
        "recipientEmails": [item.get("email") for item in invited],
        "entityTitle": story.get("title") or story.get("storyTitle") or "Untitled story",
        "actorId": str(user_doc.get("_id")),
        "actorEmail": user_doc.get("email", ""),
        "actorName": _user_display_name(user_doc),
        "createdAt": datetime.now(timezone.utc),
    })
    updated = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {}
    return jsonify({"ok": True, "story": _story_to_api(updated), "collaborators": _story_collaborators(updated)})


@app.delete("/api/stories/<story_id>/collaborators/<path:email>")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_remove_story_collaborator(story_id: str, email: str):
    user_doc = _current_user_doc()
    story = _find_owned_story_or_404(story_id, user_doc)
    if not story:
        return jsonify({"ok": False, "error": "Story not found."}), 404
    expected_status = str(story.get("status") or "Assigned").strip()
    if expected_status in LOCKED_COLLABORATOR_STATUSES:
        return jsonify({"ok": False, "error": "Collaborators are locked once a story is ready for publication."}), 409
    if not _can_manage_story_collaborators(story, user_doc):
        return jsonify({"ok": False, "error": "Only story authors or editors can remove co-authors."}), 403
    target_email = normalize_email(email)
    if not is_valid_email(target_email):
        return jsonify({"ok": False, "error": "Choose a valid collaborator email."}), 400
    current_collaborators = _story_collaborators(story)
    next_collaborators = [item for item in current_collaborators if item.get("email") != target_email]
    if len(next_collaborators) == len(current_collaborators):
        return jsonify({"ok": False, "error": "Collaborator not found."}), 404
    result = stories_col.update_one(
        {"$and": [{"_id": story["_id"]}, _story_collaborator_mutation_query(user_doc, expected_status)]},
        {"$set": {"collaborators": next_collaborators, "updatedAt": _now_iso()}},
    )
    if result.matched_count == 0:
        return jsonify({"ok": False, "error": "Story access or status changed. Refresh and try again."}), 409
    updated = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {}
    return jsonify({"ok": True, "story": _story_to_api(updated), "collaborators": _story_collaborators(updated)})

@app.post("/api/stories/<story_id>/drive-attachment")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_attach_drive_file(story_id: str):
    user_doc = _current_user_doc()
    story = _find_owned_story_or_404(story_id, user_doc)
    if not story:
        return jsonify({"ok": False, "error": "Story not found."}), 404

    if not _can_edit_story_content(story, user_doc):
        return jsonify({"ok": False, "error": "You can comment on this story, but cannot edit its attached work."}), 403
    expected_status = str(story.get("status") or "Assigned").strip()
    attachment_count, _stored_bytes = _story_attachment_usage(story)
    if attachment_count >= MAX_STORY_ATTACHMENTS:
        return jsonify({"ok": False, "error": f"A story can have at most {MAX_STORY_ATTACHMENTS} attachments."}), 409

    payload = _request_payload()
    file_id = str(payload.get("fileId") or payload.get("id") or "").strip()
    access_token, token_error = _google_access_token_for_user(user_doc)
    if token_error:
        return _google_oauth_error(token_error)
    metadata, metadata_error = _drive_file_metadata(file_id, access_token)
    if metadata_error:
        return jsonify({"ok": False, "error": metadata_error}), 400

    canonical_file_id = str(metadata.get("id") or file_id).strip()
    if not DRIVE_FILE_ID_RE.fullmatch(canonical_file_id):
        return jsonify({"ok": False, "error": "Drive file id is invalid."}), 400
    fallback_link = payload.get("url") or payload.get("webViewLink") or ""
    if fallback_link and not _safe_http_url(fallback_link):
        return jsonify({"ok": False, "error": "Drive file link is invalid."}), 400
    web_view_link = _safe_http_url(metadata.get("webViewLink") or "") or _safe_http_url(fallback_link)
    if not web_view_link:
        web_view_link = f"https://drive.google.com/open?id={canonical_file_id}"

    now = datetime.now(timezone.utc)
    drive_attachment = {
        "id": _new_attachment_id("drive"),
        "type": "drive",
        "fileId": canonical_file_id,
        "name": str(metadata.get("name") or payload.get("name") or "Drive file").strip()[:255] or "Drive file",
        "mimeType": str(metadata.get("mimeType") or payload.get("mimeType") or "").strip()[:200],
        "typeLabel": _drive_type_label(metadata.get("mimeType") or payload.get("mimeType"), metadata.get("name") or payload.get("name")),
        "webViewLink": web_view_link,
        "iconLink": _safe_http_url(metadata.get("iconLink") or payload.get("iconUrl") or ""),
        "addedBy": _attachment_added_by(user_doc),
        "addedAt": now,
        "permissionStatus": "not_shared",
        "shareResults": [],
        "lastShareAttemptAt": None,
    }
    update = {
        "updatedAt": _now_iso(),
        "attachmentType": "multiple",
        "attachmentName": drive_attachment["name"],
        "googleDocUrl": drive_attachment["webViewLink"],
    }
    unset = {
        "attachmentFileId": "",
        "attachmentContentType": "",
        "attachmentSize": "",
        "attachmentUploadedAt": "",
        "docUrl": "",
        "driveAttachment": "",
    }
    mutation_scope = _story_content_mutation_query(user_doc, expected_status)
    result = stories_col.update_one(
        {"$and": [{"_id": story["_id"]}, mutation_scope]},
        {"$set": {**update, "attachments": _merged_attachment_items(story, drive_attachment)}, "$unset": unset},
    )
    if result.matched_count == 0:
        return jsonify({"ok": False, "error": "Story access changed. Refresh and try again."}), 409

    shared, share_results, _share_message = _share_drive_attachment_with_editors(drive_attachment, user_doc)
    successful_shares = sum(1 for item in share_results if item.get("ok"))
    permission_status = "shared" if shared else ("partial" if successful_shares else "failed")
    share_attempt_at = datetime.now(timezone.utc)
    share_update = stories_col.update_one(
        {
            "$and": [
                {"_id": story["_id"]},
                mutation_scope,
                {"attachments": {"$elemMatch": {"id": drive_attachment["id"], "type": "drive", "fileId": canonical_file_id}}},
            ]
        },
        {"$set": {
            "attachments.$[attachment].permissionStatus": permission_status,
            "attachments.$[attachment].shareResults": share_results,
            "attachments.$[attachment].lastShareAttemptAt": share_attempt_at,
        }},
        array_filters=[{
            "attachment.id": drive_attachment["id"],
            "attachment.type": "drive",
            "attachment.fileId": canonical_file_id,
        }],
    )
    metadata_saved = share_update.matched_count > 0
    updated = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {}
    warning = "" if shared and metadata_saved else "Drive file attached, but it could not be shared with every editor or admin."
    return jsonify({
        "ok": True,
        "story": _story_to_api(updated),
        "warning": warning,
        "shareResults": share_results,
    })


@app.post("/api/stories/<story_id>/attachment")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_upload_story_attachment(story_id: str):
    user_doc = _current_user_doc()
    story = _find_owned_story_or_404(story_id, user_doc)
    if not story:
        return jsonify({"ok": False, "error": "Story not found."}), 404
    if not _can_edit_story_content(story, user_doc):
        return jsonify({"ok": False, "error": "You can comment on this story, but cannot edit its attached work."}), 403
    expected_status = str(story.get("status") or "Assigned").strip()

    attachment_count, stored_bytes = _story_attachment_usage(story)
    if attachment_count >= MAX_STORY_ATTACHMENTS:
        return jsonify({"ok": False, "error": f"A story can have at most {MAX_STORY_ATTACHMENTS} attachments."}), 409

    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        return jsonify({"ok": False, "error": "Choose a file to upload."}), 400

    data = uploaded.read(MAX_STORY_ATTACHMENT_BYTES + 1)
    if len(data) > MAX_STORY_ATTACHMENT_BYTES:
        return jsonify({"ok": False, "error": "Story files must be 10 MB or smaller."}), 413
    if not data:
        return jsonify({"ok": False, "error": "Choose a non-empty file."}), 400
    if stored_bytes + len(data) > MAX_STORY_STORAGE_BYTES:
        return jsonify({"ok": False, "error": "This story has reached its attachment storage limit."}), 413

    filename = (secure_filename(uploaded.filename) or "story-upload")[:180]
    if Path(filename).suffix.lower() in DANGEROUS_UPLOAD_SUFFIXES:
        return jsonify({"ok": False, "error": "This file type is not allowed for story attachments."}), 400
    content_type = re.sub(r"[\x00-\x1f\x7f]", "", str(uploaded.content_type or "application/octet-stream"))[:200]
    now = datetime.now(timezone.utc)
    attachment_id = _new_attachment_id("file")
    file_id = story_files.put(
        data,
        filename=filename,
        contentType=content_type,
        metadata={
            "workspaceId": _workspace_id_for_user(user_doc),
            "storyId": _doc_public_id(story, "storyId"),
            "attachmentId": attachment_id,
            "uploadedBy": str(user_doc.get("_id")),
            "uploadedByEmail": user_doc.get("email", ""),
        },
    )
    attachment = {
        "id": attachment_id,
        "type": "file",
        "fileId": str(file_id),
        "name": filename,
        "contentType": content_type,
        "size": len(data),
        "uploadedAt": now,
        "addedBy": _attachment_added_by(user_doc),
        "permissionStatus": "app_access",
    }
    update = {
        "updatedAt": _now_iso(),
        "attachmentType": "multiple",
        "attachmentName": filename,
        "googleDocUrl": "",
    }
    mutation_scope = _story_content_mutation_query(user_doc, expected_status)
    result = stories_col.update_one(
        {"$and": [{"_id": story["_id"]}, mutation_scope]},
        {"$set": {**update, "attachments": _merged_attachment_items(story, attachment)}, "$unset": {"docUrl": "", "driveAttachment": ""}},
    )
    if result.matched_count == 0:
        try:
            story_files.delete(file_id)
        except Exception:
            pass
        return jsonify({"ok": False, "error": "Story access changed. Refresh and try again."}), 409
    updated = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {}
    return jsonify({"ok": True, "story": _story_to_api(updated)})


def _find_attachment_item(story: dict, attachment_id: str) -> dict | None:
    target = str(attachment_id or "").strip()
    if not target:
        return None
    items = story.get("attachments") if isinstance(story.get("attachments"), list) else []
    for item in items:
        if isinstance(item, dict) and target in _attachment_identifiers(item):
            return item
    for item in _legacy_story_attachment_items(story):
        if target in _attachment_identifiers(item):
            return item
    return None


def _gridfs_file_matches_story(stored, story: dict, attachment_id: str = "") -> bool:
    metadata = getattr(stored, "metadata", None)
    if not isinstance(metadata, dict):
        return False
    expected_story_id = _doc_public_id(story, "storyId")
    if str(metadata.get("storyId") or "") != expected_story_id:
        return False
    expected_workspace_id = str(story.get("workspaceId") or "").strip()
    stored_workspace_id = str(metadata.get("workspaceId") or "").strip()
    if stored_workspace_id and stored_workspace_id != expected_workspace_id:
        return False
    stored_attachment_id = str(metadata.get("attachmentId") or "").strip()
    if attachment_id and stored_attachment_id and stored_attachment_id != attachment_id:
        return False
    return True


@app.get("/api/stories/<story_id>/attachments/<attachment_id>")
@require_auth
def api_download_story_attachment_by_id(story_id: str, attachment_id: str):
    user_doc = _current_user_doc()
    story = _find_owned_story_or_404(story_id, user_doc)
    if not story:
        return jsonify({"ok": False, "error": "Story not found."}), 404
    attachment = _find_attachment_item(story, attachment_id)
    if not attachment or attachment.get("type") != "file":
        return jsonify({"ok": False, "error": "Story file not found."}), 404
    file_id = _object_id_or_none(str(attachment.get("fileId") or attachment.get("attachmentFileId") or ""))
    if not file_id:
        return jsonify({"ok": False, "error": "Story file not found."}), 404
    try:
        stored = story_files.get(file_id)
    except Exception:
        return jsonify({"ok": False, "error": "Story file not found."}), 404
    if not _gridfs_file_matches_story(stored, story, str(attachment.get("id") or attachment_id)):
        return jsonify({"ok": False, "error": "Story file not found."}), 404
    if int(getattr(stored, "length", 0) or 0) > MAX_STORY_ATTACHMENT_BYTES:
        return jsonify({"ok": False, "error": "Stored story file exceeds the download limit."}), 413

    filename = (secure_filename(str(attachment.get("name") or getattr(stored, "filename", "") or "story-file")) or "story-file")[:180]
    content_type = re.sub(r"[\x00-\x1f\x7f]", "", str(attachment.get("contentType") or getattr(stored, "content_type", "") or "application/octet-stream"))[:200]
    return send_file(
        BytesIO(stored.read()),
        mimetype=content_type,
        as_attachment=True,
        download_name=filename,
    )


@app.delete("/api/stories/<story_id>/attachments/<attachment_id>")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_delete_story_attachment(story_id: str, attachment_id: str):
    user_doc = _current_user_doc()
    story = _find_owned_story_or_404(story_id, user_doc)
    if not story:
        return jsonify({"ok": False, "error": "Story not found."}), 404
    if not _can_edit_story_content(story, user_doc):
        return jsonify({"ok": False, "error": "You can comment on this story, but cannot edit its attached work."}), 403
    expected_status = str(story.get("status") or "Assigned").strip()
    attachment = _find_attachment_item(story, attachment_id)
    if not attachment:
        return jsonify({"ok": False, "error": "Attachment not found."}), 404
    file_id_to_delete = _object_id_or_none(str(attachment.get("fileId") or attachment.get("attachmentFileId") or "")) if attachment.get("type") == "file" else None
    unset = {}
    set_values = {"updatedAt": _now_iso()}
    if attachment.get("type") == "file" and attachment_id == "legacy-file":
        unset.update({
            "attachmentFileId": "",
            "attachmentName": "",
            "attachmentContentType": "",
            "attachmentSize": "",
            "attachmentUploadedAt": "",
        })
    if attachment.get("type") == "link":
        set_values["googleDocUrl"] = ""
        unset["docUrl"] = ""
    if attachment.get("type") == "drive":
        attachment_url = str(attachment.get("webViewLink") or attachment.get("url") or "")
        if str(story.get("googleDocUrl") or "") == attachment_url:
            set_values["googleDocUrl"] = ""
        unset["driveAttachment"] = ""
    operation = {
        "$pull": {"attachments": _attachment_pull_condition(attachment_id)},
        "$set": set_values,
    }
    if unset:
        operation["$unset"] = unset
    mutation_scope = _story_content_mutation_query(user_doc, expected_status)
    result = stories_col.update_one(
        {"$and": [{"_id": story["_id"]}, mutation_scope]},
        operation,
    )
    if result.matched_count == 0:
        return jsonify({"ok": False, "error": "Story access changed. Refresh and try again."}), 409
    if file_id_to_delete:
        try:
            stored = story_files.get(file_id_to_delete)
            if _gridfs_file_matches_story(stored, story, str(attachment.get("id") or attachment_id)):
                story_files.delete(file_id_to_delete)
        except Exception:
            pass
    updated = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {}
    return jsonify({"ok": True, "story": _story_to_api(updated)})


@app.get("/api/stories/<story_id>/attachment")
@require_auth
def api_download_story_attachment(story_id: str):
    user_doc = _current_user_doc()
    story = _find_owned_story_or_404(story_id, user_doc)
    if not story:
        return jsonify({"ok": False, "error": "Story not found."}), 404

    file_id = _object_id_or_none(str(story.get("attachmentFileId") or ""))
    if not file_id:
        return jsonify({"ok": False, "error": "Story file not found."}), 404
    try:
        stored = story_files.get(file_id)
    except Exception:
        return jsonify({"ok": False, "error": "Story file not found."}), 404
    if not _gridfs_file_matches_story(stored, story):
        return jsonify({"ok": False, "error": "Story file not found."}), 404
    if int(getattr(stored, "length", 0) or 0) > MAX_STORY_ATTACHMENT_BYTES:
        return jsonify({"ok": False, "error": "Stored story file exceeds the download limit."}), 413

    filename = (secure_filename(str(story.get("attachmentName") or getattr(stored, "filename", "") or "story-file")) or "story-file")[:180]
    content_type = re.sub(r"[\x00-\x1f\x7f]", "", str(story.get("attachmentContentType") or getattr(stored, "content_type", "") or "application/octet-stream"))[:200]
    return send_file(
        BytesIO(stored.read()),
        mimetype=content_type,
        as_attachment=True,
        download_name=filename,
    )


@app.get("/api/pitches")
@require_auth
def api_pitches():
    user_doc = _current_user_doc()
    query = _pitch_query_for_user(user_doc)
    if query is None:
        return jsonify({"ok": False, "error": "Pitch board is not available to guests."}), 403
    docs = list(pitches_col.find(query).sort([("updatedAt", -1), ("_id", -1)]))
    return jsonify({"ok": True, "pitches": [_pitch_to_api(doc) for doc in docs]})


@app.post("/api/pitches")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_create_pitch():
    user_doc = _current_user_doc()
    payload = _request_payload()
    details, detail_error = _pitch_detail_updates(payload, {"title": "", "angle": "", "section": "", "notes": ""})
    if detail_error:
        return jsonify({"ok": False, "error": detail_error}), 400
    now_iso = _now_iso()
    doc = {
        "workspaceId": _workspace_id_for_user(user_doc),
        "title": details["title"],
        "angle": details["angle"],
        "status": "In Progress",
        "section": details["section"],
        "owner": _user_display_name(user_doc),
        "ownerEmail": normalize_email(user_doc.get("email") or ""),
        "ownerUserId": str(user_doc.get("_id")),
        "submittedAt": now_iso,
        "notes": details.get("notes", ""),
        "feedback": [],
        "comments": [],
        "createdAt": now_iso,
        "updatedAt": now_iso,
    }
    res = pitches_col.insert_one(doc)
    doc["_id"] = res.inserted_id
    return jsonify({"ok": True, "pitch": _pitch_to_api(doc)})


@app.patch("/api/pitches/<pitch_id>")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_update_pitch(pitch_id: str):
    user_doc = _current_user_doc()
    pitch = _find_owned_pitch_or_404(pitch_id, user_doc)
    if not pitch:
        return jsonify({"ok": False, "error": "Pitch not found."}), 404

    payload = _request_payload()
    role = _current_user_role(user_doc)
    current_status = _normalize_pitch_status(pitch.get("status"))
    raw_current_status = str(pitch.get("status") or "In Progress")
    detail_fields = {"title", "angle", "section", "notes"}
    if any(field in payload for field in detail_fields):
        if "status" in payload or any(field in payload for field in DUE_DATE_FIELDS):
            return jsonify({"ok": False, "error": "Update pitch details separately from workflow status."}), 400
        if current_status != "In Progress":
            return jsonify({"ok": False, "error": "Only in-progress pitches can be edited."}), 409
        if role == ROLE_WRITER and not _pitch_owned_by_user(pitch, user_doc):
            return jsonify({"ok": False, "error": "Only the pitch owner can edit this pitch."}), 403
        detail_update, detail_error = _pitch_detail_updates(payload, pitch)
        if detail_error:
            return jsonify({"ok": False, "error": detail_error}), 400
        mutation_scope = _workspace_query(user_doc)
        if role == ROLE_WRITER:
            mutation_scope = _scoped_query(user_doc, _owned_pitch_query(user_doc))
        result = pitches_col.update_one(
            {"$and": [{"_id": pitch["_id"]}, mutation_scope, {"status": raw_current_status}]},
            {"$set": {**detail_update, "updatedAt": _now_iso()}},
        )
        if result.matched_count == 0:
            return jsonify({"ok": False, "error": "Pitch access or status changed. Refresh and try again."}), 409
        activity_col.insert_one({
            "workspaceId": _workspace_id_for_user(user_doc),
            "entityType": "pitch",
            "entityId": _doc_public_id(pitch, "pitchId"),
            "eventType": "pitch_edited",
            "text": "Updated pitch details.",
            "actorId": str(user_doc.get("_id")),
            "actorEmail": user_doc.get("email", ""),
            "actorName": _user_display_name(user_doc),
            "createdAt": datetime.now(timezone.utc),
        })
        updated = pitches_col.find_one({"_id": pitch["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {}
        return jsonify({"ok": True, "pitch": _pitch_to_api(updated)})

    approval_deadline = _first_text_value(*(payload.get(field) for field in DUE_DATE_FIELDS))
    approval_message = str(payload.get("approvalMessage") or payload.get("message") or "").strip()
    invite_emails = _approval_invite_emails(payload.get("inviteEmails") or payload.get("invites") or payload.get("emails"))
    if len(approval_deadline) > 80:
        return jsonify({"ok": False, "error": "Due date is too long."}), 400
    if len(approval_message) > 200:
        return jsonify({"ok": False, "error": "Approval message must be 200 characters or fewer."}), 400
    update = {"updatedAt": _now_iso()}
    next_status = ""
    if "status" in payload:
        next_status = _normalize_pitch_status(payload.get("status"))
        if next_status not in {"In Progress", "Ready for Review", "Approved", "On Hold"}:
            return jsonify({"ok": False, "error": "Choose a valid pitch status."}), 400
        if role == ROLE_WRITER:
            if not _pitch_owned_by_user(pitch, user_doc):
                return jsonify({"ok": False, "error": "Only the pitch owner can submit it for review."}), 403
            if next_status != "Ready for Review":
                return jsonify({"ok": False, "error": "Writers can only submit their own pitch for review."}), 403
            if current_status != "In Progress":
                return jsonify({"ok": False, "error": "This pitch can no longer be submitted for review."}), 409
        if role in {ROLE_ADMIN, ROLE_EDITOR} and next_status == "Approved" and current_status != "Ready for Review":
            return jsonify({"ok": False, "error": "Only pitches ready for review can be approved."}), 409
        update["status"] = next_status
    if next_status == "Approved":
        if not approval_deadline:
            return jsonify({"ok": False, "error": "Due date is required before approving a pitch."}), 400
        update["deadline"] = approval_deadline
        update["dueDate"] = approval_deadline
    if len(update) == 1:
        return jsonify({"ok": False, "error": "No editable fields provided."}), 400

    mutation_scope = _workspace_query(user_doc)
    if role == ROLE_WRITER:
        mutation_scope = _scoped_query(user_doc, _owned_pitch_query(user_doc))
    result = pitches_col.update_one(
        {"$and": [{"_id": pitch["_id"]}, mutation_scope, {"status": raw_current_status}]},
        {"$set": update},
    )
    if result.matched_count == 0:
        return jsonify({"ok": False, "error": "Pitch access or status changed. Refresh and try again."}), 409
    updated = pitches_col.find_one({"_id": pitch["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {**pitch, **update}
    payload = {"ok": True, "pitch": _pitch_to_api(updated)}
    if update.get("status") == "Approved":
        try:
            story_doc = _create_story_from_pitch(updated, user_doc, approval_deadline, approval_message)
        except Exception:
            rollback_set = {"status": raw_current_status, "updatedAt": _now_iso()}
            rollback_operation = {"$set": rollback_set}
            original_deadline = _due_date_from_doc(pitch)
            if original_deadline:
                rollback_set["deadline"] = original_deadline
                rollback_set["dueDate"] = original_deadline
            else:
                rollback_operation["$unset"] = {"deadline": "", "dueDate": ""}
            rollback = pitches_col.update_one(
                {
                    "_id": pitch["_id"],
                    "workspaceId": _workspace_id_for_user(user_doc),
                    "status": "Approved",
                    "updatedAt": update["updatedAt"],
                },
                rollback_operation,
            )
            current = pitches_col.find_one({"_id": pitch["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or updated
            error = "Pitch approval could not create its story. Refresh and try again."
            return jsonify({"ok": False, "error": error, "pitch": _pitch_to_api(current)}), (503 if rollback.matched_count else 500)
        invite_warning = _add_approval_collaborators(story_doc, invite_emails, approval_message, user_doc)
        story_doc = stories_col.find_one({"_id": story_doc["_id"]}) or story_doc
        payload["story"] = _story_to_api(story_doc)
        if invite_warning:
            payload["warning"] = invite_warning
    if "status" in update and update["status"] != pitch.get("status"):
        _record_status_activity("pitch", _doc_public_id(pitch, "pitchId"), pitch.get("status", ""), update["status"], user_doc)
    return jsonify(payload)


@app.delete("/api/pitches/<pitch_id>")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_delete_pitch(pitch_id: str):
    user_doc = _current_user_doc()
    pitch = _find_owned_pitch_or_404(pitch_id, user_doc)
    if not pitch:
        return jsonify({"ok": False, "error": "Pitch not found."}), 404
    role = _current_user_role(user_doc)
    current_status = _normalize_pitch_status(pitch.get("status"))
    raw_current_status = str(pitch.get("status") or "In Progress")
    if role == ROLE_WRITER:
        if not _pitch_owned_by_user(pitch, user_doc):
            return jsonify({"ok": False, "error": "Only the pitch owner can delete this pitch."}), 403
        if current_status != "In Progress":
            return jsonify({"ok": False, "error": "Writers can only delete their in-progress pitches."}), 409
    if role in {ROLE_ADMIN, ROLE_EDITOR} and current_status == "Approved":
        return jsonify({"ok": False, "error": "Approved pitches cannot be deleted because their story must retain its source pitch."}), 409

    mutation_scope = _workspace_query(user_doc)
    if role == ROLE_WRITER:
        mutation_scope = _scoped_query(user_doc, _owned_pitch_query(user_doc))
    result = pitches_col.delete_one({
        "$and": [{"_id": pitch["_id"]}, mutation_scope, {"status": raw_current_status}],
    })
    if result.deleted_count == 0:
        return jsonify({"ok": False, "error": "Pitch access or status changed. Refresh and try again."}), 409
    activity_col.insert_one({
        "workspaceId": _workspace_id_for_user(user_doc),
        "entityType": "pitch",
        "entityId": _doc_public_id(pitch, "pitchId"),
        "entityTitle": pitch.get("title") or "Untitled pitch",
        "eventType": "pitch_deleted",
        "text": f"Deleted pitch: {pitch.get('title') or 'Untitled pitch'}.",
        "actorId": str(user_doc.get("_id")),
        "actorEmail": user_doc.get("email", ""),
        "actorName": _user_display_name(user_doc),
        "createdAt": datetime.now(timezone.utc),
    })
    return jsonify({"ok": True})


@app.get("/api/feedback")
@require_auth
def api_feedback():
    user_doc = _current_user_doc()
    entity_type = str(request.args.get("entityType") or "").strip().lower()
    entity_id = str(request.args.get("entityId") or "").strip()
    _, error, status = _require_feedback_entity(entity_type, entity_id, user_doc)
    if error:
        return jsonify({"ok": False, "error": error}), status
    docs = list(
        feedback_col.find({
            "workspaceId": _workspace_id_for_user(user_doc),
            "entityType": entity_type,
            "entityId": entity_id,
        }).sort([("createdAt", -1), ("_id", -1)]).limit(100)
    )
    return jsonify({"ok": True, "feedback": [_feedback_to_api(doc) for doc in docs]})


@app.post("/api/feedback")
@require_roles(ROLE_ADMIN, ROLE_EDITOR, ROLE_WRITER)
def api_create_feedback():
    user_doc = _current_user_doc()
    payload = _request_payload()
    entity_type = str(payload.get("entityType") or "").strip().lower()
    entity_id = str(payload.get("entityId") or "").strip()
    _, error, status = _require_feedback_entity(entity_type, entity_id, user_doc)
    if error:
        return jsonify({"ok": False, "error": error}), status
    text = str(payload.get("text") or "").strip()
    if not text:
        return jsonify({"ok": False, "error": "Feedback text is required."}), 400
    if len(text) > 4000:
        return jsonify({"ok": False, "error": "Feedback must be 4000 characters or fewer."}), 400

    now = datetime.now(timezone.utc)
    doc = {
        "workspaceId": _workspace_id_for_user(user_doc),
        "entityType": entity_type,
        "entityId": entity_id,
        "text": text,
        "authorId": str(user_doc.get("_id")),
        "authorEmail": user_doc.get("email", ""),
        "authorName": _user_display_name(user_doc),
        "createdAt": now,
        "updatedAt": now,
    }
    res = feedback_col.insert_one(doc)
    doc["_id"] = res.inserted_id
    return jsonify({"ok": True, "feedback": _feedback_to_api(doc)})


@app.patch("/api/feedback/<feedback_id>")
@require_roles(ROLE_ADMIN, ROLE_EDITOR)
def api_update_feedback(feedback_id: str):
    user_doc = _current_user_doc()
    oid = _object_id_or_none(feedback_id)
    if not oid:
        return jsonify({"ok": False, "error": "Invalid feedback id."}), 400
    doc = feedback_col.find_one({"_id": oid, "workspaceId": _workspace_id_for_user(user_doc)})
    if not doc:
        return jsonify({"ok": False, "error": "Feedback not found."}), 404
    _, error, status = _require_feedback_entity(doc.get("entityType"), doc.get("entityId"), user_doc)
    if error:
        return jsonify({"ok": False, "error": error}), status
    text = str(_request_payload().get("text") or "").strip()
    if not text:
        return jsonify({"ok": False, "error": "Feedback text is required."}), 400
    if len(text) > 4000:
        return jsonify({"ok": False, "error": "Feedback must be 4000 characters or fewer."}), 400
    feedback_col.update_one({"_id": oid}, {"$set": {"text": text, "updatedAt": datetime.now(timezone.utc)}})
    updated = feedback_col.find_one({"_id": oid, "workspaceId": _workspace_id_for_user(user_doc)}) or {}
    return jsonify({"ok": True, "feedback": _feedback_to_api(updated)})


@app.delete("/api/feedback/<feedback_id>")
@require_roles(ROLE_ADMIN, ROLE_EDITOR)
def api_delete_feedback(feedback_id: str):
    user_doc = _current_user_doc()
    oid = _object_id_or_none(feedback_id)
    if not oid:
        return jsonify({"ok": False, "error": "Invalid feedback id."}), 400
    doc = feedback_col.find_one({"_id": oid, "workspaceId": _workspace_id_for_user(user_doc)})
    if not doc:
        return jsonify({"ok": False, "error": "Feedback not found."}), 404
    _, error, status = _require_feedback_entity(doc.get("entityType"), doc.get("entityId"), user_doc)
    if error:
        return jsonify({"ok": False, "error": error}), status
    feedback_col.delete_one({"_id": oid})
    return jsonify({"ok": True})


@app.get("/api/activity")
@require_auth
def api_activity():
    user_doc = _current_user_doc()
    entity_type = str(request.args.get("entityType") or "").strip().lower()
    entity_id = str(request.args.get("entityId") or "").strip()
    if entity_type not in {"story", "pitch"} or not entity_id:
        return jsonify({"ok": False, "error": "Activity requires story or pitch entity information."}), 400

    if entity_type == "story":
        permitted_doc = _find_owned_story_or_404(entity_id, user_doc)
        denied_message = "Story activity is not available."
    else:
        permitted_doc = _find_owned_pitch_or_404(entity_id, user_doc)
        denied_message = "Pitch activity is not available."
    if not permitted_doc:
        return jsonify({"ok": False, "error": denied_message}), 403

    docs = list(
        activity_col.find({
            "workspaceId": _workspace_id_for_user(user_doc),
            "entityType": entity_type,
            "entityId": entity_id,
            "eventType": {"$in": sorted(IMPORTANT_ACTIVITY_EVENTS)},
        }).sort([("createdAt", -1), ("_id", -1)]).limit(50)
    )
    return jsonify({"ok": True, "activity": [_activity_to_api(doc) for doc in docs]})


def _pending_story_invitations(user_doc: dict) -> list[tuple[dict, dict]]:
    email = normalize_email(user_doc.get("email") or "")
    user_id = str(user_doc.get("_id") or "")
    if not email and not user_id:
        return []
    query = {"workspaceId": _workspace_id_for_user(user_doc), "collaborators": {"$elemMatch": {
        "status": "pending",
        "$or": [{"email": email}, {"userId": user_id}],
    }}}
    invitations = []
    for story in stories_col.find(query).sort([("updatedAt", -1), ("_id", -1)]):
        for invite in _story_collaborators(story):
            if invite.get("status") != "pending":
                continue
            if normalize_email(invite.get("email") or "") == email or str(invite.get("userId") or "") == user_id:
                invitations.append((story, invite))
                break
    return invitations


def _dashboard_entity_maps(user_doc: dict, personal_only: bool = False) -> tuple[dict[str, str], dict[str, str]]:
    if personal_only:
        story_query = _scoped_query(user_doc, {"$or": _owned_story_query(user_doc)["$or"] + [_collaborator_story_query(user_doc)]})
        pitch_query = _scoped_query(user_doc, _owned_pitch_query(user_doc))
    else:
        story_query = _story_query_for_user(user_doc)
        pitch_query = _pitch_query_for_user(user_doc)
    story_titles = {}
    pitch_titles = {}
    if story_query is not None:
        for doc in stories_col.find(story_query, {"_id": 1, "id": 1, "storyId": 1, "title": 1, "storyTitle": 1}):
            story_titles[_doc_public_id(doc, "storyId")] = doc.get("title") or doc.get("storyTitle") or "Untitled story"
    if pitch_query is not None:
        for doc in pitches_col.find(pitch_query, {"_id": 1, "id": 1, "pitchId": 1, "title": 1}):
            pitch_titles[_doc_public_id(doc, "pitchId")] = doc.get("title") or "Untitled pitch"
    return story_titles, pitch_titles

def _dashboard_event_datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        parsed = None
        for fmt in ("%B %d, %Y", "%b %d, %Y", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(raw, fmt)
                break
            except ValueError:
                continue
    if parsed is None:
        return None
    return parsed.replace(tzinfo=timezone.utc) if parsed.tzinfo is None else parsed.astimezone(timezone.utc)


def _activity_actor_is_user(doc: dict, user_doc: dict) -> bool:
    actor_id = str(doc.get("actorId") or "").lower()
    actor_email = normalize_email(doc.get("actorEmail") or "")
    user_ids = {str(value).lower() for value in _owner_match_values(user_doc) if value}
    return bool((actor_id and actor_id in user_ids) or (actor_email and actor_email in user_ids))


def _dashboard_status_activity_relevant(doc: dict, role: str, personal_story_ids: set[str], personal_pitch_ids: set[str]) -> bool:
    if doc.get("eventType") != "status_change":
        return False
    entity_type = str(doc.get("entityType") or "")
    entity_id = str(doc.get("entityId") or "")
    to_status = str(doc.get("toStatus") or "")
    if entity_type == "pitch":
        to_status = _normalize_pitch_status(to_status)
        if entity_id in personal_pitch_ids:
            return to_status in {"In Progress", "Approved", "On Hold"}
        return role in {ROLE_EDITOR, ROLE_ADMIN} and to_status == "Ready for Review"
    if entity_type == "story":
        if entity_id in personal_story_ids:
            return to_status in {"Assigned", "Needs Revision", "Returned", "Ready for Publish", "Published"}
        return role in {ROLE_EDITOR, ROLE_ADMIN} and to_status == "Submitted"
    return False


def _dashboard_status_activity_text(doc: dict) -> str:
    actor = str(doc.get("actorName") or "A newsroom teammate")
    entity_type = str(doc.get("entityType") or "")
    to_status = str(doc.get("toStatus") or "")
    if entity_type == "pitch":
        to_status = _normalize_pitch_status(to_status)
        messages = {
            "Ready for Review": f"{actor} submitted this pitch for review.",
            "In Progress": f"{actor} requested more work before this pitch can move forward.",
            "Approved": f"{actor} approved this pitch and moved it to Stories.",
            "On Hold": f"{actor} placed this pitch on hold.",
        }
        return messages.get(to_status, str(doc.get("text") or "Pitch updated."))
    messages = {
        "Submitted": f"{actor} submitted this story for review.",
        "Assigned": f"{actor} assigned this story.",
        "Needs Revision": f"{actor} requested story revisions.",
        "Returned": f"{actor} returned this story for revisions.",
        "Ready for Publish": f"{actor} moved this story to teacher approval.",
        "Published": f"{actor} marked this story published.",
    }
    return messages.get(to_status, str(doc.get("text") or "Story updated."))

@app.get("/api/dashboard")
@require_auth
def api_dashboard():
    user_doc = _current_user_doc()
    role = _current_user_role(user_doc)
    user_email = normalize_email(user_doc.get("email") or "")
    tasks = []

    pending_invitations = _pending_story_invitations(user_doc) if role == ROLE_WRITER else []
    for story, invite in pending_invitations:
        story_id = _doc_public_id(story, "storyId")
        tasks.append({
            "id": f"invite:{story_id}", "kind": "invitation", "title": story.get("title") or story.get("storyTitle") or "Untitled story",
            "detail": f"{invite.get('invitedBy') or 'A newsroom teammate'} invited you to join as a co-author.",
            "time": invite.get("invitedAt") or "", "priority": "high", "entityType": "story", "entityId": story_id,
            "actions": ["accept", "decline"],
        })

    if role in {ROLE_EDITOR, ROLE_ADMIN}:
        for pitch in pitches_col.find(_scoped_query(user_doc, {"status": {"$in": ["Ready for Review", "Submitted", "Needs Review"]}})).sort([("updatedAt", -1), ("_id", -1)]).limit(20):
            if _pitch_owned_by_user(pitch, user_doc):
                continue
            tasks.append({"id": f"pitch:{_doc_public_id(pitch, 'pitchId')}", "kind": "pitch_review", "title": pitch.get("title") or "Untitled pitch", "detail": f"Pitch from {pitch.get('owner') or pitch.get('writer') or pitch.get('ownerEmail') or 'a writer'} needs review.", "time": _date_for_api(pitch.get("updatedAt") or pitch.get("submittedAt") or pitch.get("createdAt")), "priority": "normal", "entityType": "pitch", "entityId": _doc_public_id(pitch, "pitchId")})
        for story in stories_col.find(_scoped_query(user_doc, {"status": {"$in": ["Submitted", "In Review"]}})).sort([("updatedAt", -1), ("_id", -1)]).limit(20):
            if _story_primary_owned_by_user(story, user_doc):
                continue
            tasks.append({"id": f"story:{_doc_public_id(story, 'storyId')}", "kind": "story_review", "title": story.get("title") or story.get("storyTitle") or "Untitled story", "detail": f"{story.get('writer') or story.get('writerEmail') or 'A writer'} submitted this story for review.", "time": _date_for_api(story.get("submittedAt") or story.get("updatedAt")), "priority": "high" if story.get("dueSoon") else "normal", "entityType": "story", "entityId": _doc_public_id(story, "storyId")})
    elif role == ROLE_WRITER:
        for pitch in pitches_col.find(_scoped_query(user_doc, {"$and": [_owned_pitch_query(user_doc), {"status": {"$in": ["In Progress", "New"]}}]})).sort([("updatedAt", -1), ("_id", -1)]).limit(20):
            tasks.append({"id": f"pitch:{_doc_public_id(pitch, 'pitchId')}", "kind": "pitch_work", "title": pitch.get("title") or "Untitled pitch", "detail": "Continue refining this pitch, then submit it when it is ready for editor review.", "time": _date_for_api(pitch.get("updatedAt") or pitch.get("createdAt")), "priority": "normal", "entityType": "pitch", "entityId": _doc_public_id(pitch, "pitchId")})
        writer_story_query = _story_query_for_user(user_doc)
        for story in stories_col.find(writer_story_query).sort([("updatedAt", -1), ("_id", -1)]).limit(30):
            status = str(story.get("status") or "Assigned")
            if status in {"Published", "Ready for Publish", "Submitted", "In Review"}:
                continue
            detail = "Revisions were requested. Review the editor notes and resubmit." if status in {"Returned", "Needs Revision"} else "Continue reporting and attach your work before the deadline."
            tasks.append({"id": f"story:{_doc_public_id(story, 'storyId')}", "kind": "revision" if status in {"Returned", "Needs Revision"} else "story_work", "title": story.get("title") or story.get("storyTitle") or "Untitled story", "detail": detail, "time": _date_for_api(story.get("updatedAt") or story.get("createdAt")), "dueDate": _story_deadline(story), "priority": "high" if status in {"Returned", "Needs Revision"} or story.get("dueSoon") else "normal", "entityType": "story", "entityId": _doc_public_id(story, "storyId")})

    story_titles, pitch_titles = _dashboard_entity_maps(user_doc)
    personal_story_titles, personal_pitch_titles = _dashboard_entity_maps(user_doc, personal_only=True)
    personal_story_ids = set(personal_story_titles)
    personal_pitch_ids = set(personal_pitch_titles)
    entity_clauses = []
    if story_titles:
        entity_clauses.append({"entityType": "story", "entityId": {"$in": list(story_titles)}})
    if pitch_titles:
        entity_clauses.append({"entityType": "pitch", "entityId": {"$in": list(pitch_titles)}})
    personal_entity_clauses = []
    if personal_story_titles:
        personal_entity_clauses.append({"entityType": "story", "entityId": {"$in": list(personal_story_titles)}})
    if personal_pitch_titles:
        personal_entity_clauses.append({"entityType": "pitch", "entityId": {"$in": list(personal_pitch_titles)}})

    if role in {ROLE_EDITOR, ROLE_ADMIN}:
        activity_query = {"entityType": {"$in": ["story", "pitch"]}}
    elif entity_clauses:
        activity_query = {"$or": entity_clauses}
    else:
        activity_query = {"_id": None}

    activities = []
    activity_cutoff = datetime.now(timezone.utc) - timedelta(days=7)
    for doc in activity_col.find(_scoped_query(user_doc, activity_query)).sort([("createdAt", -1), ("_id", -1)]).limit(100):
        event_datetime = _dashboard_event_datetime(doc.get("createdAt"))
        if event_datetime is None or event_datetime < activity_cutoff:
            continue
        if _activity_actor_is_user(doc, user_doc):
            continue
        if not _dashboard_status_activity_relevant(doc, role, personal_story_ids, personal_pitch_ids):
            continue
        item = _activity_to_api(doc)
        item["text"] = _dashboard_status_activity_text(doc)
        item["title"] = doc.get("entityTitle") or (story_titles if item["entityType"] == "story" else pitch_titles).get(item["entityId"], "")
        item["kind"] = item.get("eventType") or "update"
        item["_sortAt"] = event_datetime.timestamp()
        activities.append(item)

    feedback_query = {"$or": personal_entity_clauses} if personal_entity_clauses else {"_id": None}
    for doc in feedback_col.find(_scoped_query(user_doc, feedback_query)).sort([("createdAt", -1), ("_id", -1)]).limit(30):
        if normalize_email(doc.get("authorEmail") or "") == user_email:
            continue
        event_datetime = _dashboard_event_datetime(doc.get("createdAt"))
        if event_datetime is None or event_datetime < activity_cutoff:
            continue
        entity_type = doc.get("entityType") or ""
        entity_id = doc.get("entityId") or ""
        activities.append({"id": f"feedback:{doc.get('_id')}", "entityType": entity_type, "entityId": entity_id, "eventType": "comment", "kind": "comment", "text": f"{doc.get('authorName') or doc.get('actorName') or 'A teammate'} commented: {doc.get('text') or ''}", "title": (story_titles if entity_type == "story" else pitch_titles).get(entity_id, ""), "time": _date_for_api(doc.get("createdAt")), "occurredAt": _datetime_for_api(doc.get("createdAt")), "actorName": doc.get("authorName") or doc.get("actorName") or "", "_sortAt": event_datetime.timestamp()})

    for story, invite in pending_invitations:
        event_datetime = _dashboard_event_datetime(invite.get("invitedAtIso") or invite.get("invitedAt"))
        if event_datetime is None or event_datetime < activity_cutoff:
            continue
        story_id = _doc_public_id(story, "storyId")
        activities.append({"id": f"invite:{story_id}", "entityType": "story", "entityId": story_id, "eventType": "invitation", "kind": "invitation", "text": f"{invite.get('invitedBy') or 'A newsroom teammate'} invited you to co-author this story.", "title": story.get("title") or story.get("storyTitle") or "Untitled story", "time": invite.get("invitedAt") or "", "occurredAt": invite.get("invitedAtIso") or "", "actions": ["accept", "decline"], "_sortAt": event_datetime.timestamp()})

    activities.sort(key=lambda item: item.get("_sortAt", 0), reverse=True)
    for item in activities:
        item.pop("_sortAt", None)
    return jsonify({"ok": True, "tasks": tasks[:40], "activity": activities[:60]})


@app.post("/api/story-invitations/<story_id>/<decision>")
@require_roles(ROLE_WRITER)
def api_respond_story_invitation(story_id: str, decision: str):
    if decision not in {"accept", "decline"}:
        return jsonify({"ok": False, "error": "Choose accept or decline."}), 400
    user_doc = _current_user_doc()
    email = normalize_email(user_doc.get("email") or "")
    clauses = [{"_id": _object_id_or_none(story_id)}, {"id": story_id}, {"storyId": story_id}]
    clauses = [clause for clause in clauses if list(clause.values())[0]]
    story = stories_col.find_one({"$and": [_workspace_query(user_doc), {"$or": clauses}, {"collaborators": {"$elemMatch": {"email": email, "status": "pending"}}}]})
    if not story:
        return jsonify({"ok": False, "error": "This invitation is no longer available."}), 404
    invitation_query = {
        "_id": story["_id"],
        "workspaceId": _workspace_id_for_user(user_doc),
        "collaborators": {"$elemMatch": {"email": email, "status": "pending"}},
    }
    if decision == "accept":
        result = stories_col.update_one(
            invitation_query,
            {"$set": {
                "collaborators.$[invite].status": "accepted",
                "collaborators.$[invite].role": "edit",
                "collaborators.$[invite].respondedAt": _now_iso(),
                "updatedAt": _now_iso(),
            }},
            array_filters=[{"invite.email": email, "invite.status": "pending"}],
        )
    else:
        result = stories_col.update_one(
            invitation_query,
            {"$pull": {"collaborators": {"email": email, "status": "pending"}}, "$set": {"updatedAt": _now_iso()}},
        )
    if result.modified_count == 0:
        return jsonify({"ok": False, "error": "This invitation is no longer available."}), 409
    activity_col.insert_one({"workspaceId": _workspace_id_for_user(user_doc), "entityType": "story", "entityId": _doc_public_id(story, "storyId"), "entityTitle": story.get("title") or story.get("storyTitle") or "Untitled story", "eventType": f"invitation_{'accepted' if decision == 'accept' else 'declined'}", "text": f"{_user_display_name(user_doc)} {'joined the story as a co-author' if decision == 'accept' else 'declined the co-author invitation'}.", "actorId": str(user_doc.get("_id")), "actorEmail": email, "actorName": _user_display_name(user_doc), "createdAt": datetime.now(timezone.utc)})
    updated = stories_col.find_one({"_id": story["_id"], "workspaceId": _workspace_id_for_user(user_doc)}) or {}
    return jsonify({"ok": True, "decision": decision, "story": _story_to_api(updated) if decision == "accept" else None})


@app.get("/api/article-records")
@require_auth
def api_article_records():
    user_doc = _current_user_doc()
    try:
        page = max(1, int(request.args.get("page", "1") or "1"))
        limit = min(max(1, int(request.args.get("limit", "20") or "20")), 5000)
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Page and limit must be whole numbers."}), 400
    skip = (page - 1) * limit
    search = str(request.args.get("search") or "").strip()
    section = str(request.args.get("section") or "").strip()
    sort = str(request.args.get("sort") or "desc").strip().lower()
    sort_direction = ASCENDING if sort == "asc" else -1

    query = {}
    clauses = []
    if search:
        interview_urls = set()
        interview_query = _interviewee_search_query(_workspace_id_for_user(user_doc), search)
        for person in interviews_col.find(
            interview_query,
            {"url": 1, "articleUrl": 1, "article_url": 1},
        ):
            person_url = str(person.get("url") or person.get("articleUrl") or person.get("article_url") or "").strip()
            interview_urls.update(_article_url_lookup_candidates(person_url))
        clauses.append(_article_search_clause(search, interview_urls))
    if section and section != "All sections":
        clauses.append({"$or": [
            {"section": section},
            {"category": section},
            {"tags": section},
            {"categories": section},
        ]})
    if clauses:
        query = {"$and": clauses} if len(clauses) > 1 else clauses[0]

    query = _scoped_query(user_doc, query)
    total = articles_col.count_documents(query)
    docs = list(
        articles_col.find(query)
        .sort([("datePublishedSort", sort_direction), ("_id", sort_direction)])
        .skip(skip)
        .limit(limit)
    )

    urls = []
    for doc in docs:
        url = str(doc.get("url") or doc.get("articleUrl") or doc.get("article_url") or "").strip()
        if url:
            urls.extend(_article_url_lookup_candidates(url))

    people_by_url: dict[str, list[dict]] = {}
    if urls:
        cursor = interviews_col.find(_scoped_query(user_doc, {
            "$or": [
                {"url": {"$in": urls}},
                {"articleUrl": {"$in": urls}},
                {"article_url": {"$in": urls}},
            ]
        }))
        for person_doc in cursor:
            person = _interview_to_api(person_doc)
            people_by_url.setdefault(_article_join_url_key(person["url"]), []).append(person)

    section_values = {}

    def add_section_value(value):
        text = str(value or "").strip()
        if text and text.lower() != "all sections":
            section_values.setdefault(text.lower(), text)

    for doc in articles_col.find(_workspace_query(user_doc), {"section": 1, "category": 1, "tags": 1, "categories": 1}):
        for value in [doc.get("section"), doc.get("category")]:
            add_section_value(value)
        for field in ["tags", "categories"]:
            for value in _coerce_list_field(doc.get(field)):
                add_section_value(value)

    total_pages = max(1, (total + limit - 1) // limit)
    return jsonify({
        "ok": True,
        "articles": [_article_to_api(doc, people_by_url) for doc in docs],
        "page": page,
        "limit": limit,
        "total": total,
        "totalPages": total_pages,
        "sections": sorted(section_values.values(), key=lambda value: value.lower()),
    })


@app.get("/api/interview-records")
@require_auth
def api_interview_records():
    docs = list(interviews_col.find(_workspace_query(_current_user_doc())).sort([("_id", -1)]))
    return jsonify({"ok": True, "people": [_interview_to_api(doc) for doc in docs]})


@app.post("/api/interview-records/<record_id>")
@require_role_at_least(ROLE_EDITOR)
def api_update_interview_record(record_id: str):
    try:
        oid = ObjectId(record_id)
    except Exception:
        return jsonify({"ok": False, "error": "Invalid record id"}), 400

    user_doc = _current_user_doc()
    workspace_id = _workspace_id_for_user(user_doc)
    existing = interviews_col.find_one({"_id": oid, "workspaceId": workspace_id})
    if not existing:
        return jsonify({"ok": False, "error": "Record not found"}), 404

    payload = _request_payload()
    if not any(key in payload for key in {"firstName", "lastName", "grade", "house", "url"}):
        return jsonify({"ok": False, "error": "No editable fields provided"}), 400

    try:
        first_name = _validated_person_text(
            payload.get("firstName", existing.get("firstName") or existing.get("first_name") or ""),
            "First name",
            100,
            required=True,
        )
        last_name = _validated_person_text(
            payload.get("lastName", existing.get("lastName") or existing.get("last_name") or ""),
            "Last name",
            100,
            required=True,
        )
        grade = _validated_person_text(payload.get("grade", existing.get("grade") or ""), "Grade", 20)
        house = _validated_person_text(payload.get("house", existing.get("house") or ""), "House", 40)
        if grade not in EXTRACTION_GRADES:
            raise ValueError("Choose a valid grade.")
        if house not in EXTRACTION_HOUSES:
            raise ValueError("Choose a valid house.")
        raw_url = payload.get("url", existing.get("url") or existing.get("articleUrl") or "")
        canonical_url = canonicalize_article_url(raw_url)
        workspace = _workspace_for_user(user_doc)
        if not workspace or not _validated_publication_url(canonical_url, workspace):
            raise ValueError("The article URL must use the workspace publication domain.")
    except ArticleExtractionError as exc:
        return jsonify({"ok": False, **exc.to_dict()}), 400
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400

    update = {
        "firstName": first_name,
        "lastName": last_name,
        "name": f"{first_name} {last_name}",
        "firstNameNorm": first_name.casefold(),
        "lastNameNorm": last_name.casefold(),
        "grade": grade,
        "house": house,
        "url": canonical_url,
        "articleUrl": canonical_url,
        "urlNorm": canonical_url,
        "updatedAt": datetime.now(timezone.utc),
    }
    try:
        interviews_col.update_one({"_id": oid, "workspaceId": workspace_id}, {"$set": update})
    except DuplicateKeyError:
        return jsonify({"ok": False, "error": "That interviewee is already saved for this article."}), 409

    doc = interviews_col.find_one({"_id": oid, "workspaceId": workspace_id}) or {}
    return jsonify({"ok": True, "record": _interview_to_api(doc)})


@app.delete("/api/interview-records/<record_id>")
@require_role_at_least(ROLE_EDITOR)
def api_delete_interview_record(record_id: str):
    try:
        oid = ObjectId(record_id)
    except Exception:
        return jsonify({"ok": False, "error": "Invalid record id"}), 400

    result = interviews_col.delete_one({"_id": oid, "workspaceId": _workspace_id_for_user(_current_user_doc())})
    if result.deleted_count == 0:
        return jsonify({"ok": False, "error": "Record not found"}), 404
    return jsonify({"ok": True})


EXTRACTION_TOKEN_SALT = "falcon-newsroom-article-extraction-v1"
EXTRACTION_TOKEN_VERSION = 1
MAX_EXTRACTION_PEOPLE = 100
EXTRACTION_GRADES = {"", "9", "10", "11", "12", "Staff"}
EXTRACTION_HOUSES = {"", "SMCS", "Global", "Humanities", "ISP"}


def _workspace_extraction_host(workspace: dict | None) -> str:
    if not workspace:
        return ""
    configured_url = _safe_http_url(workspace.get("publicationUrl") or "")
    candidates = []
    if configured_url:
        candidates.append(urlparse(configured_url).hostname or "")
    candidates.append(workspace.get("articleDomain") or "")
    for candidate in candidates:
        try:
            return normalize_publication_host(str(candidate or ""))
        except ArticleExtractionError:
            continue
    return ""


def _consume_extraction_rate_limit(user_doc: dict, now: datetime | None = None) -> tuple[bool, int]:
    if EXTRACTION_RATE_LIMIT_MAX <= 0:
        return False, 0
    window_seconds = max(60, min(EXTRACTION_RATE_LIMIT_WINDOW_SECONDS, 86400))
    current = now or datetime.now(timezone.utc)
    if current.tzinfo is None:
        current = current.replace(tzinfo=timezone.utc)
    epoch = int(current.timestamp())
    window_epoch = epoch - (epoch % window_seconds)
    window_start = datetime.fromtimestamp(window_epoch, tz=timezone.utc)
    retry_after = max(1, window_seconds - (epoch - window_epoch))
    workspace_id = _workspace_id_for_user(user_doc)
    user_id = str(user_doc.get("_id") or "")
    try:
        bucket = extraction_rate_col.find_one_and_update(
            {
                "workspaceId": workspace_id,
                "userId": user_id,
                "windowStart": window_start,
            },
            {
                "$inc": {"count": 1},
                "$setOnInsert": {
                    "createdAt": current,
                    "expiresAt": window_start + timedelta(seconds=window_seconds * 2),
                },
            },
            upsert=True,
            return_document=ReturnDocument.AFTER,
        ) or {}
    except DuplicateKeyError:
        bucket = extraction_rate_col.find_one_and_update(
            {
                "workspaceId": workspace_id,
                "userId": user_id,
                "windowStart": window_start,
            },
            {"$inc": {"count": 1}},
            return_document=ReturnDocument.AFTER,
        ) or {}
    return int(bucket.get("count") or 0) > EXTRACTION_RATE_LIMIT_MAX, retry_after


def _extraction_serializer() -> URLSafeTimedSerializer:
    return URLSafeTimedSerializer(app.secret_key, salt=EXTRACTION_TOKEN_SALT)


def _metadata_text(value, max_length: int) -> str:
    text = str(value or "").strip()
    text = re.sub(r"[\x00-\x1f\x7f]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()[:max_length]


def _metadata_list(value, max_items: int, max_length: int) -> list[str]:
    values = value if isinstance(value, list) else _coerce_list_field(value)
    output = []
    seen = set()
    for item in values:
        text = _metadata_text(item, max_length)
        key = text.casefold()
        if text and key not in seen:
            seen.add(key)
            output.append(text)
        if len(output) >= max_items:
            break
    return output


def _sanitized_extracted_article(article, canonical_url: str) -> dict:
    source = article if isinstance(article, dict) else {}
    return {
        "url": canonical_url,
        "title": _metadata_text(source.get("title"), 300),
        "authors": _metadata_list(source.get("authors") or source.get("author"), 12, 100),
        "tags": _metadata_list(source.get("tags") or source.get("categories"), 16, 80),
        "datePublished": _metadata_text(
            source.get("datePublished") or source.get("publishedAt") or source.get("date"),
            120,
        ),
    }


def _safe_extraction_warnings(value) -> list[dict[str, str]]:
    items = value if isinstance(value, list) else []
    warnings = []
    for item in items[:10]:
        if not isinstance(item, dict):
            continue
        message = _metadata_text(item.get("message"), 500)
        if message:
            warnings.append({
                "code": _metadata_text(item.get("code"), 80),
                "message": message,
            })
    return warnings


def _enrich_extracted_people(people, workspace: dict) -> list[dict]:
    active_metadata = workspace.get("namesDatabase") if isinstance(workspace.get("namesDatabase"), dict) else {}
    active_batch_id = str(active_metadata.get("activeBatchId") or "").strip()
    workspace_id = str(workspace.get("publicId") or workspace.get("_id") or "").strip()
    output = []
    for item in (people if isinstance(people, list) else [])[:50]:
        if not isinstance(item, dict):
            continue
        first_name = _metadata_text(item.get("firstName") or item.get("first_name"), 100)
        last_name = _metadata_text(item.get("lastName") or item.get("last_name"), 100)
        if not first_name or not last_name:
            continue
        person = {"firstName": first_name, "lastName": last_name, "grade": "", "house": ""}
        if active_batch_id:
            exact_query = {
                "workspaceId": workspace_id,
                "uploadBatchId": active_batch_id,
                "firstName": {"$regex": f"^{re.escape(first_name)}$", "$options": "i"},
                "lastName": {"$regex": f"^{re.escape(last_name)}$", "$options": "i"},
            }
            matches = list(names_col.find(exact_query, {"grade": 1, "house": 1}).limit(2))
            if len(matches) == 1:
                grade = _metadata_text(matches[0].get("grade"), 20)
                house = _metadata_text(matches[0].get("house"), 40)
                person["grade"] = grade if grade in EXTRACTION_GRADES else ""
                person["house"] = house if house in EXTRACTION_HOUSES else ""
                person["directoryMatch"] = "exact"
            elif len(matches) > 1:
                person["directoryMatch"] = "ambiguous"
        output.append(person)
    return output


def _validated_person_text(value, field_label: str, max_length: int, required: bool = False) -> str:
    if not isinstance(value, str):
        if value is None and not required:
            return ""
        raise ValueError(f"{field_label} must be text.")
    text = value.strip()
    if required and not text:
        raise ValueError(f"{field_label} is required for every interviewee.")
    if len(text) > max_length:
        raise ValueError(f"{field_label} must be {max_length} characters or fewer.")
    if any(ord(char) < 32 or ord(char) == 127 for char in text):
        raise ValueError(f"{field_label} contains unsupported characters.")
    return text


def _validated_extraction_people(value, today=None) -> list[dict]:
    if not isinstance(value, list) or not value:
        raise ValueError("Add at least one interviewee before saving.")
    if len(value) > MAX_EXTRACTION_PEOPLE:
        raise ValueError(f"Save no more than {MAX_EXTRACTION_PEOPLE} interviewees at once.")
    current_date = today or datetime.now(timezone.utc).date()
    output = []
    seen = set()
    for item in value:
        if not isinstance(item, dict):
            raise ValueError("Each interviewee must be an object.")
        first_name = _validated_person_text(item.get("firstName"), "First name", 100, required=True)
        last_name = _validated_person_text(item.get("lastName"), "Last name", 100, required=True)
        grade = _validated_person_text(item.get("grade", ""), "Grade", 20)
        house = _validated_person_text(item.get("house", ""), "House", 40)
        if grade not in EXTRACTION_GRADES:
            raise ValueError("Choose a valid grade.")
        if house not in EXTRACTION_HOUSES:
            raise ValueError("Choose a valid house.")
        raw_date = _validated_person_text(item.get("dateAdded"), "Date added", 10, required=True)
        try:
            parsed_date = datetime.strptime(raw_date, "%Y-%m-%d").date()
        except ValueError:
            raise ValueError("Date added must use YYYY-MM-DD format.") from None
        if parsed_date.year < 2000 or parsed_date > current_date:
            raise ValueError("Date added must be between January 1, 2000 and today.")
        key = (first_name.casefold(), last_name.casefold())
        if key in seen:
            continue
        seen.add(key)
        output.append({
            "firstName": first_name,
            "lastName": last_name,
            "firstNameNorm": key[0],
            "lastNameNorm": key[1],
            "grade": grade,
            "house": house,
            "dateAdded": parsed_date.isoformat(),
        })
    if not output:
        raise ValueError("Add at least one unique interviewee before saving.")
    return output


def _extractor_error_status(code: str) -> int:
    if code == "response_too_large":
        return 413
    if code == "invalid_content_type":
        return 415
    if code == "fetch_timeout":
        return 504
    if code in {
        "dns_failed", "fetch_failed", "http_error", "peer_address_unavailable",
        "peer_address_mismatch", "too_many_redirects", "invalid_redirect", "https_downgrade",
    }:
        return 502
    if code in {"blocked_address", "publication_host_mismatch", "invalid_html"}:
        return 422
    return 400


def _extracted_article_date_sort(value: str):
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except ValueError:
        try:
            return datetime.strptime(raw[:10], "%Y-%m-%d").replace(tzinfo=timezone.utc)
        except ValueError:
            return None


def _upsert_extracted_article(workspace_id: str, canonical_url: str, article: dict, actor: dict) -> dict:
    now = datetime.now(timezone.utc)
    candidates = sorted(_article_url_lookup_candidates(canonical_url))
    lookup = {
        "workspaceId": workspace_id,
        "$or": [
            {"urlNorm": canonical_url},
            {"canonicalUrl": canonical_url},
            {"url": {"$in": candidates}},
            {"articleUrl": {"$in": candidates}},
            {"article_url": {"$in": candidates}},
        ],
    }
    existing = articles_col.find_one(lookup)
    if existing:
        fill_if_empty = {
            "title": article.get("title") or "",
            "articleTitle": article.get("title") or "",
            "authors": article.get("authors") or [],
            "author": ", ".join(article.get("authors") or []),
            "byline": ", ".join(article.get("authors") or []),
            "tags": article.get("tags") or [],
            "categories": article.get("tags") or [],
            "datePublished": article.get("datePublished") or "",
        }
        updates = {
            "url": existing.get("url") or canonical_url,
            "articleUrl": existing.get("articleUrl") or existing.get("url") or canonical_url,
            "canonicalUrl": canonical_url,
            "urlNorm": canonical_url,
            "updatedAt": now,
        }
        for field, value in fill_if_empty.items():
            if value and not existing.get(field):
                updates[field] = value
        date_sort = _extracted_article_date_sort(article.get("datePublished") or "")
        if date_sort and not existing.get("datePublishedSort"):
            updates["datePublishedSort"] = date_sort
        articles_col.update_one({"_id": existing["_id"], "workspaceId": workspace_id}, {"$set": updates})
        return articles_col.find_one({"_id": existing["_id"], "workspaceId": workspace_id}) or {**existing, **updates}

    authors = article.get("authors") or []
    tags = article.get("tags") or []
    insert_doc = {
        "workspaceId": workspace_id,
        "url": canonical_url,
        "articleUrl": canonical_url,
        "canonicalUrl": canonical_url,
        "urlNorm": canonical_url,
        "title": article.get("title") or "Untitled article",
        "articleTitle": article.get("title") or "Untitled article",
        "authors": authors,
        "author": ", ".join(authors),
        "byline": ", ".join(authors),
        "tags": tags,
        "categories": tags,
        "section": tags[0] if tags else "",
        "category": tags[0] if tags else "",
        "datePublished": article.get("datePublished") or "",
        "status": "Published",
        "addedBy": _user_display_name(actor),
        "addedByUserId": str(actor.get("_id") or ""),
        "createdAt": now,
        "updatedAt": now,
    }
    date_sort = _extracted_article_date_sort(article.get("datePublished") or "")
    if date_sort:
        insert_doc["datePublishedSort"] = date_sort
    try:
        result = articles_col.insert_one(insert_doc)
        insert_doc["_id"] = result.inserted_id
        return insert_doc
    except DuplicateKeyError:
        existing = articles_col.find_one({"workspaceId": workspace_id, "urlNorm": canonical_url})
        if not existing:
            raise
        return existing


@app.post("/api/extract")
@require_role_at_least(ROLE_EDITOR)
def api_extract():
    user_doc = _current_user_doc()
    workspace = _workspace_for_user(user_doc)
    publication_host = _workspace_extraction_host(workspace)
    if not workspace or not publication_host:
        return jsonify({"ok": False, "error": "Configure a valid workspace publication URL before extracting articles."}), 409

    limited, retry_after = _consume_extraction_rate_limit(user_doc)
    if limited:
        response = jsonify({"ok": False, "error": "Too many article extraction requests. Please wait and try again."})
        response.headers["Retry-After"] = str(retry_after)
        return response, 429

    payload = _request_payload()
    raw_url = payload.get("url") if isinstance(payload.get("url"), str) else ""
    try:
        canonical_input = canonicalize_article_url(raw_url)
        result = extract_article(
            canonical_input,
            publication_host,
            gemini_api_key=os.getenv("GEMINI_API_KEY", "").strip() or None,
        )
        canonical_url = canonicalize_article_url(result.get("article_url") or canonical_input)
        article = _sanitized_extracted_article(result.get("article"), canonical_url)
        people = _enrich_extracted_people(result.get("people"), workspace)
        warnings = _safe_extraction_warnings(result.get("warnings"))
        extraction = result.get("intervieweeExtraction") if isinstance(result.get("intervieweeExtraction"), dict) else {}
        extraction_status = _metadata_text(extraction.get("status"), 80) or "manual"
        extraction_info = {
            "attempted": bool(extraction.get("attempted")),
            "status": extraction_status,
            "mode": extraction_status,
            "model": _metadata_text(extraction.get("model"), 100),
            "people": people,
            "warnings": warnings,
        }
        token = _extraction_serializer().dumps({
            "v": EXTRACTION_TOKEN_VERSION,
            "workspaceId": _workspace_id_for_user(user_doc),
            "userId": str(user_doc.get("_id") or ""),
            "authVersion": _auth_version_for_doc(user_doc),
            "url": canonical_url,
            "article": article,
        })
        return jsonify({
            "ok": True,
            "articleUrl": canonical_url,
            "article_url": canonical_url,
            "article": article,
            "people": people,
            "warnings": warnings,
            "intervieweeExtraction": extraction_info,
            "extractionMode": extraction_status,
            "extractionToken": token,
        })
    except ArticleExtractionError as exc:
        return jsonify({"ok": False, **exc.to_dict()}), _extractor_error_status(exc.code)
    except Exception:
        return jsonify({"ok": False, "error": "The article could not be extracted right now."}), 502


@app.post("/api/save")
@require_role_at_least(ROLE_EDITOR)
def api_save():
    user_doc = _current_user_doc()
    payload = _request_payload()
    token = payload.get("extractionToken") if isinstance(payload.get("extractionToken"), str) else ""
    if not token or len(token) > 8192:
        return jsonify({"ok": False, "error": "Run article extraction again before saving."}), 400
    try:
        ticket = _extraction_serializer().loads(token, max_age=max(60, EXTRACTION_TOKEN_MAX_AGE_SECONDS))
    except SignatureExpired:
        return jsonify({"ok": False, "error": "This extraction expired. Run it again before saving."}), 410
    except BadSignature:
        return jsonify({"ok": False, "error": "This extraction token is invalid. Run extraction again."}), 400
    if not isinstance(ticket, dict) or ticket.get("v") != EXTRACTION_TOKEN_VERSION:
        return jsonify({"ok": False, "error": "This extraction token is invalid. Run extraction again."}), 400

    workspace_id = _workspace_id_for_user(user_doc)
    if (
        str(ticket.get("workspaceId") or "") != workspace_id
        or str(ticket.get("userId") or "") != str(user_doc.get("_id") or "")
        or int(ticket.get("authVersion") or 0) != _auth_version_for_doc(user_doc)
    ):
        return jsonify({"ok": False, "error": "This extraction does not belong to your current session."}), 403

    raw_article_url = payload.get("articleUrl") if isinstance(payload.get("articleUrl"), str) else ""
    try:
        canonical_url = canonicalize_article_url(raw_article_url)
    except ArticleExtractionError as exc:
        return jsonify({"ok": False, **exc.to_dict()}), 400
    if canonical_url != str(ticket.get("url") or ""):
        return jsonify({"ok": False, "error": "The article URL changed after extraction. Run extraction again."}), 409
    workspace = _workspace_for_user(user_doc)
    if not workspace or not _validated_publication_url(canonical_url, workspace):
        return jsonify({"ok": False, "error": "The article URL must use the workspace publication domain."}), 422

    try:
        people = _validated_extraction_people(payload.get("people"))
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400

    article = _sanitized_extracted_article(ticket.get("article"), canonical_url)
    article_doc = _upsert_extracted_article(workspace_id, canonical_url, article, user_doc)
    now = datetime.now(timezone.utc)
    created = 0
    updated = 0
    saved_records = []
    for person in people:
        key = {
            "workspaceId": workspace_id,
            "urlNorm": canonical_url,
            "firstNameNorm": person["firstNameNorm"],
            "lastNameNorm": person["lastNameNorm"],
        }
        update = {
            "$set": {
                "url": canonical_url,
                "articleUrl": canonical_url,
                "firstName": person["firstName"],
                "lastName": person["lastName"],
                "name": f"{person['firstName']} {person['lastName']}",
                "grade": person["grade"],
                "house": person["house"],
                "dateAdded": person["dateAdded"],
                "addedBy": _user_display_name(user_doc),
                "addedByUserId": str(user_doc.get("_id") or ""),
                "addedByEmail": normalize_email(user_doc.get("email") or ""),
                "updatedAt": now,
            },
            "$setOnInsert": {"createdAt": now},
        }
        try:
            result = interviews_col.update_one(key, update, upsert=True)
        except DuplicateKeyError:
            result = interviews_col.update_one(key, {"$set": update["$set"]})
        if result.upserted_id is not None:
            created += 1
        else:
            updated += 1
        saved = interviews_col.find_one(key)
        if saved:
            saved_records.append(_interview_to_api(saved))

    return jsonify({
        "ok": True,
        "created": created,
        "updated": updated,
        "saved": len(saved_records),
        "message": f"{len(saved_records)} interviewee record{'s' if len(saved_records) != 1 else ''} saved.",
        "article": _article_to_api(article_doc),
        "people": saved_records,
    })


NAME_HEADER_ALIASES = {
    "firstName": {"firstname", "first", "givenname", "given"},
    "lastName": {"lastname", "last", "surname", "familyname", "family"},
    "grade": {"grade", "gradelevel", "classyear", "graduationyear", "year"},
    "house": {"house", "academy", "program"},
    "type": {"type", "persontype", "category", "studentstaff", "status"},
    "email": {"email", "emailaddress", "schoolemail"},
    "title": {"title", "jobtitle", "position"},
}
MAX_NAMES_ROWS = 25000
MAX_NAMES_COLUMNS = 100
MAX_NAMES_CELL_CHARS = 500
MAX_XLSX_ARCHIVE_ENTRIES = 1000
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


def _upload_cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    return str(value).strip()


def _validated_upload_row(row) -> list[str]:
    values = list(row)
    if len(values) > MAX_NAMES_COLUMNS:
        raise ValueError(f"Names files can contain at most {MAX_NAMES_COLUMNS} columns.")
    normalized = [_upload_cell_text(value) for value in values]
    if any(len(value) > MAX_NAMES_CELL_CHARS for value in normalized):
        raise ValueError(f"Names file cells can contain at most {MAX_NAMES_CELL_CHARS} characters.")
    return normalized


def _validate_xlsx_archive(raw: bytes):
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
                raise ValueError("The Excel workbook contains too many archive entries.")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValueError("Encrypted Excel workbooks are not supported.")
            if sum(max(0, int(entry.file_size)) for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError("The expanded Excel workbook is too large.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Could not read this Excel file. Save it as a valid .xlsx workbook.") from exc


def _normalized_upload_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", _upload_cell_text(value).lower())


def _read_names_upload(upload) -> tuple[list[str], list[list[str]], str]:
    filename = secure_filename(upload.filename or "")[:180]
    suffix = Path(filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ValueError("Choose a CSV or .xlsx Excel file.")

    raw = upload.read(MAX_NAMES_UPLOAD_BYTES + 1)
    if not raw:
        raise ValueError("The selected file is empty.")
    if len(raw) > MAX_NAMES_UPLOAD_BYTES:
        raise ValueError("The names file must be 10 MB or smaller.")

    rows = []
    if suffix == ".csv":
        try:
            decoded = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Could not read this CSV. Export it using UTF-8 encoding.") from exc
        try:
            for row_number, row in enumerate(csv.reader(StringIO(decoded), strict=True)):
                if row_number > MAX_NAMES_ROWS:
                    raise ValueError(f"Names files can contain at most {MAX_NAMES_ROWS:,} data rows.")
                rows.append(_validated_upload_row(row))
        except csv.Error as exc:
            raise ValueError("Could not parse this CSV file.") from exc
    else:
        _validate_xlsx_archive(raw)
        workbook = None
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            worksheet = workbook.active
            if int(worksheet.max_row or 0) > MAX_NAMES_ROWS + 1:
                raise ValueError(f"Names files can contain at most {MAX_NAMES_ROWS:,} data rows.")
            if int(worksheet.max_column or 0) > MAX_NAMES_COLUMNS:
                raise ValueError(f"Names files can contain at most {MAX_NAMES_COLUMNS} columns.")
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True)):
                if row_number > MAX_NAMES_ROWS:
                    raise ValueError(f"Names files can contain at most {MAX_NAMES_ROWS:,} data rows.")
                rows.append(_validated_upload_row(row))
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Could not read this Excel file. Save it as a valid .xlsx workbook.") from exc
        finally:
            if workbook is not None:
                workbook.close()

    if not rows or not any(rows[0]):
        raise ValueError("The file must begin with a header row.")
    return rows[0], rows[1:], filename


def _names_documents(headers: list[str], rows: list[list[str]]) -> tuple[list[dict], int, int]:
    headers = _validated_upload_row(headers)
    normalized_headers = [_normalized_upload_header(header) for header in headers]
    column_indexes = {}
    for field, aliases in NAME_HEADER_ALIASES.items():
        column_indexes[field] = next((index for index, header in enumerate(normalized_headers) if header in aliases), None)

    if column_indexes["firstName"] is None or column_indexes["lastName"] is None:
        raise ValueError("Include firstName and lastName columns in the header row.")

    documents = []
    skipped = 0
    duplicates = 0
    seen_names = set()
    mapped_indexes = {index for index in column_indexes.values() if index is not None}

    for raw_row in rows:
        row = _validated_upload_row(raw_row)
        def value_at(index):
            return _upload_cell_text(row[index]) if index is not None and index < len(row) else ""

        first_name = value_at(column_indexes["firstName"])
        last_name = value_at(column_indexes["lastName"])
        if not first_name or not last_name or len(first_name) > 120 or len(last_name) > 120:
            skipped += 1
            continue

        name_key = (first_name.casefold(), last_name.casefold())
        if name_key in seen_names:
            duplicates += 1
            skipped += 1
            continue
        seen_names.add(name_key)

        document = {
            "firstName": first_name,
            "lastName": last_name,
            "grade": value_at(column_indexes["grade"]),
            "house": value_at(column_indexes["house"]),
            "type": value_at(column_indexes["type"]),
            "email": value_at(column_indexes["email"]),
            "title": value_at(column_indexes["title"]),
        }
        details = {}
        for index, header in enumerate(headers):
            if index in mapped_indexes:
                continue
            key = re.sub(r"[.$]", "_", _upload_cell_text(header)).strip()[:80]
            value = value_at(index)
            if key and value:
                details[key] = value
        if details:
            document["details"] = details
        documents.append(document)

    if not documents:
        raise ValueError("No valid names were found. Fill in firstName and lastName for at least one row.")
    return documents, skipped, duplicates


def _active_names_query(workspace: dict) -> dict:
    workspace_id = str(workspace.get("publicId") or workspace.get("_id") or "").strip()
    metadata = workspace.get("namesDatabase") if isinstance(workspace.get("namesDatabase"), dict) else {}
    active_batch_id = str(metadata.get("activeBatchId") or "").strip()
    query = {"workspaceId": workspace_id}
    if active_batch_id:
        query["uploadBatchId"] = active_batch_id
    return query


@app.get("/api/admin/names")
@require_roles(ROLE_ADMIN)
def api_admin_names():
    workspace = _workspace_for_user(_current_user_doc()) or {}
    metadata = workspace.get("namesDatabase") if isinstance(workspace.get("namesDatabase"), dict) else {}
    return jsonify({
        "ok": True,
        "count": names_col.count_documents(_active_names_query(workspace)),
        "updatedAt": str(metadata.get("updatedAt") or ""),
        "updatedBy": str(metadata.get("updatedBy") or ""),
        "sourceFile": str(metadata.get("sourceFile") or ""),
        "fields": ["firstName", "lastName", "grade", "house", "type"],
    })


@app.post("/api/admin/names/upload")
@require_roles(ROLE_ADMIN)
def api_admin_upload_names():
    upload = request.files.get("file")
    if not upload:
        return jsonify({"ok": False, "error": "Choose a names file to upload."}), 400

    current_user = _current_user_doc()
    workspace = _workspace_for_user(current_user)
    if not workspace:
        return jsonify({"ok": False, "error": "Workspace not found."}), 404

    try:
        headers, rows, filename = _read_names_upload(upload)
        documents, skipped, duplicates = _names_documents(headers, rows)
        uploaded_at = _now_iso()
        uploaded_by = _user_display_name(current_user)
        workspace_id = str(workspace.get("publicId") or workspace.get("_id") or "").strip()
        upload_batch_id = secrets.token_urlsafe(24)
        for document in documents:
            document["workspaceId"] = workspace_id
            document["uploadBatchId"] = upload_batch_id
            document["uploadedAt"] = uploaded_at
            document["uploadedBy"] = uploaded_by

        if not users_col.find_one({"_id": current_user["_id"], "workspaceId": workspace_id, "role": ROLE_ADMIN}):
            return jsonify({"ok": False, "error": "Admin access changed. Refresh and try again."}), 403

        metadata = {
            "activeBatchId": upload_batch_id,
            "count": len(documents),
            "skipped": skipped,
            "duplicates": duplicates,
            "sourceFile": filename,
            "updatedAt": uploaded_at,
            "updatedBy": uploaded_by,
        }
        batch_inserted = False
        batch_activated = False
        try:
            batch_inserted = True
            names_col.insert_many(documents, ordered=True)
            result = workspaces_col.update_one({"_id": workspace["_id"]}, {"$set": {
                "namesDatabase": metadata,
                "updatedAt": uploaded_at,
            }})
            if result.matched_count == 0:
                raise RuntimeError("Workspace changed while activating the names batch.")
            batch_activated = True

            active_workspace = workspaces_col.find_one({"_id": workspace["_id"]}) or {}
            active_metadata = active_workspace.get("namesDatabase") if isinstance(active_workspace.get("namesDatabase"), dict) else {}
            if str(active_metadata.get("activeBatchId") or "") == upload_batch_id:
                try:
                    names_col.delete_many({
                        "workspaceId": workspace_id,
                        "uploadBatchId": {"$ne": upload_batch_id},
                        "uploadedAt": {"$lt": uploaded_at},
                    })
                except Exception:
                    pass
        except Exception:
            if batch_inserted and not batch_activated:
                names_col.delete_many({"workspaceId": workspace_id, "uploadBatchId": upload_batch_id})
            raise

        return jsonify({"ok": True, **metadata, "collection": NAMES_COLLECTION})
    except ValueError as exc:
        return jsonify({"ok": False, "error": str(exc)}), 400
    except Exception:
        return jsonify({"ok": False, "error": "Could not replace the names database. The existing list was kept."}), 500


if __name__ == "__main__":
    debug_enabled = FLASK_ENV != "production" and os.getenv("FLASK_DEBUG", "0").strip().lower() in {"1", "true", "yes", "on"}
    app.run(
        host=os.getenv("V3_AUTH_HOST", "127.0.0.1"),
        port=int(os.getenv("V3_AUTH_PORT", "5003")),
        debug=debug_enabled,
        use_reloader=False,
    )
